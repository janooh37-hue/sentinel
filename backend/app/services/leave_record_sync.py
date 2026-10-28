"""Append a date-corrected record revision without resubmitting its approval chain.

The leave service owns the transaction. This path deliberately renders artifacts,
not another leave, and moves the existing assignments rather than issuing new work.
"""

from __future__ import annotations

import json
import logging
import re
import shutil
from collections import Counter
from contextlib import ExitStack
from copy import deepcopy
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.errors import ValidationFailedError
from app.config import get_settings
from app.core import signature_layout
from app.core.docx_engine import aztec_corner_for
from app.db.models import (
    AuditLog,
    Book,
    BookRevisionAccess,
    BookVersion,
    Document,
    Employee,
    Leave,
    SignatureArtifactRevision,
    Submitter,
    User,
)
from app.services import (
    artifact_service,
    book_service,
    document_service,
    included_papers_service,
    signature_placement_service,
    user_signature_service,
)

log = logging.getLogger(__name__)

_DATE_KEYS = {
    "Leave Application Form": ("start_date", "end_date", "total_days"),
    "Administrative Leave Form": ("start_date", "end_date", "duration"),
    "Leave Permit Form": ("date",),
}
_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")
_PRINTED_DATE = re.compile(r"\b\d{2}[/-]\d{2}[/-]\d{4}\b")


class _Unreproducible(Exception):
    """The record's paper cannot be re-rendered as it was; skip, never print it differently."""


def _absolute(path: str, data_dir: Path) -> Path:
    candidate = Path(path)
    return candidate if candidate.is_absolute() else data_dir / candidate


def _generation_version(book: Book) -> BookVersion:
    """The newest version a generation (not a leave edit) produced."""
    return next(
        (version for version in reversed(book.versions) if version.trigger != "leave-edit"),
        book.versions[0],
    )


def _signature_shape(path: Path) -> tuple[Counter[str], int]:
    """Marked signature roles, and the count of other images, in a rendered DOCX body."""
    try:
        marked = Counter[str](d.role for d in signature_layout.inspect_signature_drawings(path))
    except signature_layout.SignatureIdentityError as exc:
        raise _Unreproducible("The paper's signature markers are malformed") from exc
    body = DocxDocument(str(path)).element.body
    unmarked = sum(
        1
        for docpr in body.iter(qn("wp:docPr"))
        if signature_layout.parse_marker_descr(docpr.get("descr")) is None
    )
    return marked, unmarked


def _printed_dates(path: Path) -> list[str]:
    """Every printed date, in document order."""
    body = DocxDocument(str(path)).element.body
    return _PRINTED_DATE.findall("\n".join(t.text or "" for t in body.iter(qn("w:t"))))


def _require_same_paper(
    original: Path, rendered: Path, old_dates: set[str], new_dates: set[str]
) -> None:
    """Only the leave dates may differ: same signatures, same request/issue dates."""
    old_marked, old_unmarked = _signature_shape(original)
    new_marked, new_unmarked = _signature_shape(rendered)
    if old_marked:
        same = old_marked == new_marked and old_unmarked == new_unmarked
    else:
        # Before signature markers (2026-09-15) every signature was an unmarked image.
        same = old_unmarked == new_unmarked + sum(new_marked.values())
    if not same:
        raise _Unreproducible("Re-rendered signatures differ from the record's paper")
    # Same template, same token order: each printed date either stays or is a
    # leave date that moved (a request date may equal a leave date).
    old_printed, new_printed = _printed_dates(original), _printed_dates(rendered)
    if len(old_printed) != len(new_printed) or any(
        old != new and (old not in old_dates or new not in new_dates)
        for old, new in zip(old_printed, new_printed, strict=True)
    ):
        raise _Unreproducible("Re-rendered request/issue dates differ from the record's paper")


def _recovered_snapshot(
    db: Session, fields: dict[str, Any], reference: Path, *, issued: datetime
) -> dict[str, Any]:
    """Issue date and signature choices of a paper rendered before ``_signing``
    existed, read from the paper itself; ``_require_same_paper`` verifies them."""
    roles = _signature_shape(reference)[0]
    submitter_id: int | None = None
    if roles["submitter"]:
        raw = str(fields.get("submitter_id") or "")
        candidate = int(raw) if raw.isdigit() else None
        if candidate is None or db.get(Submitter, candidate) is None:
            raise _Unreproducible("The paper's submitter signature cannot be attributed")
        submitter_id = candidate
    return {
        "issued": issued.strftime("%d/%m/%Y"),
        "employee": roles["employee"] > 0,
        "submitter_id": submitter_id,
    }


def _date_like(value: object, new: date) -> str:
    if isinstance(value, str):
        for fmt in _DATE_FORMATS:
            try:
                datetime.strptime(value, fmt)
            except ValueError:
                continue
            return new.strftime(fmt)
    raise ValueError("Unrecognized stored date format")


def _updated_fields(version: BookVersion, leave: Leave) -> dict[str, Any]:
    fields = deepcopy(version.fields or {})
    keys = _DATE_KEYS[version.template_id or ""]
    fields[keys[0]] = _date_like(fields.get(keys[0]), leave.start_date)
    if len(keys) > 1:
        fields[keys[1]] = _date_like(fields.get(keys[1]), leave.end_date)
    if len(keys) > 2:
        key = keys[2]
        previous_count = fields.get(key)
        if type(previous_count) not in (int, str):
            raise ValueError("Unrecognized stored day-count type")
        count = (leave.end_date - leave.start_date).days + 1
        fields[key] = str(count) if isinstance(previous_count, str) else count
    return fields


def sync_leave_record(
    db: Session,
    leave: Leave,
    *,
    old_start: date,
    old_end: date,
    old_days: int,
    actor: str | None,
) -> None:
    """Sync only the current leave owner's mapped form, atomically with Leaves.

    Duty Resumption can itself own the leave. An unmapped owner is audited and
    left untouched: we never select a different record or change ownership.
    Failed conversion/publication rolls back record mutations and removes only
    this attempt's files. The caller must roll back its outer leave transaction.
    """
    try:
        with ExitStack() as cleanup:
            with db.begin_nested():
                _sync_leave_record(
                    db,
                    leave,
                    old_start=old_start,
                    old_end=old_end,
                    old_days=old_days,
                    actor=actor,
                    cleanup=cleanup,
                )
            cleanup.pop_all()
    except Exception as exc:
        log.exception("record sync failed for leave %s", leave.id)
        raise ValidationFailedError(
            "GENERATION_PDF_FAILED",
            "The record PDF could not be regenerated; the leave edit was not saved",
        ) from exc


def _sync_leave_record(
    db: Session,
    leave: Leave,
    *,
    old_start: date,
    old_end: date,
    old_days: int,
    actor: str | None,
    cleanup: ExitStack,
) -> None:
    """Prepare safe field updates, then append files and workflow metadata."""
    owner = book_service.leave_owner(db, leave.id)
    if owner is None:
        return
    previous = max(owner.versions, key=lambda version: version.version_no)
    payload = {
        "ref_number": owner.ref_number,
        "leave_id": leave.id,
        "version_no": previous.version_no,
        "actor": actor,
        "from": {"start": str(old_start), "end": str(old_end), "days": old_days},
        "to": {"start": str(leave.start_date), "end": str(leave.end_date), "days": leave.days},
        "source": "leaves",
        "note": "edited from Leaves",
    }

    def audit(*, reason: str | None = None) -> None:
        db.add(
            AuditLog(
                actor=actor,
                action="book.leave_dates_not_synced" if reason else "book.leave_dates_synced",
                entity_type="book",
                entity_id=str(owner.id),
                payload=json.dumps({**payload, **({"reason": reason} if reason else {})}),
            )
        )

    if not previous.fields:
        audit(reason="Word-authored record has no structured fields")
        return
    if previous.template_id not in _DATE_KEYS:
        audit(reason="No date mapping for this template")
        return
    if previous.signature_revision > 1:
        audit(reason="Corrected signature placement cannot be replayed")
        return
    if previous.template_id == "Leave Permit Form" and leave.start_date != leave.end_date:
        audit(reason="A Leave Permit Form cannot represent a multi-day leave")
        return
    try:
        fields = _updated_fields(previous, leave)
    except ValueError as exc:
        audit(reason=str(exc))
        return
    original = db.get(Document, previous.document_id) if previous.document_id else None
    if original is None or not original.docx_path:
        audit(reason="Original document/output directory is unavailable")
        return
    if (
        previous.template_id == "Administrative Leave Form"
        and not str(fields.get("admin_leaves_this_month") or "").strip()
    ):
        # Generation counted this month's leaves at the time; that count is not stored.
        audit(reason="The paper's monthly administrative-leave count was not stored")
        return

    signer = None
    signature = None
    # Signed means an electronic sign_book artifact. An approved version whose
    # manager signature was embedded at generation re-renders through the
    # generated path below with the same embed flags.
    signed = bool(previous.signed_pdf_path or previous.signed_at)
    if signed:
        source = previous.signed_base_pdf_path or previous.signed_pdf_path or ""
        if "book_attachments" in Path(source.replace("\\", "/")).parts:
            audit(reason="A scanned signature cannot be replayed")
            return
        signer = db.get(User, previous.signed_by_user_id) if previous.signed_by_user_id else None
        signature = user_signature_service.resolve_signature(signer) if signer else None
        if signer is None or signature is None or not signature.is_file():
            audit(reason="Previous signer or saved signature cannot be resolved")
            return

    settings = get_settings()

    def relative(path: Path | None) -> str | None:
        if path is None:
            return None
        try:
            return str(path.relative_to(settings.data_dir))
        except ValueError:
            return str(path)

    output = _absolute(original.docx_path, settings.data_dir)
    out_dir = output.parent
    employee = db.get(Employee, original.employee_id) if original.employee_id else None
    keys = _DATE_KEYS[previous.template_id][:2]
    old_dates = {str(previous.fields.get(key)) for key in keys}
    new_dates = {str(fields.get(key)) for key in keys}

    # The paper's frozen issue date and signature choices. Pre-snapshot papers
    # recover them for rendering only: persisting a snapshot would enroll the
    # record in the sign gate (book_service.is_enrolled).
    snapshot = fields.get(document_service._SIGNING_KEY)
    snapshot = dict(snapshot) if isinstance(snapshot, dict) else None
    signed_reference: Path | None = None
    signed_snapshot: dict[str, Any] | None = None
    try:
        render_snapshot = snapshot or _recovered_snapshot(
            db, fields, output, issued=_generation_version(owner).created_at
        )
        if signed:
            retained = (
                db.scalar(
                    select(SignatureArtifactRevision.docx_path).where(
                        SignatureArtifactRevision.version_id == previous.id,
                        SignatureArtifactRevision.revision == previous.signature_revision,
                    )
                )
                if previous.signature_revision
                else None
            )
            signed_reference = _absolute(retained, settings.data_dir) if retained else None
            if snapshot is None:
                if signed_reference is None or previous.signed_at is None:
                    raise _Unreproducible(
                        "The signed paper predates snapshots and was not retained"
                    )
                signed_snapshot = _recovered_snapshot(
                    db,
                    fields,
                    signed_reference,
                    issued=previous.signed_at.replace(tzinfo=UTC).astimezone(),
                )
    except _Unreproducible as skip:
        audit(reason=str(skip))
        return

    raw_submitter = render_snapshot.get("submitter_id")
    submitter_id = (
        raw_submitter
        if isinstance(raw_submitter, int) and db.get(Submitter, raw_submitter) is not None
        else None
    )
    embed = {
        "employee": render_snapshot.get("employee") is True,
        "manager": previous.manager_sig_embedded,
    }
    data = document_service._build_template_data(
        db,
        template_id=previous.template_id,
        employee=employee,
        employee_id=original.employee_id,
        fields={**fields, document_service._SIGNING_KEY: render_snapshot},
        manager_id=owner.doc_manager_id,
        submitter_id=submitter_id,
        embed_signature=embed,
        current_user=None,
    )
    # Generation left ``today`` to the renderer's clock, which defaults it to
    # ``date``; pin both to the paper instead of the edit day.
    data.setdefault("today", data.get("date") or render_snapshot["issued"])
    primary_data = data
    if document_service._submitter_signs_employee_cell(previous.template_id, submitter_id, embed):
        assert submitter_id is not None
        primary_data = {
            **data,
            "employee_sig_path": document_service._submitter_sign_path(db, submitter_id),
        }
    try:
        with db.begin_nested():
            submission_id = str(uuid4())
            actor_user = db.scalar(select(User).where(User.email == actor)) if actor else None

            def render(
                template_id: str, template_data: dict[str, Any], role: str
            ) -> tuple[Document, artifact_service.ArtifactResult]:
                artifact = artifact_service.produce_from_template(
                    template_id=template_id,
                    data=template_data,
                    destination=out_dir
                    / document_service._build_docx_filename(
                        template_id, (employee.name_en if employee else "") or "", datetime.now()
                    ),
                    stamps=artifact_service.StampPlan(
                        grade=owner.grade,
                        reference=owner.ref_number,
                        header_reference=owner.classification_code is None,
                        aztec_corner=aztec_corner_for(template_id),
                    ),
                    converter=document_service.convert_docx_to_pdf,
                )
                cleanup.callback(artifact_service.cleanup_created, artifact, allowed_root=out_dir)
                if artifact.conversion.pdf_path is None:
                    raise RuntimeError(f"PDF conversion failed for {template_id}")
                document = Document(
                    employee_id=original.employee_id,
                    template_id=template_id,
                    ref_number=owner.ref_number,
                    docx_path=relative(artifact.docx_path),
                    pdf_path=relative(artifact.conversion.pdf_path),
                    leave_id=leave.id if role == "primary" else None,
                    submission_id=submission_id,
                    role=role,
                )
                db.add(document)
                db.flush()
                return document, artifact

            primary, artifact = render(previous.template_id, primary_data, "primary")
            _require_same_paper(output, artifact.docx_path, old_dates, new_dates)
            companion = (
                db.scalar(
                    select(Document).where(
                        Document.submission_id == original.submission_id,
                        Document.role == "companion",
                        Document.template_id == "Leave Undertaking",
                    )
                )
                if original.submission_id
                else None
            )
            if companion is not None:
                _, companion_artifact = render("Leave Undertaking", data, "companion")
                if companion.docx_path:
                    _require_same_paper(
                        _absolute(companion.docx_path, settings.data_dir),
                        companion_artifact.docx_path,
                        old_dates,
                        new_dates,
                    )

            document_service._freeze_legacy_current_approval_context(db, owner, previous)
            version = BookVersion(
                book=owner,
                version_no=previous.version_no + 1,
                document_id=primary.id,
                template_id=previous.template_id,
                fields=fields,
                approval_context=deepcopy(previous.approval_context),
                trigger="leave-edit",
                status=previous.status,
                manager_sig_embedded=previous.manager_sig_embedded,
                signed_by_user_id=previous.signed_by_user_id,
                signed_at=previous.signed_at,
                signature_revision=0,
                created_by_user_id=actor_user.id if actor_user else None,
            )
            db.add(version)
            owner.doc_path = primary.docx_path
            db.flush()
            assert artifact.conversion.pdf_path is not None
            package = included_papers_service.publish_generated_package(
                db,
                owner,
                version,
                primary,
                artifact.conversion.pdf_path,
                invalidate_revision=True,
                data_dir=settings.data_dir,
            )
            cleanup.callback(
                included_papers_service.cleanup_published_package,
                package,
                allowed_root=settings.data_dir / "book_packages" / str(owner.id),
            )
            if signed:
                assert signer is not None and signature is not None
                names = [signer.display_name] if signer.display_name else []
                signer_employee = (
                    db.get(Employee, signer.employee_id) if signer.employee_id else None
                )
                if signer_employee is not None:
                    names.extend(n for n in (signer_employee.name_ar, signer_employee.name_en) if n)
                signed_artifact = document_service.render_signed_artifact(
                    db,
                    version=version,
                    signer_signature_path=str(signature),
                    signer_names=names,
                    signer_user_id=signer.id,
                    output_dir=out_dir,
                    signing_snapshot=signed_snapshot,
                )
                cleanup.callback(
                    artifact_service.cleanup_created, signed_artifact, allowed_root=out_dir
                )
                if signed_reference is not None:
                    _require_same_paper(
                        signed_reference, signed_artifact.docx_path, old_dates, new_dates
                    )
                signed_primary = signed_artifact.conversion.pdf_path
                if signed_primary is None:
                    raise RuntimeError("Signed PDF conversion failed")
                signed_path = included_papers_service.publish_signed_package(
                    db,
                    owner,
                    version,
                    signed_primary,
                    physical_scan=False,
                    data_dir=settings.data_dir,
                )
                published = settings.data_dir / signed_path
                for path in {version.signed_base_pdf_path, signed_path}:
                    if path is not None:
                        cleanup.callback((settings.data_dir / path).unlink, missing_ok=True)
                signature_dir = settings.data_dir / "signature_artifacts" / str(version.id)
                if not signature_dir.exists():
                    cleanup.callback(shutil.rmtree, signature_dir, ignore_errors=True)
                signature_placement_service.capture_initial_revision(
                    db,
                    version=version,
                    source_kind=signature_placement_service.SOURCE_KIND_APPROVAL,
                    signer_user_id=signer.id,
                    docx_path=signed_artifact.docx_path,
                    primary_pdf_path=signed_artifact.conversion.pdf_path,
                    published_pdf_path=published,
                )

            # Keep exactly one live assignment/decision chain. Retained decisions expose
            # the previous revision's history after its actual steps move to the new one.
            for step in list(previous.approval_steps):
                book_service.retain_revision_access(db, previous, step)
            db.flush()
            for grant in list(previous.revision_access):
                db.add(
                    BookRevisionAccess(
                        version_id=version.id,
                        user_id=grant.user_id,
                        kind=grant.kind,
                        state=grant.state,
                        note=grant.note,
                        assigned_at=grant.assigned_at,
                        decided_at=grant.decided_at,
                        revoked_at=grant.revoked_at,
                        revoked_by_user_id=grant.revoked_by_user_id,
                        revocation_reason=grant.revocation_reason,
                    )
                )
            for step in list(previous.approval_steps):
                step.version = version
            db.flush()
    except _Unreproducible as skip:
        # Nothing of this attempt survives: no revision, no files.
        cleanup.close()
        audit(reason=str(skip))
        return
    payload["version_no"] = version.version_no
    audit()
