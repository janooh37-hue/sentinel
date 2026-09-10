"""Secure, in-memory parsing and generation for vehicle import workbooks."""

from __future__ import annotations

import io
import json
import mimetypes
import posixpath
import re
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath
from typing import Any, TypedDict, cast
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.datetime import CALENDAR_WINDOWS_1900, from_excel
from PIL import Image, UnidentifiedImageError

from app.core.extraction.vehicle_licence import normalize_arabic_digits
from app.core.vehicle_plate import parse_plate

MAX_COMPRESSED_BYTES = 25 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_ZIP_MEMBERS = 5_000
MAX_ROWS = 1_000
MAX_IMAGES = 1_000

CANONICAL_COLUMNS = (
    "plate_code",
    "plate_number",
    "traffic_code",
    "type_ar",
    "type_en",
    "class_ar",
    "class_en",
    "vin",
    "site",
    "make",
    "model",
    "model_year",
    "colour",
    "license_start",
    "license_expiry",
    "insurance_expiry",
    "inmate_capacity",
    "passenger_capacity",
    "accessories_ar",
    "accessories_en",
    "notes_ar",
    "notes_en",
    "contract_note_ar",
    "contract_note_en",
)


class _VehicleClassAliases(TypedDict):
    ar: dict[str, str]
    en: dict[str, str]


class _VehicleClassCatalog(TypedDict):
    presets: list[dict[str, str]]
    aliases: _VehicleClassAliases


_VEHICLE_CLASS_CATALOG = cast(
    _VehicleClassCatalog,
    json.loads(
        Path(__file__).with_name("vehicle_classes.json").read_text(encoding="utf-8")
    ),
)
_VEHICLE_CLASSES = tuple(
    (preset["ar"], preset["en"]) for preset in _VEHICLE_CLASS_CATALOG["presets"]
)
_VEHICLE_CLASS_BY_AR = {ar: (ar, en) for ar, en in _VEHICLE_CLASSES}
_VEHICLE_CLASS_BY_EN = {en.casefold(): (ar, en) for ar, en in _VEHICLE_CLASSES}
_VEHICLE_CLASS_ALIASES_AR = _VEHICLE_CLASS_CATALOG["aliases"]["ar"]
_VEHICLE_CLASS_ALIASES_EN = {
    alias.casefold(): canonical.casefold()
    for alias, canonical in _VEHICLE_CLASS_CATALOG["aliases"]["en"].items()
}

_CREATE_REQUIRED_COLUMNS = frozenset(
    {
        "plate_number",
        "traffic_code",
        "type_ar",
        "type_en",
        "class_ar",
        "class_en",
        "license_start",
        "license_expiry",
    }
)
_CANONICAL_COLUMN_GUIDANCE = {
    "plate_code": (
        "Plate category code before the plate number (1 to 3 digits).",
        "رمز فئة اللوحة الذي يسبق رقمها (من رقم إلى ثلاثة أرقام).",
    ),
    "plate_number": (
        "Official plate number (1 to 6 digits); retain significant leading zeros.",
        "رقم لوحة المركبة الرسمي (من رقم إلى ستة أرقام)؛ احتفظ بالأصفار البادئة المهمة.",
    ),
    "traffic_code": (
        "Vehicle traffic-file code (4 to 12 digits).",
        "الرمز المروري للمركبة (من 4 إلى 12 رقمًا).",
    ),
    "type_ar": (
        "Vehicle type name in Arabic.",
        "اسم نوع المركبة باللغة العربية.",
    ),
    "type_en": (
        "Vehicle type name in English.",
        "اسم نوع المركبة باللغة الإنجليزية.",
    ),
    "class_ar": (
        "Vehicle class in Arabic; use a preset or an exact custom-class name.",
        "فئة المركبة باللغة العربية؛ استخدم فئة معتمدة أو اسم فئة مخصصة دقيقًا.",
    ),
    "class_en": (
        "Vehicle class in English; use the matching preset or exact custom-class name.",
        "فئة المركبة باللغة الإنجليزية؛ استخدم الفئة المعتمدة المطابقة أو اسم فئة مخصصة دقيقًا.",
    ),
    "vin": (
        "Vehicle identification number (VIN) or chassis number.",
        "رقم تعريف المركبة (VIN) أو رقم القاعدة.",
    ),
    "site": (
        "Optional workbook site reference; select the required assigned site during review.",
        "مرجع اختياري للموقع داخل الملف؛ اختر الموقع المعيّن المطلوب أثناء المراجعة.",
    ),
    "make": (
        "Vehicle manufacturer.",
        "الشركة المصنّعة للمركبة.",
    ),
    "model": (
        "Vehicle model.",
        "طراز المركبة.",
    ),
    "model_year": (
        "Four-digit vehicle model year.",
        "سنة صنع المركبة من أربعة أرقام.",
    ),
    "colour": (
        "Vehicle colour.",
        "لون المركبة.",
    ),
    "license_start": (
        "Licence validity start date.",
        "تاريخ بداية سريان ترخيص المركبة.",
    ),
    "license_expiry": (
        "Licence expiry date.",
        "تاريخ انتهاء ترخيص المركبة.",
    ),
    "insurance_expiry": (
        "Insurance expiry date.",
        "تاريخ انتهاء تأمين المركبة.",
    ),
    "inmate_capacity": (
        "Inmate capacity as a non-negative whole number.",
        "سعة النزلاء كعدد صحيح غير سالب.",
    ),
    "passenger_capacity": (
        "Passenger capacity as a non-negative whole number.",
        "سعة الركاب كعدد صحيح غير سالب.",
    ),
    "accessories_ar": (
        "Vehicle accessories description in Arabic.",
        "وصف ملحقات المركبة باللغة العربية.",
    ),
    "accessories_en": (
        "Vehicle accessories description in English.",
        "وصف ملحقات المركبة باللغة الإنجليزية.",
    ),
    "notes_ar": (
        "General vehicle notes in Arabic.",
        "ملاحظات المركبة العامة باللغة العربية.",
    ),
    "notes_en": (
        "General vehicle notes in English.",
        "ملاحظات المركبة العامة باللغة الإنجليزية.",
    ),
    "contract_note_ar": (
        "Contract details or note in Arabic.",
        "تفاصيل العقد أو ملاحظته باللغة العربية.",
    ),
    "contract_note_en": (
        "Contract details or note in English.",
        "تفاصيل العقد أو ملاحظته باللغة الإنجليزية.",
    ),
}

_LEGACY_COLUMNS = {
    2: "combined_plate",
    3: "traffic_code",
    4: "license_start",
    5: "license_expiry",
    6: "type_ar",
    7: "vin",
    8: "class_ar",
    9: "inmate_capacity",
    10: "passenger_capacity",
    11: "accessories_ar",
    12: "notes_ar",
    13: "contract_note_ar",
}
_DATE_FIELDS = frozenset({"license_start", "license_expiry", "insurance_expiry"})
_INTEGER_FIELDS = frozenset({"model_year", "inmate_capacity", "passenger_capacity"})
_IDENTIFIER_FIELDS = frozenset({"plate_code", "plate_number", "traffic_code"})
_BLANK_MARKERS = frozenset({"", "-", "—"})
_LEGACY_TITLE_RE = re.compile(r"كشف\s+تفصيلي\s+لمركبات", re.IGNORECASE)
_DTD_RE = re.compile(br"<!\s*(?:DOCTYPE|ENTITY)\b", re.IGNORECASE)
_DRIVE_RE = re.compile(r"^[A-Za-z]:")
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_VML_NS = "urn:schemas-microsoft-com:vml"
_OFFICE_NS = "urn:schemas-microsoft-com:office:office"
_EXCEL_NS = "urn:schemas-microsoft-com:office:excel"


class VehicleXlsxError(ValueError):
    """An upload-level workbook rejection with a stable API error code."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True, slots=True)
class VehicleXlsxIssue:
    row_id: str | None
    field: str | None
    code: str
    message: str


@dataclass(frozen=True, slots=True)
class VehicleXlsxSection:
    id: str
    sheet: str
    title: str


@dataclass(slots=True)
class VehicleXlsxRow:
    row_id: str
    section_id: str
    sheet: str
    row_number: int
    raw: dict[str, str | None]
    values: dict[str, str | int | None]
    image_ids: list[str] = field(default_factory=list)


@dataclass(frozen=True, slots=True)
class VehicleXlsxImage:
    image_id: str
    row_id: str | None
    original_name: str
    media_type: str
    kind: str | None
    data: bytes


@dataclass(slots=True)
class ParsedVehicleWorkbook:
    layout: str
    sections: list[VehicleXlsxSection]
    rows: list[VehicleXlsxRow]
    images: list[VehicleXlsxImage]
    warnings: list[VehicleXlsxIssue]


def _bad_file(message: str) -> VehicleXlsxError:
    return VehicleXlsxError("VEHICLE_IMPORT_BAD_FILE", message)


def _safe_member_name(name: str) -> bool:
    path = PurePosixPath(name)
    return bool(
        name
        and "\\" not in name
        and not name.startswith("/")
        and not _DRIVE_RE.match(name)
        and ".." not in path.parts
    )


def _parse_xml(raw: bytes, *, part: str) -> ET.Element:
    if _DTD_RE.search(raw):
        raise _bad_file(f"Unsafe XML declarations are not allowed ({part}).")
    try:
        return ET.fromstring(raw)
    except ET.ParseError as exc:
        raise _bad_file(f"Corrupt XML part: {part}.") from exc


def _validate_package(data: bytes) -> dict[str, bytes]:
    if not data or len(data) > MAX_COMPRESSED_BYTES:
        code = "VEHICLE_IMPORT_TOO_LARGE" if len(data) > MAX_COMPRESSED_BYTES else "VEHICLE_IMPORT_BAD_FILE"
        message = (
            f"Workbook exceeds {MAX_COMPRESSED_BYTES // (1024 * 1024)} MiB."
            if code == "VEHICLE_IMPORT_TOO_LARGE"
            else "The uploaded workbook is empty."
        )
        raise VehicleXlsxError(code, message)
    source = io.BytesIO(data)
    if not zipfile.is_zipfile(source):
        raise _bad_file("The upload is not a valid .xlsx package.")
    source.seek(0)
    try:
        with zipfile.ZipFile(source) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS:
                raise VehicleXlsxError(
                    "VEHICLE_IMPORT_TOO_LARGE",
                    f"Workbook contains more than {MAX_ZIP_MEMBERS} package members.",
                )
            expanded = 0
            image_count = 0
            xml_parts: dict[str, bytes] = {}
            for member in members:
                if not _safe_member_name(member.filename):
                    raise _bad_file("The workbook contains an unsafe archive path.")
                if member.flag_bits & 0x1:
                    raise _bad_file("Encrypted workbooks are not supported.")
                expanded += member.file_size
                if expanded > MAX_EXPANDED_BYTES:
                    raise VehicleXlsxError(
                        "VEHICLE_IMPORT_TOO_LARGE",
                        f"Expanded workbook exceeds {MAX_EXPANDED_BYTES // (1024 * 1024)} MiB.",
                    )
                if member.filename.startswith("xl/media/") and not member.is_dir():
                    image_count += 1
                if member.filename.lower().endswith((".xml", ".rels", ".vml")):
                    raw = archive.read(member)
                    _parse_xml(raw, part=member.filename)
                    xml_parts[member.filename] = raw
            if image_count > MAX_IMAGES:
                raise VehicleXlsxError(
                    "VEHICLE_IMPORT_TOO_LARGE",
                    f"Workbook contains more than {MAX_IMAGES} embedded images.",
                )
            required = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"}
            if not required.issubset({member.filename for member in members}):
                raise _bad_file("The ZIP package is not an Excel workbook.")
            bad_member = archive.testzip()
            if bad_member is not None:
                raise _bad_file(f"Corrupt workbook package member: {bad_member}.")
            row_count = sum(
                1
                for name, raw in xml_parts.items()
                if name.startswith("xl/worksheets/") and name.endswith(".xml")
                for node in _parse_xml(raw, part=name).iter(f"{{{_MAIN_NS}}}row")
            )
            if row_count > MAX_ROWS:
                raise VehicleXlsxError(
                    "VEHICLE_IMPORT_TOO_LARGE",
                    f"Workbook contains more than {MAX_ROWS} rows.",
                )
            for name, raw in xml_parts.items():
                if not name.endswith(".rels"):
                    continue
                root = _parse_xml(raw, part=name)
                if any(rel.get("TargetMode", "").casefold() == "external" for rel in root):
                    raise _bad_file("External workbook relationships are not allowed.")
            return xml_parts
    except VehicleXlsxError:
        raise
    except (OSError, RuntimeError, ValueError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise _bad_file("The workbook package is corrupt or unsupported.") from exc


def _plain_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        try:
            decimal = Decimal(str(value))
            return format(decimal, "f")
        except InvalidOperation:
            return str(value)
    text = normalize_arabic_digits(str(value)).strip()
    return text or None


def _normalized_text(value: object) -> str | None:
    text = _plain_text(value)
    if text is None or text in _BLANK_MARKERS:
        return None
    return " ".join(text.split())


def _vehicle_class_preset(
    ar_value: object,
    en_value: object,
) -> tuple[str, str] | None:
    ar = _normalized_text(ar_value)
    if ar is not None:
        canonical_ar = _VEHICLE_CLASS_ALIASES_AR.get(ar, ar)
        if match := _VEHICLE_CLASS_BY_AR.get(canonical_ar):
            return match
    en = _normalized_text(en_value)
    if en is not None:
        normalized_en = en.casefold()
        canonical_en = _VEHICLE_CLASS_ALIASES_EN.get(normalized_en, normalized_en)
        if match := _VEHICLE_CLASS_BY_EN.get(canonical_en):
            return match
    return None


def _identifier(value: object) -> str | None:
    text = _normalized_text(value)
    if text is None:
        return None
    text = re.sub(r"\s+", "", text)
    if re.fullmatch(r"[+-]?\d+(?:\.0+)?[Ee][+-]?\d+", text):
        try:
            numeric = Decimal(text)
        except InvalidOperation:
            return text
        if numeric == numeric.to_integral_value():
            return format(numeric.quantize(Decimal(1)), "f")
    if re.fullmatch(r"\d+\.0+", text):
        return text.partition(".")[0]
    return text


def _integer(value: object) -> int | None:
    text = _normalized_text(value)
    if text is None:
        return None
    match = re.fullmatch(r"(\d+)(?:\s*(?:نزيل|نزلاء|راكب|ركاب|passengers?|inmates?))?", text, re.IGNORECASE)
    if match is None:
        raise ValueError("Expected a non-negative whole number.")
    return int(match.group(1))


def _date_value(value: object, *, epoch: datetime) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, int | float) and not isinstance(value, bool):
        try:
            parsed = from_excel(value, epoch=epoch)
        except (OverflowError, TypeError, ValueError) as exc:
            raise ValueError("Invalid Excel date serial.") from exc
        if isinstance(parsed, datetime):
            return parsed.date().isoformat()
        if isinstance(parsed, date):
            return parsed.isoformat()
        raise ValueError("Invalid Excel date serial.")
    text = _normalized_text(value)
    if text is None:
        return None
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        try:
            parsed = from_excel(float(text), epoch=epoch)
        except (OverflowError, TypeError, ValueError) as exc:
            raise ValueError("Invalid Excel date serial.") from exc
        return parsed.date().isoformat() if isinstance(parsed, datetime) else parsed.isoformat()
    iso_match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    day_first = re.fullmatch(r"(\d{1,2})[./\\-](\d{1,2})[./\\-](\d{4})", text)
    try:
        if iso_match is not None:
            return date(*(int(part) for part in iso_match.groups())).isoformat()
        if day_first is not None:
            day, month, year = (int(part) for part in day_first.groups())
            return date(year, month, day).isoformat()
    except ValueError as exc:
        raise ValueError("Invalid calendar date.") from exc
    raise ValueError("Expected YYYY-MM-DD or DD/MM/YYYY date format.")


def normalize_import_values(
    values: Mapping[str, object],
    *,
    epoch: datetime = CALENDAR_WINDOWS_1900,
) -> tuple[dict[str, str | int | None], list[tuple[str, str]]]:
    """Normalize editable import values without guessing invalid cells."""
    normalized: dict[str, str | int | None] = {}
    errors: list[tuple[str, str]] = []
    for field_name in CANONICAL_COLUMNS:
        value = values.get(field_name)
        try:
            if field_name in _DATE_FIELDS:
                normalized[field_name] = _date_value(value, epoch=epoch)
            elif field_name in _INTEGER_FIELDS:
                normalized[field_name] = _integer(value)
            elif field_name in _IDENTIFIER_FIELDS:
                normalized[field_name] = _identifier(value)
            elif field_name == "vin":
                text = _identifier(value)
                normalized[field_name] = text.upper() if text is not None else None
            else:
                normalized[field_name] = _normalized_text(value)
        except ValueError as exc:
            normalized[field_name] = None
            errors.append((field_name, str(exc)))

    class_preset = _vehicle_class_preset(
        normalized.get("class_ar"),
        normalized.get("class_en"),
    )
    if class_preset is not None:
        normalized["class_ar"], normalized["class_en"] = class_preset

    combined = values.get("combined_plate")
    if combined is not None and _normalized_text(combined) is not None:
        parsed = parse_plate(_plain_text(combined) or "")
        if parsed is None:
            errors.append(("plate_number", "Unrecognized plate number."))
            normalized["plate_code"] = None
            normalized["plate_number"] = None
        else:
            normalized["plate_code"], normalized["plate_number"] = parsed
    else:
        plate_number = normalized.get("plate_number")
        plate_code = normalized.get("plate_code")
        if isinstance(plate_number, str):
            parsed = parse_plate(plate_number)
            if parsed is None:
                errors.append(("plate_number", "Unrecognized plate number."))
            else:
                parsed_code, parsed_number = parsed
                if plate_code is not None and parsed_code is not None and plate_code != parsed_code:
                    errors.append(("plate_code", "Plate code conflicts with the combined plate."))
                else:
                    normalized["plate_code"] = plate_code or parsed_code
                    normalized["plate_number"] = parsed_number
        if plate_code is not None and (not isinstance(plate_code, str) or not re.fullmatch(r"\d{1,3}", plate_code)):
            errors.append(("plate_code", "Plate code must contain 1 to 3 digits."))
        if plate_number is not None and (
            not isinstance(normalized.get("plate_number"), str)
            or not re.fullmatch(r"\d{1,6}", str(normalized.get("plate_number")))
        ):
            errors.append(("plate_number", "Plate number must contain 1 to 6 digits."))
    traffic_code = normalized.get("traffic_code")
    if traffic_code is not None and (
        not isinstance(traffic_code, str) or not re.fullmatch(r"\d{4,12}", traffic_code)
    ):
        errors.append(("traffic_code", "Traffic code must contain 4 to 12 digits."))
    return normalized, errors


def _cell_raw(cell: Cell) -> str | None:
    return _plain_text(cell.value)


def _row_from_cells(
    *,
    worksheet: Any,
    row_number: int,
    section: VehicleXlsxSection,
    mapping: Mapping[int, str],
    epoch: datetime,
    layout: str,
) -> tuple[VehicleXlsxRow | None, list[VehicleXlsxIssue]]:
    raw_source: dict[str, object] = {}
    raw: dict[str, str | None] = {column: None for column in CANONICAL_COLUMNS}
    formula_fields: list[str] = []
    for column_number, field_name in mapping.items():
        cell = worksheet.cell(row=row_number, column=column_number)
        if cell.data_type == "f":
            formula_fields.append("plate_number" if field_name == "combined_plate" else field_name)
        raw_source[field_name] = cell.value
        if field_name in raw:
            raw[field_name] = _cell_raw(cell)
    if "combined_plate" in raw_source:
        raw["plate_number"] = _plain_text(raw_source["combined_plate"])
    if not any(_normalized_text(value) is not None for value in raw_source.values()):
        return None, []
    row_id = f"{section.id}-row-{row_number}"
    values, errors = normalize_import_values(raw_source, epoch=epoch)
    if layout == "legacy":
        if values.get("type_ar") is not None and values.get("type_en") is None:
            values["type_en"] = values["type_ar"]
        if values.get("class_ar") is not None and values.get("class_en") is None:
            values["class_en"] = values["class_ar"]
    issues = [
        VehicleXlsxIssue(
            row_id=row_id,
            field=field_name,
            code="VEHICLE_IMPORT_INVALID_FIELD",
            message="Formula cells are not allowed in imported fields.",
        )
        for field_name in formula_fields
    ]
    issues.extend(
        VehicleXlsxIssue(
            row_id=row_id,
            field=field_name,
            code="VEHICLE_IMPORT_INVALID_FIELD",
            message=message,
        )
        for field_name, message in errors
    )
    return (
        VehicleXlsxRow(
            row_id=row_id,
            section_id=section.id,
            sheet=worksheet.title,
            row_number=row_number,
            raw=raw,
            values=values,
        ),
        issues,
    )


def _legacy_sections(workbook: Any) -> tuple[list[VehicleXlsxSection], list[tuple[Any, int, int, VehicleXlsxSection]]]:
    sections: list[VehicleXlsxSection] = []
    spans: list[tuple[Any, int, int, VehicleXlsxSection]] = []
    for worksheet in workbook.worksheets:
        title_rows: list[tuple[int, str]] = []
        for row_number in range(1, worksheet.max_row + 1):
            title = " ".join(
                text
                for cell in worksheet[row_number]
                if (text := _normalized_text(cell.value)) is not None
            )
            if _LEGACY_TITLE_RE.search(title):
                title_rows.append((row_number, title))
        for index, (title_row, title) in enumerate(title_rows):
            section = VehicleXlsxSection(
                id=f"section-{len(sections) + 1}",
                sheet=worksheet.title,
                title=title,
            )
            sections.append(section)
            end = title_rows[index + 1][0] - 1 if index + 1 < len(title_rows) else worksheet.max_row
            spans.append((worksheet, title_row + 2, end, section))
    return sections, spans


def _standard_sections(
    workbook: Any,
) -> tuple[
    list[VehicleXlsxSection],
    list[tuple[Any, int, int, VehicleXlsxSection, dict[int, str]]],
    list[VehicleXlsxIssue],
]:
    sections: list[VehicleXlsxSection] = []
    spans: list[tuple[Any, int, int, VehicleXlsxSection, dict[int, str]]] = []
    warnings: list[VehicleXlsxIssue] = []
    canonical = set(CANONICAL_COLUMNS)
    for worksheet in workbook.worksheets:
        for row_number in range(1, min(worksheet.max_row, 25) + 1):
            headers = {
                column_number: (_normalized_text(worksheet.cell(row_number, column_number).value) or "")
                for column_number in range(1, worksheet.max_column + 1)
            }
            recognized = {value for value in headers.values() if value in canonical}
            if "plate_number" not in recognized or len(recognized) < 2:
                continue
            section = VehicleXlsxSection(
                id=f"section-{len(sections) + 1}",
                sheet=worksheet.title,
                title=worksheet.title,
            )
            sections.append(section)
            mapping: dict[int, str] = {}
            seen: set[str] = set()
            for column_number, header in headers.items():
                if not header:
                    continue
                if header not in canonical:
                    warnings.append(
                        VehicleXlsxIssue(
                            row_id=None,
                            field=header,
                            code="VEHICLE_IMPORT_INVALID_FIELD",
                            message=f"Unknown template column {header!r} was ignored.",
                        )
                    )
                elif header in seen:
                    warnings.append(
                        VehicleXlsxIssue(
                            row_id=None,
                            field=header,
                            code="VEHICLE_IMPORT_INVALID_FIELD",
                            message=f"Duplicate template column {header!r} was ignored.",
                        )
                    )
                else:
                    seen.add(header)
                    mapping[column_number] = header
            spans.append((worksheet, row_number + 1, worksheet.max_row, section, mapping))
            break
    return sections, spans, warnings


def _relationships(xml_parts: Mapping[str, bytes], part: str) -> dict[str, tuple[str, str | None]]:
    raw = xml_parts.get(part)
    if raw is None:
        return {}
    root = _parse_xml(raw, part=part)
    return {
        str(rel.get("Id")): (str(rel.get("Target")), rel.get("TargetMode"))
        for rel in root.findall(f"{{{_REL_NS}}}Relationship")
        if rel.get("Id") and rel.get("Target")
    }


def _part_target(source_part: str, target: str) -> str:
    if "\\" in target or _DRIVE_RE.match(target):
        raise _bad_file("The workbook contains an unsafe relationship target.")
    resolved = (
        posixpath.normpath(target.lstrip("/"))
        if target.startswith("/")
        else posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))
    )
    if not _safe_member_name(resolved):
        raise _bad_file("The workbook contains an unsafe relationship target.")
    return resolved

def _sheet_parts(xml_parts: Mapping[str, bytes]) -> list[tuple[str, str]]:
    workbook_root = _parse_xml(xml_parts["xl/workbook.xml"], part="xl/workbook.xml")
    rels = _relationships(xml_parts, "xl/_rels/workbook.xml.rels")
    result: list[tuple[str, str]] = []
    for sheet in workbook_root.findall(f".//{{{_MAIN_NS}}}sheet"):
        rel_id = sheet.get(f"{{{_DOC_REL_NS}}}id")
        name = sheet.get("name")
        if not rel_id or not name or rel_id not in rels:
            continue
        target, mode = rels[rel_id]
        if mode is not None:
            continue
        result.append((name, _part_target("xl/workbook.xml", target)))
    return result


def _image_type(name: str, data: bytes) -> str:
    guessed = mimetypes.guess_type(name)[0]
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
            detected = Image.MIME.get(image.format or "")
    except (OSError, UnidentifiedImageError) as exc:
        raise _bad_file(f"Embedded image {Path(name).name!r} is corrupt.") from exc
    return detected or guessed or "application/octet-stream"


def _normal_drawing_anchors(
    xml_parts: Mapping[str, bytes],
    archive: zipfile.ZipFile,
    drawing_part: str,
) -> list[tuple[int | None, str, str, bytes]]:
    raw = xml_parts.get(drawing_part)
    if raw is None:
        return []
    root = _parse_xml(raw, part=drawing_part)
    rel_part = posixpath.join(
        posixpath.dirname(drawing_part), "_rels", f"{posixpath.basename(drawing_part)}.rels"
    )
    rels = _relationships(xml_parts, rel_part)
    result: list[tuple[int | None, str, str, bytes]] = []
    anchor_tags = {
        f"{{{_DRAWING_NS}}}oneCellAnchor",
        f"{{{_DRAWING_NS}}}twoCellAnchor",
        f"{{{_DRAWING_NS}}}absoluteAnchor",
    }
    for anchor in (node for node in root if node.tag in anchor_tags):
        row_number: int | None = None
        from_node = anchor.find(f"{{{_DRAWING_NS}}}from")
        if from_node is not None:
            row_node = from_node.find(f"{{{_DRAWING_NS}}}row")
            if row_node is not None and row_node.text is not None:
                row_number = int(row_node.text) + 1
        if anchor.tag == f"{{{_DRAWING_NS}}}twoCellAnchor":
            to_node = anchor.find(f"{{{_DRAWING_NS}}}to")
            to_row = to_node.find(f"{{{_DRAWING_NS}}}row") if to_node is not None else None
            if (
                row_number is None
                or to_row is None
                or to_row.text is None
                or int(to_row.text) + 1 != row_number
            ):
                row_number = None
        blips = anchor.findall(f".//{{{_A_NS}}}blip")
        if len(blips) != 1:
            result.append((None, "embedded-image", "application/octet-stream", b""))
            continue
        rel_id = blips[0].get(f"{{{_DOC_REL_NS}}}embed")
        if rel_id is None or rel_id not in rels:
            result.append((None, "embedded-image", "application/octet-stream", b""))
            continue
        target, mode = rels[rel_id]
        if mode is not None:
            result.append((None, "embedded-image", "application/octet-stream", b""))
            continue
        image_part = _part_target(drawing_part, target)
        try:
            data = archive.read(image_part)
        except KeyError as exc:
            raise _bad_file("An embedded image relationship is broken.") from exc
        name = posixpath.basename(image_part)
        result.append((row_number, name, _image_type(name, data), data))
    return result


def _vml_anchors(
    xml_parts: Mapping[str, bytes],
    archive: zipfile.ZipFile,
    vml_part: str,
) -> list[tuple[int | None, str, str, bytes]]:
    raw = xml_parts.get(vml_part)
    if raw is None:
        return []
    root = _parse_xml(raw, part=vml_part)
    rel_part = posixpath.join(
        posixpath.dirname(vml_part), "_rels", f"{posixpath.basename(vml_part)}.rels"
    )
    rels = _relationships(xml_parts, rel_part)
    result: list[tuple[int | None, str, str, bytes]] = []
    for shape in root.findall(f".//{{{_VML_NS}}}shape"):
        fills = shape.findall(f"{{{_VML_NS}}}fill")
        image_nodes = shape.findall(f"{{{_VML_NS}}}imagedata")
        rel_ids = [
            rel_id
            for node in (*fills, *image_nodes)
            if (rel_id := node.get(f"{{{_OFFICE_NS}}}relid")) is not None
        ]
        if not rel_ids:
            continue
        client = shape.find(f"{{{_EXCEL_NS}}}ClientData")
        rows = client.findall(f"{{{_EXCEL_NS}}}Row") if client is not None else []
        row_number = (
            int(rows[0].text) + 1
            if len(rows) == 1 and rows[0].text is not None and rows[0].text.isdigit()
            else None
        )
        for rel_id in rel_ids:
            if rel_id not in rels:
                result.append((None, "embedded-image", "application/octet-stream", b""))
                continue
            target, mode = rels[rel_id]
            if mode is not None:
                result.append((None, "embedded-image", "application/octet-stream", b""))
                continue
            image_part = _part_target(vml_part, target)
            try:
                data = archive.read(image_part)
            except KeyError as exc:
                raise _bad_file("A VML image relationship is broken.") from exc
            name = posixpath.basename(image_part)
            result.append((row_number, name, _image_type(name, data), data))
    return result


def _extract_images(
    data: bytes,
    xml_parts: Mapping[str, bytes],
    *,
    row_lookup: Mapping[tuple[str, int], str],
    legacy_sheets: set[str],
) -> tuple[list[VehicleXlsxImage], list[VehicleXlsxIssue]]:
    images: list[VehicleXlsxImage] = []
    warnings: list[VehicleXlsxIssue] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for sheet_name, sheet_part in _sheet_parts(xml_parts):
            rel_part = posixpath.join(
                posixpath.dirname(sheet_part), "_rels", f"{posixpath.basename(sheet_part)}.rels"
            )
            rels = _relationships(xml_parts, rel_part)
            anchors: list[tuple[int | None, str, str, bytes]] = []
            for target, mode in rels.values():
                if mode is not None:
                    continue
                target_part = _part_target(sheet_part, target)
                if "/drawings/" not in target_part:
                    continue
                if target_part.endswith(".vml"):
                    anchors.extend(_vml_anchors(xml_parts, archive, target_part))
                elif target_part.endswith(".xml"):
                    anchors.extend(_normal_drawing_anchors(xml_parts, archive, target_part))
            for row_number, original_name, media_type, image_data in anchors:
                image_id = f"image-{len(images) + 1}"
                row_id = row_lookup.get((sheet_name, row_number)) if row_number is not None else None
                kind = "license" if sheet_name in legacy_sheets else None
                if row_id is None:
                    warnings.append(
                        VehicleXlsxIssue(
                            row_id=None,
                            field=None,
                            code="VEHICLE_IMPORT_UNASSIGNED_IMAGE",
                            message=(
                                f"Embedded image {original_name!r} could not be resolved to one import row."
                            ),
                        )
                    )
                images.append(
                    VehicleXlsxImage(
                        image_id=image_id,
                        row_id=row_id,
                        original_name=original_name,
                        media_type=media_type,
                        kind=kind,
                        data=image_data,
                    )
                )
    if len(images) > MAX_IMAGES:
        raise VehicleXlsxError(
            "VEHICLE_IMPORT_TOO_LARGE",
            f"Workbook contains more than {MAX_IMAGES} embedded image anchors.",
        )
    return images, warnings


def parse_vehicle_workbook(data: bytes) -> ParsedVehicleWorkbook:
    """Parse a legacy or canonical vehicle workbook without extracting its ZIP."""
    xml_parts = _validate_package(data)
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=False, keep_links=False)
    except Exception as exc:
        raise _bad_file("The workbook could not be opened.") from exc

    sections, legacy_spans = _legacy_sections(workbook)
    rows: list[VehicleXlsxRow] = []
    warnings: list[VehicleXlsxIssue] = []
    legacy_sheets: set[str] = set()
    if sections:
        layout = "legacy"
        for worksheet, start, end, section in legacy_spans:
            legacy_sheets.add(worksheet.title)
            for row_number in range(start, end + 1):
                parsed_row, issues = _row_from_cells(
                    worksheet=worksheet,
                    row_number=row_number,
                    section=section,
                    mapping=_LEGACY_COLUMNS,
                    epoch=workbook.epoch,
                    layout=layout,
                )
                if parsed_row is not None:
                    rows.append(parsed_row)
                    warnings.extend(issues)
    else:
        layout = "standard"
        sections, standard_spans, header_warnings = _standard_sections(workbook)
        warnings.extend(header_warnings)
        for worksheet, start, end, section, mapping in standard_spans:
            for row_number in range(start, end + 1):
                parsed_row, issues = _row_from_cells(
                    worksheet=worksheet,
                    row_number=row_number,
                    section=section,
                    mapping=mapping,
                    epoch=workbook.epoch,
                    layout=layout,
                )
                if parsed_row is not None:
                    rows.append(parsed_row)
                    warnings.extend(issues)
    workbook.close()
    if not sections:
        raise _bad_file("No legacy section headers or canonical vehicle header row were found.")
    if len(rows) > MAX_ROWS:
        raise VehicleXlsxError(
            "VEHICLE_IMPORT_TOO_LARGE", f"Workbook contains more than {MAX_ROWS} vehicle rows."
        )

    row_lookup = {(row.sheet, row.row_number): row.row_id for row in rows}
    images, image_warnings = _extract_images(
        data,
        xml_parts,
        row_lookup=row_lookup,
        legacy_sheets=legacy_sheets,
    )
    warnings.extend(image_warnings)
    rows_by_id = {row.row_id: row for row in rows}
    for image in images:
        if image.row_id is not None and image.row_id in rows_by_id:
            rows_by_id[image.row_id].image_ids.append(image.image_id)
    return ParsedVehicleWorkbook(
        layout=layout,
        sections=sections,
        rows=rows,
        images=images,
        warnings=warnings,
    )


def build_vehicle_import_template() -> bytes:
    """Return the blank canonical import workbook with bilingual guidance."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vehicles"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    for column_number, column_name in enumerate(CANONICAL_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_number, value=column_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[cell.column_letter].width = max(14, min(24, len(column_name) + 2))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(CANONICAL_COLUMNS)).coordinate}"
    sheet.sheet_view.rightToLeft = False
    sheet.row_dimensions[1].height = 30

    instructions = workbook.create_sheet("Instructions - التعليمات")
    instructions.sheet_view.rightToLeft = True
    overview = (
        ("Vehicle import template", "نموذج استيراد المركبات"),
        (
            "Enter one vehicle per row on the Vehicles sheet.",
            "أدخل مركبة واحدة في كل صف في ورقة Vehicles.",
        ),
        ("Do not rename canonical column headers.", "لا تغيّر أسماء أعمدة النموذج."),
        (
            "Create: fill every column marked Yes / نعم. An assigned site is also required during review.",
            "الإنشاء: عبّئ كل عمود مميّز بـ Yes / نعم. يجب أيضًا تعيين الموقع أثناء المراجعة.",
        ),
        (
            "Update: a blank cell preserves the existing stored value; enter a value only when it should change.",
            "التحديث: تترك الخلية الفارغة القيمة الحالية المخزنة كما هي؛ أدخل قيمة فقط عند الحاجة إلى تغييرها.",
        ),
        ("Dates: YYYY-MM-DD or DD/MM/YYYY.", "التواريخ: YYYY-MM-DD أو DD/MM/YYYY."),
        (
            "Plate numbers are text; keep significant leading zeros.",
            "أرقام اللوحات نصية؛ احتفظ بالأصفار البادئة المهمة.",
        ),
        (
            "Known legacy class spellings become exact bilingual presets; unknown class text is preserved for custom review.",
            "تُحوّل صيغ الفئات القديمة المعروفة إلى الفئات الثنائية المعتمدة؛ ويُحتفظ بنص الفئة غير المعروفة للمراجعة كفئة مخصصة.",
        ),
        (
            "Place images on the same row as their vehicle.",
            "ضع الصور في صف المركبة نفسه.",
        ),
        (
            "Choose photo or licence for every image during review.",
            "اختر صورة مركبة أو رخصة لكل صورة أثناء المراجعة.",
        ),
        (
            "Site assignment is confirmed during import review.",
            "يتم تأكيد تعيين الموقع أثناء مراجعة الاستيراد.",
        ),
    )
    for row_number, (english, arabic) in enumerate(overview, start=1):
        instructions.cell(row=row_number, column=2, value=english)
        instructions.cell(row=row_number, column=3, value=arabic)

    guidance_header_row = len(overview) + 2
    guidance_headers = (
        "Canonical column / اسم العمود",
        "English guidance",
        "الإرشادات العربية",
        "Create required? / مطلوب عند الإنشاء؟",
    )
    for column_number, value in enumerate(guidance_headers, start=1):
        cell = instructions.cell(row=guidance_header_row, column=column_number, value=value)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_number, column_name in enumerate(
        CANONICAL_COLUMNS,
        start=guidance_header_row + 1,
    ):
        english, arabic = _CANONICAL_COLUMN_GUIDANCE[column_name]
        instructions.cell(row=row_number, column=1, value=column_name)
        instructions.cell(row=row_number, column=2, value=english)
        instructions.cell(row=row_number, column=3, value=arabic)
        instructions.cell(
            row=row_number,
            column=4,
            value="Yes / نعم" if column_name in _CREATE_REQUIRED_COLUMNS else "No / لا",
        )

    for row in instructions.iter_rows():
        for cell in row:
            if cell.row != guidance_header_row:
                cell.font = Font(name="Arial", bold=cell.row == 1)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["B"].width = 58
    instructions.column_dimensions["C"].width = 58
    instructions.column_dimensions["D"].width = 28
    instructions.freeze_panes = f"A{guidance_header_row + 1}"
    instructions.auto_filter.ref = (
        f"A{guidance_header_row}:D{guidance_header_row + len(CANONICAL_COLUMNS)}"
    )

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


__all__ = [
    "CANONICAL_COLUMNS",
    "MAX_COMPRESSED_BYTES",
    "MAX_EXPANDED_BYTES",
    "MAX_IMAGES",
    "MAX_ROWS",
    "MAX_ZIP_MEMBERS",
    "ParsedVehicleWorkbook",
    "VehicleXlsxError",
    "VehicleXlsxImage",
    "VehicleXlsxIssue",
    "VehicleXlsxRow",
    "VehicleXlsxSection",
    "build_vehicle_import_template",
    "normalize_import_values",
    "parse_vehicle_workbook",
]
