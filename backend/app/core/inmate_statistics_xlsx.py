"""Render the monthly inmate conduct violations register as workbook bytes."""

from __future__ import annotations

import io
from calendar import month_name
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from typing import Any, Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.constants import ARABIC_MONTHS


@dataclass(frozen=True)
class RegisterRow:
    row_no: int
    name: str
    uid: str
    nationality_label: str
    violation_date: date
    duty_unit: str
    details_text: str


@dataclass(frozen=True)
class RegisterSection:
    key: str
    rows: tuple[RegisterRow, ...]


_SECTION_KEYS: Final[tuple[str, ...]] = ("citizens", "expats", "pending")
_SECTION_TITLES: Final[dict[str, str]] = {
    "citizens": "المواطنون",
    "expats": "الوافدون",
    "pending": "قيد الإكمال",
}
_HEADERS: Final[dict[str, tuple[str, ...]]] = {
    "ar": (
        "ت",
        "الإسم",
        "الرقم الموحد",
        "تاريخ المخالفة",
        "السربة",
        "تفاصيل المخالفة",
    ),
    "en": (
        "No.",
        "Name",
        "Unified No.",
        "Violation date",
        "Duty company",
        "Violation details",
    ),
}
_NATIONALITY_HEADERS: Final[dict[str, str]] = {"ar": "الجنسية", "en": "Nationality"}
_SUMMARY_LABELS: Final[dict[str, dict[str, str]]] = {
    "ar": {
        "title": "ملخص سجل مخالفات المحتجزين",
        "citizens": "المواطنون",
        "expats": "الوافدون",
        "pending": "قيد الإكمال",
        "total": "إجمالي مخالفات الشهر",
        "status": "الحالة",
        "closed": "مغلق",
        "legacy": "تقرير سابق محفوظ",
        "open": "مفتوح — إسقاط حي",
        "closed_at": "تاريخ الإغلاق",
        "closed_by": "أغلق بواسطة",
        "forced": "إغلاق إداري مع سجلات ناقصة",
        "reason": "السبب",
        "count": "عدد المخالفات",
    },
    "en": {
        "title": "Inmate Violations Register Summary",
        "citizens": "Citizens",
        "expats": "Non-citizens",
        "pending": "Pending completion",
        "total": "Month total",
        "status": "Status",
        "closed": "Closed",
        "legacy": "Legacy report",
        "open": "Open — live projection",
        "closed_at": "Closed at",
        "closed_by": "Closed by",
        "forced": "Administrative close with incomplete records",
        "reason": "Reason",
        "count": "Violation count",
    },
}
_HEADER_FILL: Final = PatternFill(fill_type="solid", fgColor="1F4E78")
_SUBTOTAL_FILL: Final = PatternFill(fill_type="solid", fgColor="D9EAF7")
_HEADER_FONT: Final = Font(name="Arial", size=11, bold=True, color="FFFFFF")
_BODY_FONT: Final = Font(name="Arial", size=11)
_BOLD_FONT: Final = Font(name="Arial", size=11, bold=True)
_THIN_SIDE: Final = Side(style="thin", color="7F7F7F")
_THIN_BORDER: Final = Border(
    left=_THIN_SIDE,
    right=_THIN_SIDE,
    top=_THIN_SIDE,
    bottom=_THIN_SIDE,
)
_CENTER: Final = Alignment(horizontal="center", vertical="center")
_RIGHT: Final = Alignment(horizontal="right", vertical="center")
_WRAPPED: Final = Alignment(horizontal="right", vertical="top", wrap_text=True)


def _require_language(language: str) -> None:
    if language not in _HEADERS:
        raise ValueError(f"unsupported register language: {language!r}")


def _section_map(sections: Sequence[RegisterSection]) -> dict[str, RegisterSection]:
    result = {key: RegisterSection(key=key, rows=()) for key in _SECTION_KEYS}
    seen: set[str] = set()
    for section in sections:
        if section.key not in result:
            raise ValueError(f"unsupported register section: {section.key!r}")
        if section.key in seen:
            raise ValueError(f"duplicate register section: {section.key!r}")
        result[section.key] = section
        seen.add(section.key)
    return result


def _headers(section_key: str, language: str) -> tuple[str, ...]:
    headers = list(_HEADERS[language])
    if section_key == "expats":
        headers.insert(3, _NATIONALITY_HEADERS[language])
    return tuple(headers)


def _style_header(sheet: Any, columns: int) -> None:
    sheet.row_dimensions[1].height = 25
    for column in range(1, columns + 1):
        cell = sheet.cell(1, column)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _THIN_BORDER
        cell.alignment = _CENTER


def _style_body_cell(cell: Any, *, wrapped: bool = False) -> None:
    cell.font = _BODY_FONT
    cell.border = _THIN_BORDER
    cell.alignment = _WRAPPED if wrapped else _RIGHT


def _write_section_sheet(workbook: Any, section: RegisterSection, language: str) -> None:
    sheet = workbook.create_sheet(_SECTION_TITLES[section.key])
    headers = _headers(section.key, language)
    sheet.append(headers)
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    _style_header(sheet, len(headers))

    nationality_column = 4 if section.key == "expats" else None
    date_column = 5 if section.key == "expats" else 4
    details_column = len(headers)

    for entry in section.rows:
        values: list[object] = [entry.row_no, entry.name, entry.uid]
        if nationality_column is not None:
            values.append(entry.nationality_label)
        values.extend((entry.violation_date, entry.duty_unit, entry.details_text))
        sheet.append(values)
        row = sheet.max_row
        sheet.row_dimensions[row].height = 34
        for column in range(1, len(headers) + 1):
            _style_body_cell(sheet.cell(row, column), wrapped=column == details_column)
        sheet.cell(row, 1).number_format = "0"
        sheet.cell(row, 1).alignment = _CENTER
        sheet.cell(row, 3).number_format = "@"
        sheet.cell(row, 3).alignment = _CENTER
        sheet.cell(row, date_column).number_format = "yyyy-mm-dd"
        sheet.cell(row, date_column).alignment = _CENTER
        if nationality_column is not None:
            sheet.cell(row, nationality_column).alignment = _CENTER

    count_row = sheet.max_row + 1
    count_label = _SUMMARY_LABELS[language]["count"]
    sheet.cell(count_row, 1, f"{count_label}: {len(section.rows)}")
    sheet.merge_cells(
        start_row=count_row,
        start_column=1,
        end_row=count_row,
        end_column=len(headers),
    )
    count_cell = sheet.cell(count_row, 1)
    count_cell.font = _BOLD_FONT
    count_cell.fill = _SUBTOTAL_FILL
    count_cell.border = _THIN_BORDER
    count_cell.alignment = _RIGHT

    widths = (
        (8, 28, 24, 18, 22, 52)
        if section.key != "expats"
        else (
            8,
            28,
            24,
            18,
            18,
            22,
            52,
        )
    )
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def _month_label(year: int, month: int, language: str) -> str:
    name = ARABIC_MONTHS[month - 1] if language == "ar" else month_name[month]
    return f"{name} {year}"


def _excel_datetime(value: datetime) -> datetime:
    if value.utcoffset() is None:
        return value
    return value.astimezone(UTC).replace(tzinfo=None)


def _write_summary_sheet(
    workbook: Any,
    *,
    year: int,
    month: int,
    sections: dict[str, RegisterSection],
    closed_at: datetime | None,
    closed_by_name: str | None,
    legacy_force_reason: str | None,
    language: str,
    legacy: bool,
) -> None:
    labels = _SUMMARY_LABELS[language]
    sheet = workbook.create_sheet("الملخص")
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = f"{labels['title']} — {_month_label(year, month, language)}"
    _style_header(sheet, 2)

    counts = {key: len(sections[key].rows) for key in _SECTION_KEYS}
    # The pending group's line prints only when the group is non-empty, exactly
    # as its worksheet does; the two populations always print.
    rows: list[tuple[str, object]] = [
        (labels[key], counts[key]) for key in _SECTION_KEYS if key != "pending" or counts[key] > 0
    ]
    rows.extend(
        (
            (labels["total"], sum(counts[key] for key in _SECTION_KEYS)),
            (labels["status"], labels["legacy"] if legacy else labels["closed"] if closed_at is not None else labels["open"]),
        )
    )
    if closed_at is not None:
        rows.extend(
            (
                (labels["closed_at"], _excel_datetime(closed_at)),
                (labels["closed_by"], closed_by_name or ""),
            )
        )
        if legacy_force_reason is not None:
            rows.extend(
                (
                    (labels["forced"], ""),
                    (labels["reason"], legacy_force_reason),
                )
            )

    for label, value in rows:
        sheet.append((label, value))
        row = sheet.max_row
        for column in (1, 2):
            _style_body_cell(sheet.cell(row, column), wrapped=True)
        sheet.cell(row, 1).font = _BOLD_FONT
        if isinstance(value, int):
            sheet.cell(row, 2).number_format = "0"
        elif isinstance(value, datetime):
            sheet.cell(row, 2).number_format = "yyyy-mm-dd hh:mm"

    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 36


def build_register_workbook(
    *,
    year: int,
    month: int,
    sections: Sequence[RegisterSection],
    closed_at: datetime | None,
    closed_by_name: str | None,
    legacy_force_reason: str | None,
    language: str,
    legacy: bool = False,
) -> bytes:
    """Render one monthly register without reading a database or filesystem."""

    _require_language(language)
    section_map = _section_map(sections)
    workbook = Workbook()
    workbook.remove(workbook.active)

    for key in _SECTION_KEYS:
        section = section_map[key]
        if key != "pending" or section.rows:
            _write_section_sheet(workbook, section, language)
    _write_summary_sheet(
        workbook,
        year=year,
        month=month,
        sections=section_map,
        closed_at=closed_at,
        closed_by_name=closed_by_name,
        legacy_force_reason=legacy_force_reason,
        language=language,
        legacy=legacy,
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def register_filename(year: int, month: int, language: str) -> str:
    """Return the download name for a monthly register."""

    _require_language(language)
    if language == "ar":
        return f"سجل مخالفات المحتجزين {ARABIC_MONTHS[month - 1]} {year}.xlsx"
    return f"Inmate Violations Register {month_name[month]} {year}.xlsx"
