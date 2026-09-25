"""Behavior tests for signature_placement_service — the public contract of
approval-signature-placement plan §§5, 7. The Word-COM measurement/movement
primitives themselves are proven separately and empirically against real
Word in ``test_artifact_word_windows.py::test_signature_placement_word_smoke``
(the isolated, opt-in Word gate); this file exercises everything ABOVE that
boundary — identity/history bookkeeping, authorization, and the
compare-and-set publish transaction — with the Word layer mocked, so it runs
in every ordinary test invocation without Microsoft Word installed.
"""

from __future__ import annotations

from pathlib import Path

import fitz
import pytest
from sqlalchemy.orm import Session

from app.api.errors import AppError, ConflictError, NotFoundError
from app.config import Settings
from app.core.signature_layout import SignatureGeometry, SignatureLayout, SignaturePageInfo
from app.db.models import (
    Book,
    BookCategory,
    BookVersion,
    Document,
    SignatureArtifactRevision,
    User,
)
from app.services import _pdf_executor
from app.services import signature_placement_service as sps


def _pdf(path: Path, labels: list[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = fitz.open()
    try:
        for label in labels:
            page = doc.new_page(width=300, height=400)
            page.insert_text((40, 80), label)
        doc.save(path)
    finally:
        doc.close()
    return path


def _record(
    db: Session, tmp_path: Path, *, manager_sig_embedded: bool = False
) -> tuple[Book, BookVersion, Document]:
    if db.get(BookCategory, "HR") is None:
        db.add(BookCategory(id="HR", prefix="HR"))
    db.flush()
    generated = _pdf(tmp_path / "generated.pdf", ["FORM"])
    docx_path = tmp_path / "source.docx"
    docx_path.write_bytes(b"docx-placeholder")
    document = Document(
        template_id="General Book",
        ref_number="HR-1",
        docx_path=docx_path.relative_to(tmp_path).as_posix(),
        pdf_path=generated.relative_to(tmp_path).as_posix(),
        submission_id="submission-1",
        role="primary",
    )
    db.add(document)
    db.flush()
    book = Book(category_id="HR", ref_number="HR-1", approval_state="approved")
    db.add(book)
    db.flush()
    version = BookVersion(
        book_id=book.id,
        version_no=1,
        document_id=document.id,
        template_id="General Book",
        fields={},
        status="approved",
        manager_sig_embedded=manager_sig_embedded,
    )
    db.add(version)
    db.commit()
    return book, version, document


def _fake_layout(*, page: int, x: float, y: float, sha: str, signature_id: str) -> SignatureLayout:
    return SignatureLayout(
        source_sha256=sha,
        pages=(
            SignaturePageInfo(page=page, width_pt=612.0, height_pt=792.0, anchor_paragraph_index=0),
        ),
        drawings={
            signature_id: SignatureGeometry(page=page, x=x, y=y, width_pt=90.0, height_pt=40.0)
        },
    )


# ---------------------------------------------------------------------------
# can_correct — authorization
# ---------------------------------------------------------------------------


def test_can_correct_original_signer_succeeds() -> None:
    version = BookVersion(signature_revision=3)
    active = SignatureArtifactRevision(revision=3, signer_user_id=7)
    assert sps.can_correct(
        version=version, active_revision=active, actor_user_id=7, actor_is_admin=False
    )


def test_can_correct_different_approver_fails() -> None:
    version = BookVersion(signature_revision=3)
    active = SignatureArtifactRevision(revision=3, signer_user_id=7)
    assert not sps.can_correct(
        version=version, active_revision=active, actor_user_id=8, actor_is_admin=False
    )


def test_can_correct_admin_always_succeeds_even_without_active_revision() -> None:
    version = BookVersion(signature_revision=0)
    assert sps.can_correct(
        version=version, active_revision=None, actor_user_id=1, actor_is_admin=True
    )


def test_can_correct_auto_embedded_no_signer_is_admin_only() -> None:
    version = BookVersion(signature_revision=1)
    active = SignatureArtifactRevision(revision=1, signer_user_id=None)
    assert not sps.can_correct(
        version=version, active_revision=active, actor_user_id=1, actor_is_admin=False
    )
    assert sps.can_correct(
        version=version, active_revision=active, actor_user_id=1, actor_is_admin=True
    )


def test_can_correct_stale_revision_fails_even_for_original_signer() -> None:
    version = BookVersion(signature_revision=4)  # a later correction already landed
    active = SignatureArtifactRevision(revision=3, signer_user_id=7)
    assert not sps.can_correct(
        version=version, active_revision=active, actor_user_id=7, actor_is_admin=False
    )


# ---------------------------------------------------------------------------
# capture_initial_revision / invalidate_revision
# ---------------------------------------------------------------------------


def test_capture_initial_revision_noop_without_manager_drawing(
    db_session: Session, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    settings = Settings(data_dir=tmp_path, templates_dir=tmp_path)
    monkeypatch.setattr(sps, "get_settings", lambda: settings)
    version = BookVersion(id=1, book_id=1, version_no=1, signature_revision=0)
    unmarked = tmp_path / "unmarked.docx"
    from docx import Document as DocxFile

    DocxFile().save(str(unmarked))

    row = sps.capture_initial_revision(
        db_session,
        version=version,
        source_kind=sps.SOURCE_KIND_APPROVAL,
        signer_user_id=5,
        docx_path=unmarked,
        primary_pdf_path=None,
        published_pdf_path=None,
    )
    assert row is None
    assert version.signature_revision == 0


def test_capture_initial_revision_captures_and_sets_revision(
    db_session: Session, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    settings = Settings(data_dir=tmp_path, templates_dir=tmp_path)
    monkeypatch.setattr(sps, "get_settings", lambda: settings)
    version = BookVersion(id=1, book_id=1, version_no=1, signature_revision=0)

    from docx import Document as DocxFile

    from app.core.docx_engine import stamp_signature_above_name

    letter = tmp_path / "letter.docx"
    doc = DocxFile()
    doc.add_paragraph("Body")
    doc.add_paragraph("Ali Manager")
    doc.save(str(letter))
    sig = tmp_path / "sig.png"
    from PIL import Image

    Image.new("RGBA", (60, 30), (0, 0, 0, 255)).save(sig)
    assert stamp_signature_above_name(letter, str(sig), ["Ali Manager"], size_mm=32.0, boldness=1)

    row = sps.capture_initial_revision(
        db_session,
        version=version,
        source_kind=sps.SOURCE_KIND_APPROVAL,
        signer_user_id=5,
        docx_path=letter,
        primary_pdf_path=None,
        published_pdf_path=None,
    )
    assert row is not None
    assert row.revision == 1
    assert row.action == sps.ACTION_INITIAL
    assert row.signer_user_id == 5
    assert version.signature_revision == 1
    retained = settings.data_dir / row.docx_path
    assert retained.is_file()
    assert retained != letter  # a COPY, source untouched
    assert letter.is_file()

    # Idempotent: a second call after signature_revision advanced is a no-op.
    again = sps.capture_initial_revision(
        db_session,
        version=version,
        source_kind=sps.SOURCE_KIND_APPROVAL,
        signer_user_id=5,
        docx_path=letter,
        primary_pdf_path=None,
        published_pdf_path=None,
    )
    assert again is None


def test_invalidate_revision_resets_to_zero_keeps_history_row(
    db_session: Session,
) -> None:
    version = BookVersion(id=1, book_id=1, version_no=1, signature_revision=2)
    sps.invalidate_revision(db_session, version=version)
    assert version.signature_revision == 0


# ---------------------------------------------------------------------------
# move_signature — authorization, CAS, and the publish transaction
# ---------------------------------------------------------------------------


def _active_row(
    db: Session, version: BookVersion, tmp_path: Path, *, signer_user_id: int | None
) -> SignatureArtifactRevision:
    tracked_docx = tmp_path / "tracked.docx"
    tracked_docx.write_bytes(b"tracked-docx-bytes")
    tracked_pdf = _pdf(tmp_path / "tracked.pdf", ["TRACKED"])
    row = SignatureArtifactRevision(
        version_id=version.id,
        revision=1,
        source_kind=sps.SOURCE_KIND_APPROVAL,
        signer_user_id=signer_user_id,
        docx_path=tracked_docx.relative_to(tmp_path).as_posix(),
        primary_pdf_path=tracked_pdf.relative_to(tmp_path).as_posix(),
        published_pdf_path=None,
        previous_published_pdf_path=None,
        docx_sha256=sps._sha256_file(tracked_docx),
        manifest=[],
        action=sps.ACTION_INITIAL,
        signature_id=None,
        before_geometry=None,
        after_geometry=None,
        actor_user_id=signer_user_id,
    )
    db.add(row)
    version.signature_revision = 1
    db.commit()
    return row


def test_move_signature_wrong_actor_is_forbidden(
    db_session: Session, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    settings = Settings(data_dir=tmp_path, templates_dir=tmp_path)
    monkeypatch.setattr(sps, "get_settings", lambda: settings)
    _book, version, document = _record(db_session, tmp_path)
    active = _active_row(db_session, version, tmp_path, signer_user_id=7)
    signer = User(id=7, email="signer@x.ae", password_hash="x", role="manager", status="active")
    other = User(id=8, email="other@x.ae", password_hash="x", role="operator", status="active")
    db_session.add_all([signer, other])
    db_session.commit()

    with pytest.raises(AppError) as error:
        sps.move_signature(
            db_session,
            document.id,
            user=other,
            signature_id="whatever",
            signature_revision=1,
            package_revision=0,
            source_sha256=active.docx_sha256,
            page=1,
            x=0.1,
            y=0.1,
        )
    assert error.value.http_status == 403
    db_session.refresh(version)
    assert version.signature_revision == 1  # untouched


def test_move_signature_admin_without_active_revision_is_not_found(
    db_session: Session, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """`can_correct` authorizes an administrator unconditionally, even with
    no active revision — that predicate answers "would this actor be
    allowed", not "does a target exist". `move_signature` must still refuse
    cleanly (404, not an internal assertion crash) when there is nothing to
    correct yet."""
    settings = Settings(data_dir=tmp_path, templates_dir=tmp_path)
    monkeypatch.setattr(sps, "get_settings", lambda: settings)
    _book, _version, document = _record(db_session, tmp_path)
    admin = User(id=1, email="admin@x.ae", password_hash="x", role="admin", status="active")
    db_session.add(admin)
    db_session.commit()

    with pytest.raises(NotFoundError) as error:
        sps.move_signature(
            db_session,
            document.id,
            user=admin,
            signature_id="whatever",
            signature_revision=0,
            package_revision=0,
            source_sha256="0" * 64,
            page=1,
            x=0.1,
            y=0.1,
        )
    assert error.value.code == "SIGNATURE_SOURCE_UNAVAILABLE"


def test_move_signature_stale_revision_conflicts_before_any_word_work(
    db_session: Session, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    settings = Settings(data_dir=tmp_path, templates_dir=tmp_path)
    monkeypatch.setattr(sps, "get_settings", lambda: settings)
    _book, version, document = _record(db_session, tmp_path)
    active = _active_row(db_session, version, tmp_path, signer_user_id=7)
    signer = User(id=7, email="signer2@x.ae", password_hash="x", role="manager", status="active")
    db_session.add(signer)
    db_session.commit()

    called = False

    def _should_not_run(*_args: object, **_kwargs: object) -> None:
        nonlocal called
        called = True
        raise AssertionError("measure_signature_position must not run on a stale revision")

    monkeypatch.setattr(_pdf_executor, "measure_signature_position", _should_not_run)

    with pytest.raises(ConflictError):
        sps.move_signature(
            db_session,
            document.id,
            user=signer,
            signature_id="whatever",
            signature_revision=99,  # stale
            package_revision=0,
            source_sha256=active.docx_sha256,
            page=1,
            x=0.1,
            y=0.1,
        )
    assert called is False


def test_move_signature_success_appends_revision_and_publishes(
    db_session: Session, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    settings = Settings(data_dir=tmp_path, templates_dir=tmp_path)
    monkeypatch.setattr(sps, "get_settings", lambda: settings)
    monkeypatch.setattr("app.services.included_papers_service.get_settings", lambda: settings)
    book, version, document = _record(db_session, tmp_path)
    active = _active_row(db_session, version, tmp_path, signer_user_id=7)
    signer = User(id=7, email="signer3@x.ae", password_hash="x", role="manager", status="active")
    db_session.add(signer)
    db_session.commit()

    signature_id = "sig-under-test"
    before_layout = _fake_layout(
        page=1, x=0.2, y=0.2, sha=active.docx_sha256, signature_id=signature_id
    )
    moved_pdf = _pdf(tmp_path / "moved-output.pdf", ["MOVED-FORM"])

    def _fake_measure(docx_path: Path) -> SignatureLayout:
        return before_layout

    def _fake_move(
        source: Path,
        destination: Path,
        *,
        signature_id: str,
        page: int,
        x: float,
        y: float,
        layout: SignatureLayout,
        before_pdf: Path,
    ) -> tuple[SignatureLayout, Path]:
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(b"moved-docx-bytes")
        after_layout = _fake_layout(page=page, x=x, y=y, sha="after-sha", signature_id=signature_id)
        return after_layout, moved_pdf

    monkeypatch.setattr(_pdf_executor, "measure_signature_position", _fake_measure)
    monkeypatch.setattr(_pdf_executor, "move_signature_position", _fake_move)

    description = sps.move_signature(
        db_session,
        document.id,
        user=signer,
        signature_id=signature_id,
        signature_revision=1,
        package_revision=0,
        source_sha256=active.docx_sha256,
        page=1,
        x=0.35,
        y=0.4,
    )

    db_session.refresh(version)
    db_session.refresh(book)
    assert version.signature_revision == 2
    assert book.included_papers_revision == 1
    assert description.signature_revision == 2
    assert description.package_revision == 1

    rows = (
        db_session.query(SignatureArtifactRevision)
        .filter(SignatureArtifactRevision.version_id == version.id)
        .order_by(SignatureArtifactRevision.revision)
        .all()
    )
    assert [r.revision for r in rows] == [1, 2]
    new_row = rows[1]
    assert new_row.action == sps.ACTION_MOVE
    assert new_row.signature_id == signature_id
    assert new_row.signer_user_id == 7  # carried from the original signer, not the correcting admin
    assert new_row.before_geometry == {"page": 1, "x": 0.2, "y": 0.2}
    assert new_row.after_geometry == {"page": 1, "x": 0.35, "y": 0.4}
    assert (settings.data_dir / new_row.docx_path).is_file()
    assert new_row.published_pdf_path is not None

    from app.db.models import AuditLog

    audit = db_session.query(AuditLog).filter(AuditLog.action == "signature.position_changed").one()
    assert audit.entity_id == str(document.id)


def test_move_signature_layout_unsupported_page_rejected(
    db_session: Session, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    settings = Settings(data_dir=tmp_path, templates_dir=tmp_path)
    monkeypatch.setattr(sps, "get_settings", lambda: settings)
    _book, version, document = _record(db_session, tmp_path)
    active = _active_row(db_session, version, tmp_path, signer_user_id=7)
    signer = User(id=7, email="signer4@x.ae", password_hash="x", role="manager", status="active")
    db_session.add(signer)
    db_session.commit()

    signature_id = "sig-unsupported"
    layout = _fake_layout(page=1, x=0.2, y=0.2, sha=active.docx_sha256, signature_id=signature_id)

    def _raise_unsupported(*_args: object, **_kwargs: object) -> None:
        from app.core.signature_layout import SignatureLayoutUnsupportedError

        raise SignatureLayoutUnsupportedError("no safe anchor on page 9")

    monkeypatch.setattr(_pdf_executor, "measure_signature_position", lambda _p: layout)
    monkeypatch.setattr(_pdf_executor, "move_signature_position", _raise_unsupported)

    from app.core.signature_layout import SignatureLayoutUnsupportedError

    with pytest.raises(SignatureLayoutUnsupportedError):
        sps.move_signature(
            db_session,
            document.id,
            user=signer,
            signature_id=signature_id,
            signature_revision=1,
            package_revision=0,
            source_sha256=active.docx_sha256,
            page=9,
            x=0.1,
            y=0.1,
        )
    db_session.refresh(version)
    assert version.signature_revision == 1  # untouched — failure cleaned up, nothing committed
    retained_root = settings.data_dir / "signature_artifacts" / str(version.id)
    # No revision-2 directory left behind.
    assert not any(
        p.name.startswith("2-") for p in (retained_root.iterdir() if retained_root.is_dir() else [])
    )


# ---------------------------------------------------------------------------
# reassign_signature — admin-only employee picker (SignaturePlacementPage)
# ---------------------------------------------------------------------------


def test_reassign_signature_non_admin_is_forbidden(
    db_session: Session, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    settings = Settings(data_dir=tmp_path, templates_dir=tmp_path)
    monkeypatch.setattr(sps, "get_settings", lambda: settings)
    book, version, document = _record(db_session, tmp_path)
    _active_row(db_session, version, tmp_path, signer_user_id=7)
    non_admin = User(id=7, email="signer5@x.ae", password_hash="x", role="manager", status="active")
    db_session.add(non_admin)
    db_session.commit()

    with pytest.raises(AppError):
        sps.reassign_signature(
            db_session,
            document.id,
            user=non_admin,
            signature_id="sig-x",
            signature_revision=1,
            package_revision=0,
            source_sha256="whatever",
            employee_id="G1042",
        )
    db_session.refresh(version)
    assert version.signature_revision == 1  # untouched


def test_reassign_signature_success_swaps_image_and_appends_revision(
    db_session: Session, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    from app.db.models import Employee

    settings = Settings(data_dir=tmp_path, templates_dir=tmp_path)
    monkeypatch.setattr(sps, "get_settings", lambda: settings)
    monkeypatch.setattr("app.services.included_papers_service.get_settings", lambda: settings)
    book, version, document = _record(db_session, tmp_path)
    active = _active_row(db_session, version, tmp_path, signer_user_id=7)
    admin = User(id=9, email="admin1@x.ae", password_hash="x", role="admin", status="active")
    other_employee = Employee(id="G1042", name_en="Muhannad", name_ar="مهند", position="Head")
    db_session.add_all([admin, other_employee])
    db_session.commit()

    sig_file = tmp_path / "vault" / "G1042.png"
    sig_file.parent.mkdir(parents=True, exist_ok=True)
    sig_file.write_bytes(b"other-employee-signature-bytes")
    monkeypatch.setattr(
        "app.core.signature.employee_signature_str", lambda _vault, _emp_id: str(sig_file)
    )

    signature_id = "sig-under-test"
    layout = _fake_layout(page=1, x=0.2, y=0.2, sha=active.docx_sha256, signature_id=signature_id)
    reassigned_pdf = _pdf(tmp_path / "reassigned-output.pdf", ["REASSIGNED-FORM"])

    def _fake_reassign(
        source: Path,
        destination: Path,
        *,
        signature_id: str,
        image_bytes: bytes,
        layout: SignatureLayout,
        before_pdf: Path,
    ) -> tuple[SignatureLayout, Path]:
        assert image_bytes == b"other-employee-signature-bytes"
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(b"reassigned-docx-bytes")
        return layout, reassigned_pdf

    monkeypatch.setattr(_pdf_executor, "measure_signature_position", lambda _p: layout)
    monkeypatch.setattr(_pdf_executor, "reassign_signature_image", _fake_reassign)

    description = sps.reassign_signature(
        db_session,
        document.id,
        user=admin,
        signature_id=signature_id,
        signature_revision=1,
        package_revision=0,
        source_sha256=active.docx_sha256,
        employee_id="G1042",
    )

    db_session.refresh(version)
    db_session.refresh(book)
    assert version.signature_revision == 2
    assert book.included_papers_revision == 1
    assert description.signature_revision == 2

    rows = (
        db_session.query(SignatureArtifactRevision)
        .filter(SignatureArtifactRevision.version_id == version.id)
        .order_by(SignatureArtifactRevision.revision)
        .all()
    )
    new_row = rows[-1]
    assert new_row.action == sps.ACTION_REASSIGN
    assert new_row.signer_user_id is None  # G1042 has no linked User row in this test
    assert (settings.data_dir / new_row.docx_path).is_file()

    from app.db.models import AuditLog

    audit = db_session.query(AuditLog).filter(AuditLog.action == "signature.reassigned").one()
    assert audit.entity_id == str(document.id)


def test_reassign_signature_unknown_employee_not_found(
    db_session: Session, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    settings = Settings(data_dir=tmp_path, templates_dir=tmp_path)
    monkeypatch.setattr(sps, "get_settings", lambda: settings)
    book, version, document = _record(db_session, tmp_path)
    _active_row(db_session, version, tmp_path, signer_user_id=7)
    admin = User(id=9, email="admin2@x.ae", password_hash="x", role="admin", status="active")
    db_session.add(admin)
    db_session.commit()

    with pytest.raises(NotFoundError):
        sps.reassign_signature(
            db_session,
            document.id,
            user=admin,
            signature_id="sig-x",
            signature_revision=1,
            package_revision=0,
            source_sha256="whatever",
            employee_id="NO-SUCH-EMPLOYEE",
        )
