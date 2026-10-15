from docx import Document

from app.core.arabic_rtl import html_to_docx


def test_break_after_avoid_keeps_only_that_paragraph_with_the_next():
    doc = Document()
    html_to_docx(
        '<p style="page-break-after:avoid">A</p><p>B</p><p style="break-after: avoid">C</p>',
        doc.add_paragraph(),
    )
    keeps = {p.text: p.paragraph_format.keep_with_next for p in doc.paragraphs}
    assert keeps == {"A": True, "B": None, "C": True}
