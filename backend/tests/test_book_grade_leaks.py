"""Personal records stay private in related-record projections and actions."""

from datetime import date, datetime
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from fastapi.testclient import TestClient

from app.api.deps import get_current_user
from app.db.models import (
    Book,
    BookCategory,
    BookVersion,
    Document,
    Employee,
    ItemPermit,
    LedgerEntry,
    Permit,
    UserPermission,
)
from app.db.session import get_db
from app.main import create_app
from app.services import (
    announce_service,
    employee_detail_service,
    item_permit_service,
    permit_service,
)
from app.services.document_reader import DocumentRead
from app.services.scan_triage_service import (
    ExternalClassification,
    ReturnedFormClassification,
    classify_text,
)
from tests.conftest import make_user


@pytest.fixture()
def records(api_db):
    creator = make_user(api_db, email="creator@test.ae")
    other = make_user(api_db, email="other@test.ae")
    admin = make_user(api_db, role="admin", email="admin@test.ae")
    employee = Employee(id="G100", name_en="Test Employee", name_ar="موظف", status="Active")
    api_db.add_all([employee, BookCategory(id="GS", prefix="GS", name_en="General Records")])
    api_db.flush()
    books, documents, logs = [], [], []
    for number, level in ((42, "personal"), (43, "general")):
        book = Book(
            category_id="GS",
            ref_number=f"GS-{number:04d}",
            subject=f"{level} record",
            confidentiality=level,
            created_by_user_id=creator.id,
            employee_id=employee.id,
            direction="outgoing",
            approval_state="approved",
            created_at=datetime(2026, 8, 10, 9),
        )
        api_db.add(book)
        api_db.flush()
        version = BookVersion(
            book_id=book.id,
            version_no=1,
            template_id="Employment Certificate",
            fields={},
            status="approved",
            created_by_user_id=creator.id,
        )
        api_db.add(version)
        document = Document(
            employee_id=employee.id,
            template_id="Employment Certificate",
            ref_number=book.ref_number,
            docx_path="output/fake.docx",
            submission_id=f"00000000-0000-0000-0000-{number:012d}",
            created_at=book.created_at,
        )
        log = LedgerEntry(
            entry_date=date(2026, 8, 10),
            direction="outgoing",
            channel="document",
            counterparty=creator.email,
            subject=book.subject,
            source_kind="document_generated",
            owner_user_id=None,
            related_book_id=book.id,
            related_employee_id=employee.id,
            created_at=book.created_at,
        )
        api_db.add_all([document, log])
        api_db.flush()
        version.document_id = document.id
        books.append(book)
        documents.append(document)
        logs.append(log)
    api_db.commit()
    actor = {"user": other}
    app = create_app()
    app.dependency_overrides[get_db] = lambda: api_db
    app.dependency_overrides[get_current_user] = lambda: actor["user"]
    return SimpleNamespace(
        db=api_db,
        creator=creator,
        other=other,
        admin=admin,
        employee=employee,
        personal=books[0],
        general=books[1],
        documents=documents,
        logs=logs,
        actor=actor,
        client=TestClient(app, raise_server_exceptions=True),
    )


@pytest.fixture(params=["other", "creator", "admin"])
def viewer(records, request):
    user = getattr(records, request.param)
    records.actor["user"] = user
    return user, request.param != "other"


def test_shared_ledger_log_and_entry_do_not_leak(records, viewer):
    _, can_open_personal = viewer
    expected = records.logs if can_open_personal else records.logs[1:]
    response = records.client.get("/api/v1/ledger/log")
    assert response.status_code == 200, response.text
    assert {row["id"] for row in response.json()["items"]} == {row.id for row in expected}
    assert response.json()["total"] == len(expected)
    for log, visible in zip(records.logs, [can_open_personal, True], strict=True):
        for path in (f"/api/v1/ledger/log/{log.id}", f"/api/v1/ledger/{log.id}"):
            response = records.client.get(path)
            assert response.status_code == (200 if visible else 404), response.text
            if visible:
                assert response.json()["related_book_id"] == log.related_book_id
            else:
                expected_code = (
                    "LOG_RECORD_NOT_FOUND" if "/log/" in path else "LEDGER_ENTRY_NOT_FOUND"
                )
                assert response.json()["error"]["code"] == expected_code


def test_employee_activity_hides_document_and_ledger(records, viewer):
    _, can_open_personal = viewer
    response = records.client.get(
        "/api/v1/employees/activity", params={"employee_id": records.employee.id}
    )
    assert response.status_code == 200, response.text
    body = response.json()
    documents = records.documents if can_open_personal else records.documents[1:]
    logs = records.logs if can_open_personal else records.logs[1:]
    assert {(row["kind"], row["source_id"]) for row in body["items"]} == {
        *(("document", row.id) for row in documents),
        *(("ledger", row.id) for row in logs),
    }
    assert {row["reference"] for row in body["items"] if row["kind"] == "document"} == {
        row.ref_number for row in documents
    }
    assert body["total"] == len(documents) + len(logs)


def test_employee_detail_hides_documents_ledger_and_stats(records, viewer):
    user, can_open_personal = viewer
    detail = employee_detail_service.get_employee_detail(records.db, records.employee.id, user=user)
    assert detail is not None
    documents = records.documents if can_open_personal else records.documents[1:]
    logs = records.logs if can_open_personal else records.logs[1:]
    assert {row.id for row in detail.recent_documents} == {row.id for row in documents}
    assert {row.ref_number for row in detail.recent_documents} == {
        row.ref_number for row in documents
    }
    assert {row.id for row in detail.recent_ledger} == {row.id for row in logs}
    assert detail.stats.documents == len(documents)
    assert detail.stats.ledger_count == len(logs)
    assert len(detail.recent_activity) == len(documents) + len(logs)


def test_scan_reference_only_matches_visible_records(records, viewer):
    user, can_open_personal = viewer
    for book, visible in ((records.personal, can_open_personal), (records.general, True)):
        read = DocumentRead(text=f"Ref: {book.ref_number}", text_source="ocr")
        result = classify_text(read, db=records.db, employees=[], user=user)
        if visible:
            assert isinstance(result, ReturnedFormClassification)
            assert result.match.book.book_id == book.id
        else:
            assert isinstance(result, ExternalClassification)
            assert all(not candidate.live_matches for candidate in result.reference_candidates)
        system = classify_text(read, db=records.db, employees=[], user=None)
        assert isinstance(system, ReturnedFormClassification)
        assert system.match.book.book_id == book.id


def test_announcement_cannot_send_hidden_record(records, viewer, monkeypatch):
    user, can_open_personal = viewer
    records.db.add(UserPermission(user_id=user.id, capability="messages.broadcast", effect="grant"))
    records.db.commit()
    groups = Mock(return_value=[SimpleNamespace(id="test@g.us", name="Test Group")])
    pdf = Mock(return_value=("test.pdf", b"fake pdf"))
    send = Mock(
        return_value=announce_service.AnnouncementResult(
            announcement_id=1, sent=1, failed=0, results=[]
        )
    )
    monkeypatch.setattr(announce_service, "groups_available", groups)
    monkeypatch.setattr(announce_service, "resolve_book_pdf", pdf)
    monkeypatch.setattr(announce_service, "send_announcement", send)
    for book, visible in ((records.personal, can_open_personal), (records.general, True)):
        groups.reset_mock()
        pdf.reset_mock()
        send.reset_mock()
        response = records.client.post(
            "/api/v1/announcements/send", data={"group_ids": ["test@g.us"], "book_id": book.id}
        )
        assert response.status_code == (200 if visible else 404), response.text
        if visible:
            assert response.json()["sent"] == 1
            pdf.assert_called_once_with(records.db, book.id)
            assert send.call_args.kwargs["book_id"] == book.id
        else:
            assert response.json()["error"]["code"] == "BOOK_NOT_FOUND"
            groups.assert_not_called()
            pdf.assert_not_called()
            send.assert_not_called()


@pytest.mark.parametrize("kind", ["permit", "item_permit"])
def test_permit_read_redacts_hidden_letter(records, viewer, kind):
    user, can_open_personal = viewer
    for book, visible in ((records.personal, can_open_personal), (records.general, True)):
        if kind == "permit":
            row = Permit(
                company="Test Company",
                zones=["green"],
                book_id=book.id,
                start_date=date(2026, 8, 10),
                end_date=date(2026, 8, 10),
                validity_value=1,
                validity_unit="day",
            )
            service = permit_service
        else:
            row = ItemPermit(
                employee_id=records.employee.id,
                book_id=book.id,
                recipient="Test Recipient",
                zones=["green"],
                site="Test Site",
                items=[{"name": "Cable", "quantity": 1}],
            )
            service = item_permit_service
        records.db.add(row)
        records.db.commit()
        records.db.refresh(row)
        result = service.to_read(row, db=records.db, user=user)
        assert (result.book_id, result.book_ref, result.approval_state) == (
            (book.id, book.ref_number, book.approval_state) if visible else (None, None, None)
        )


def test_permit_api_redacts_another_users_personal_letter(records):
    row = Permit(
        company="Test Company",
        zones=["green"],
        book_id=records.personal.id,
        start_date=date(2026, 8, 10),
        end_date=date(2026, 8, 10),
        validity_value=1,
        validity_unit="day",
    )
    records.db.add(row)
    records.db.commit()
    response = records.client.get(f"/api/v1/permits/{row.id}")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["id"] == row.id
    assert (body["book_id"], body["book_ref"], body["approval_state"]) == (None, None, None)


def test_employee_detail_api_hides_another_users_personal_document(records):
    response = records.client.get(f"/api/v1/employees/{records.employee.id}/detail")
    assert response.status_code == 200, response.text
    body = response.json()
    assert {row["id"] for row in body["recent_documents"]} == {records.documents[1].id}
    assert {row["ref_number"] for row in body["recent_documents"]} == {records.general.ref_number}
    assert body["stats"]["documents"] == 1


@pytest.mark.parametrize("level", ("personal", "secret", "directed"))
def test_non_general_violation_never_queues_or_delivers_whatsapp(records, monkeypatch, level):
    from app.services import inmate_violation_whatsapp as whatsapp

    book = records.personal
    book.confidentiality = level
    version = book.versions[0]
    version.template_id = whatsapp.TEMPLATE_ID
    version.status = "approved"
    records.db.commit()
    monkeypatch.setattr(whatsapp, "get_settings", lambda: SimpleNamespace(openwa_enabled=True))
    monkeypatch.setattr(whatsapp, "get_group", lambda db: ("test-group", "Group"))
    resolve_pdf = Mock()
    send = Mock()
    monkeypatch.setattr(announce_service, "resolve_version_pdf", resolve_pdf)
    monkeypatch.setattr(announce_service, "send_announcement", send)
    whatsapp.queue_send(records.db, version)
    assert not records.db.info.get(whatsapp._PENDING)
    # Delivery must also recheck: approval may have queued it while General.
    whatsapp._deliver(records.db, version.id)
    resolve_pdf.assert_not_called()
    send.assert_not_called()


@pytest.mark.parametrize("sealed", (False, True))
def test_register_redacts_hidden_sources_without_changing_counts(
    records, monkeypatch, tmp_path, sealed
):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.db.models import InmateViolationPeriod
    from app.services import inmate_statistics_service as register

    for book in (records.personal, records.general):
        version = book.versions[0]
        version.template_id = register.TEMPLATE_ID
        version.fields = {
            "report_date": "2026-08-10",
            "inmate_names": ["Counted inmate"],
            "imported_approved": True,
            "completion": {
                "inmates": [{"name": "Counted inmate", "nationality": "India", "wing": "A"}],
                "violation_details": f"<p>{book.confidentiality} violation content</p>",
            },
        }
    records.db.commit()
    if sealed:
        entries = register.build_month(records.db, 2026, 8).entries
        monkeypatch.setattr(register, "get_settings", lambda: SimpleNamespace(data_dir=tmp_path))
        stored = tmp_path / "stored.xlsx"
        stored.write_bytes(b"unredacted stored workbook must not be served")
        period = InmateViolationPeriod(
            year=2026,
            month=8,
            closed_at=datetime(2026, 9, 1),
            export_path=stored.name,
        )
        records.db.add(period)
        records.db.flush()
        register._freeze(records.db, period, entries)
        records.db.commit()
    path = "/api/v1/inmate-violations/statistics/2026/8"
    response = records.client.get(path)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["counts"]["total"] == 2
    hidden = next(
        entry for entry in body["entries"] if entry["id"].startswith(f"{records.personal.id}:")
    )
    assert hidden["details_text"] == ""
    assert hidden["source_ref_number"] is None
    assert hidden["source_book_id"] is None
    assert hidden["completion_book_id"] is None
    visible = next(
        entry for entry in body["entries"] if entry["source_book_id"] == records.general.id
    )
    assert visible["details_text"] == "general violation content"
    assert (
        body["projection_fingerprint"]
        == register.build_month(records.db, 2026, 8).projection_fingerprint
    )
    export = records.client.get(f"{path}/export")
    assert export.status_code == 200, export.text
    workbook = load_workbook(BytesIO(export.content))
    cells = [
        cell.value for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row
    ]
    assert "personal violation content" not in cells
    assert "general violation content" in cells
    assert cells.count("Counted inmate") == 2
    records.actor["user"] = records.creator
    creator = records.client.get(path)
    assert creator.status_code == 200, creator.text
    assert creator.json()["counts"]["total"] == 2
    assert {entry["source_book_id"] for entry in creator.json()["entries"]} == {
        records.personal.id,
        records.general.id,
    }


def test_hidden_import_completion_requires_write_access(records):
    from app.api.errors import NotFoundError
    from app.services import inmate_statistics_service as register

    with pytest.raises(NotFoundError) as exc:
        register.complete_import(
            records.db,
            records.personal.id,
            actor=records.other,
            report_time="09:00",
            reporter_id=records.employee.id,
            violation_details="Not allowed",
            inmates=[{"name": "Inmate"}],
        )
    assert exc.value.code == "BOOK_NOT_FOUND"
