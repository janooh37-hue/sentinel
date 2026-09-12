"""Observable checks for the monthly inmate violations workbook."""

from __future__ import annotations

from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.core.inmate_statistics_xlsx import (
    RegisterRow,
    RegisterSection,
    build_register_workbook,
    register_filename,
)


def _row(
    row_no: int,
    *,
    name: str,
    uid: str,
    nationality: str,
    violation_date: date,
) -> RegisterRow:
    return RegisterRow(
        row_no=row_no,
        name=name,
        uid=uid,
        nationality_label=nationality,
        violation_date=violation_date,
        duty_unit="السرية الأولى",
        details_text="مخالفة تعليمات العنبر",
    )


def _sections(*, pending: bool = True) -> tuple[RegisterSection, ...]:
    citizen = _row(
        1,
        name="مواطن تجريبي",
        uid="784-1993-4471820-3",
        nationality="الإمارات العربية المتحدة",
        violation_date=date(2026, 7, 4),
    )
    expat = _row(
        1,
        name="وافد تجريبي",
        uid="123456789",
        nationality="الهند",
        violation_date=date(2026, 7, 9),
    )
    pending_rows = (
        (
            _row(
                1,
                name="قيد الإكمال",
                uid="987654321",
                nationality="",
                violation_date=date(2026, 7, 12),
            ),
        )
        if pending
        else ()
    )
    return (
        RegisterSection("pending", pending_rows),
        RegisterSection("expats", (expat,)),
        RegisterSection("citizens", (citizen,)),
    )


def _workbook(
    *,
    pending: bool = True,
    language: str = "ar",
    closed_at: datetime | None = None,
    closed_by_name: str | None = None,
    legacy_force_reason: str | None = None,
) -> Any:
    payload = build_register_workbook(
        year=2026,
        month=7,
        sections=_sections(pending=pending),
        closed_at=closed_at,
        closed_by_name=closed_by_name,
        legacy_force_reason=legacy_force_reason,
        language=language,
    )
    return load_workbook(BytesIO(payload))


def _summary_values(sheet: Any) -> dict[object, object]:
    return {
        sheet.cell(row, 1).value: sheet.cell(row, 2).value for row in range(2, sheet.max_row + 1)
    }


def test_three_section_closed_month_has_fixed_register_geometry() -> None:
    workbook = _workbook(
        closed_at=datetime(2026, 8, 1, 10, 30, tzinfo=UTC),
        closed_by_name="مدير النظام",
        legacy_force_reason="اعتماد الشهر رغم نقص الجنسية",
    )

    assert workbook.sheetnames == ["المواطنون", "الوافدون", "قيد الإكمال", "الملخص"]
    assert all(sheet.sheet_view.rightToLeft for sheet in workbook.worksheets)
    assert all(sheet.freeze_panes == "A2" for sheet in workbook.worksheets)

    citizens = workbook["المواطنون"]
    expats = workbook["الوافدون"]
    assert [expats.cell(1, column).value for column in range(1, 8)] == [
        "ت",
        "الإسم",
        "الرقم الموحد",
        "الجنسية",
        "تاريخ المخالفة",
        "السربة",
        "تفاصيل المخالفة",
    ]
    uid = citizens["C2"]
    assert uid.value == "784-1993-4471820-3"
    assert uid.data_type == "s"
    assert uid.number_format == "@"
    violation_date = citizens["D2"]
    assert violation_date.is_date
    assert violation_date.value.date() == date(2026, 7, 4)
    assert citizens["A3"].value == "عدد المخالفات: 1"
    assert expats["A3"].value == "عدد المخالفات: 1"
    assert workbook["قيد الإكمال"]["A3"].value == "عدد المخالفات: 1"

    summary = _summary_values(workbook["الملخص"])
    assert summary["المواطنون"] == 1
    assert summary["الوافدون"] == 1
    assert summary["قيد الإكمال"] == 1
    assert summary["إجمالي مخالفات الشهر"] == 3
    assert summary["الحالة"] == "مغلق"
    assert summary["تاريخ الإغلاق"] == datetime(2026, 8, 1, 10, 30)
    assert summary["أغلق بواسطة"] == "مدير النظام"
    assert "إغلاق إداري مع سجلات ناقصة" in summary
    assert summary["السبب"] == "اعتماد الشهر رغم نقص الجنسية"


def test_empty_pending_sheet_is_omitted_and_open_month_is_live() -> None:
    workbook = _workbook(pending=False, language="en")

    assert workbook.sheetnames == ["المواطنون", "الوافدون", "الملخص"]
    assert [workbook["المواطنون"].cell(1, column).value for column in range(1, 7)] == [
        "No.",
        "Name",
        "Unified No.",
        "Violation date",
        "Duty company",
        "Violation details",
    ]
    assert workbook["المواطنون"]["A3"].value == "Violation count: 1"
    summary = _summary_values(workbook["الملخص"])
    # The pending line follows its worksheet: absent when the group is empty.
    assert "Pending completion" not in summary
    assert summary["Month total"] == 2
    assert summary["Status"] == "Open — live projection"


def test_register_filenames_use_the_requested_language() -> None:
    assert register_filename(2026, 7, "ar") == ("سجل مخالفات المحتجزين يوليو 2026.xlsx")
    assert register_filename(2026, 7, "en") == ("Inmate Violations Register July 2026.xlsx")
