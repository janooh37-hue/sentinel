"""Reply, derivation, live references, confidentiality floors and reference sends."""

from __future__ import annotations

import json
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta, timezone
from pathlib import Path

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.config import Settings
from app.db.models import (
    AuditLog,
    Book,
    BookCategory,
    BookEditSession,
    BookLink,
    BookRecipient,
    BookVersion,
    Document,
    Employee,
    User,
    UserPermission,
)
from app.db.session import get_db
from app.main import create_app
from app.services import (
    artifact_service,
    document_service,
    perm_service,
    push_service,
    word_book_service,
)
from tests.conftest import make_user

API = "/api/v1/books"
TEMPLATES = Path(__file__).resolve().parents[1] / "templates"


@dataclass
class LinkApi:
    db: Session
    client: TestClient
    users: dict[str, User]
    records: dict[str, Book]
    actor: User

    def as_user(self, name: str) -> None:
        self.actor = self.users[name]

    def derive(self, original: str, **payload):
        return self.client.post(
            f"{API}/{self.records[original].id}/derive", json={"kind": "reply", **payload}
        )

    def thread(self, book_id: int) -> list[dict]:
        response = self.client.get(f"{API}/{book_id}/thread")
        assert response.status_code == 200, response.text
        return response.json()

    def working_document(self, book_id: int):
        session = self.db.scalar(select(BookEditSession).where(BookEditSession.book_id == book_id))
        assert session is not None
        return WordDocument(session.working_path)


def _error(response, status: int, code: str) -> dict:
    assert response.status_code == status, response.text
    error = response.json()["error"]
    assert error["code"] == code, response.text
    return error


@pytest.fixture
def links_api(
    api_db: Session, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> Iterator[LinkApi]:
    users = {
        name: make_user(api_db, role=role, email=f"{name}@links.test")
        for name, role in (
            ("creator", "operator"),
            ("sender", "operator"),
            ("receiver", "operator"),
            ("cleared", "operator"),
            ("admin", "admin"),
            ("reporter", "inmate_reporter"),
        )
    }
    api_db.add(BookCategory(id="HR", prefix="HR", requires_approval=False))
    api_db.add(BookCategory(id="GS", prefix="GS", requires_approval=False))
    for name in ("creator", "sender", "receiver"):
        employee_id = f"G-links-{name}"
        api_db.add(Employee(id=employee_id, name_en=name, name_ar=f"اسم {name}"))
        users[name].employee_id = employee_id
    api_db.add_all(
        UserPermission(user_id=users[name].id, capability="books.secret", effect="grant")
        for name in ("creator", "cleared")
    )
    api_db.add_all(
        UserPermission(user_id=users[name].id, capability="books.create", effect="grant")
        for name in ("creator", "receiver")
    )
    api_db.add_all(
        UserPermission(user_id=users[name].id, capability="books.edit", effect="grant")
        for name in ("creator", "receiver")
    )
    api_db.commit()
    records = {}
    for index, level in enumerate(("general", "directed", "personal", "secret")):
        record = Book(
            category_id="HR",
            ref_number=f"HR-links-{level}",
            subject="موضوع أصلي",
            confidentiality=level,
            created_by_user_id=users["creator"].id,
            created_at=datetime(2026, 9, 21) + timedelta(days=index),
            classification_code="5/1",
        )
        document = Document(
            template_id="General Book",
            ref_number=record.ref_number,
            submission_id=f"links-{level}",
        )
        api_db.add_all([record, document])
        api_db.flush()
        api_db.add(
            BookVersion(
                book_id=record.id,
                version_no=1,
                template_id="General Book",
                document_id=document.id,
                fields={"date": "21-09-2026", "body": "<p>نص أصلي للاقتباس</p>"},
                status="none",
                created_by_user_id=users["creator"].id,
            )
        )
        records[level] = record
    api_db.commit()
    settings = Settings(data_dir=tmp_path / "data", templates_dir=TEMPLATES)
    for service in (word_book_service, artifact_service, document_service):
        monkeypatch.setattr(service, "get_settings", lambda: settings)
    monkeypatch.setattr(push_service, "send_to_user", lambda *args: None)
    app = create_app()
    app.dependency_overrides[get_db] = lambda: api_db
    with TestClient(app) as client:
        api = LinkApi(api_db, client, users, records, users["creator"])
        app.dependency_overrides[get_current_user] = lambda: api.actor
        yield api


@pytest.mark.parametrize("actor", ["sender", "reporter"])
def test_links_require_edit_capability(links_api: LinkApi, actor: str) -> None:
    api = links_api
    source, target = api.records["general"], api.records["directed"]
    target.confidentiality = "general"
    row = BookLink(book_id=source.id, target_book_id=target.id, kind="reference", attached=True)
    api.db.add(row)
    api.db.commit()
    api.as_user(actor)
    assert perm_service.has_capability(api.db, api.actor, "books.view")
    assert not perm_service.has_capability(api.db, api.actor, "books.edit")
    if actor == "sender":
        assert api.thread(source.id)[0]["can_unlink"] is False
    for response in (
        api.client.post(f"{API}/{source.id}/links", json={"target_book_id": target.id}),
        api.client.delete(f"{API}/{source.id}/links/{target.id}"),
    ):
        _error(response, 403, "FORBIDDEN")
    api.db.refresh(row)
    assert row.attached is True


@pytest.mark.parametrize("kind", ["reference", "reply"])
@pytest.mark.parametrize("hidden", ["personal", "secret", "directed", "deleted"])
def test_unlink_hidden_target_preserves_link(links_api: LinkApi, kind: str, hidden: str) -> None:
    api = links_api
    source = api.records["general"]
    target = api.records["directed" if hidden == "deleted" else hidden]
    if hidden == "deleted":
        target.deleted_at = datetime(2026, 10, 1)
    row = BookLink(book_id=source.id, target_book_id=target.id, kind=kind, attached=True)
    api.db.add(row)
    api.db.commit()
    api.as_user("receiver")
    response = api.client.delete(f"{API}/{source.id}/links/{target.id}")
    _error(response, 404, "LINK_NOT_FOUND")
    api.db.refresh(row)
    assert row.attached is True
    _error(api.client.delete(f"{API}/{source.id}/links/999999"), 404, "LINK_NOT_FOUND")


def test_link_hidden_target_and_read_only_source(links_api: LinkApi) -> None:
    api = links_api
    api.as_user("receiver")
    source = api.records["general"]
    for level in ("personal", "secret", "directed"):
        _error(
            api.client.post(
                f"{API}/{source.id}/links", json={"target_book_id": api.records[level].id}
            ),
            404,
            "BOOK_NOT_FOUND",
        )
    directed = api.records["directed"]
    api.db.add(
        BookRecipient(
            book_id=directed.id,
            user_id=api.actor.id,
            sent_by_user_id=api.users["sender"].id,
            purpose="info",
        )
    )
    api.db.commit()
    assert api.client.get(f"{API}/{directed.id}/thread").status_code == 200
    for response in (
        api.client.post(f"{API}/{directed.id}/links", json={"target_book_id": source.id}),
        api.client.delete(f"{API}/{directed.id}/links/{source.id}"),
    ):
        _error(response, 404, "BOOK_NOT_FOUND")
    assert api.db.scalar(select(BookLink)) is None


def test_thread_filters_hidden_both_directions_and_orders_records(links_api: LinkApi) -> None:
    api = links_api
    root = api.records["general"]
    api.db.add_all(
        [
            BookLink(
                book_id=root.id,
                target_book_id=api.records["secret"].id,
                kind="reference",
                attached=True,
            ),
            BookLink(
                book_id=api.records["personal"].id,
                target_book_id=root.id,
                kind="based_on",
                attached=True,
            ),
            BookLink(
                book_id=api.records["directed"].id,
                target_book_id=root.id,
                kind="reply",
                attached=False,
            ),
        ]
    )
    api.db.commit()
    api.as_user("admin")
    items = api.thread(root.id)
    assert [item["book_id"] for item in items] == [
        api.records[level].id for level in ("directed", "personal", "secret")
    ]
    assert [item["direction"] for item in items] == ["incoming", "incoming", "outgoing"]
    assert [item["can_unlink"] for item in items] == [False, False, True]
    assert all(item["current_document_id"] is not None for item in items)
    assert all(item["addressee_user_id"] is None for item in items)
    api.as_user("receiver")
    assert api.thread(root.id) == []
    api.db.add(
        BookRecipient(
            book_id=api.records["directed"].id,
            user_id=api.actor.id,
            sent_by_user_id=api.users["creator"].id,
            purpose="info",
        )
    )
    api.db.commit()
    assert [item["book_id"] for item in api.thread(root.id)] == [api.records["directed"].id]
    assert api.thread(api.records["directed"].id)[0]["can_unlink"] is False


def test_link_errors_delete_and_reattach_preserve_derivation(links_api: LinkApi) -> None:
    api = links_api
    source, target = api.records["general"], api.records["directed"]
    url = f"{API}/{source.id}/links"
    _error(api.client.post(url, json={"target_book_id": source.id}), 422, "LINK_SELF")
    response = api.client.post(url, json={"target_book_id": target.id})
    assert response.status_code == 201, response.text
    assert response.json()["kind"] == "reference"
    _error(api.client.post(url, json={"target_book_id": target.id}), 409, "LINK_EXISTS")
    assert api.client.delete(f"{url}/{target.id}").status_code == 204
    assert api.thread(source.id) == []
    _error(api.client.delete(f"{url}/{target.id}"), 404, "LINK_NOT_FOUND")
    # A detached reply is not deleted from the thread, and reattaching retains its kind.
    source.confidentiality = "directed"
    api.db.add(BookLink(book_id=source.id, target_book_id=target.id, kind="reply", attached=True))
    api.db.commit()
    assert api.client.delete(f"{url}/{target.id}").status_code == 204
    item = api.thread(source.id)[0]
    assert item["kind"] == "reply" and item["attached"] is False
    assert item["can_unlink"] is False
    response = api.client.post(url, json={"target_book_id": target.id})
    assert response.status_code == 201, response.text
    assert response.json()["kind"] == "reply" and response.json()["attached"] is True


@pytest.mark.parametrize("kind", ["reply", "based_on"])
@pytest.mark.parametrize(
    ("original", "level", "attached", "minimum"),
    [
        ("directed", "general", False, "directed"),
        ("personal", "general", True, "directed"),
        ("secret", "general", False, "directed"),
        ("secret", "directed", True, "secret"),
        ("secret", "personal", True, "secret"),
    ],
)
def test_derive_floor_saves_nothing(
    links_api: LinkApi, kind: str, original: str, level: str, attached: bool, minimum: str
) -> None:
    api = links_api
    ids_before = set(api.db.scalars(select(Book.id)))
    error = _error(
        api.derive(original, kind=kind, confidentiality=level, attach_reference=attached),
        422,
        "LEVEL_BELOW_FLOOR",
    )
    assert error["details"] == {"min": minimum}
    assert set(api.db.scalars(select(Book.id))) == ids_before
    assert api.db.scalar(select(BookLink)) is None


@pytest.mark.parametrize("attached", [True, False])
def test_secret_derive_and_later_relabel_floor(links_api: LinkApi, attached: bool) -> None:
    api = links_api
    level = "secret" if attached else "directed"
    response = api.derive("secret", confidentiality=level, attach_reference=attached)
    assert response.status_code == 201, response.text
    book_id = response.json()["book_id"]
    assert api.db.get(Book, book_id).confidentiality == level
    url = f"{API}/{book_id}/confidentiality"
    error = _error(
        api.client.put(url, json={"confidentiality": "general"}),
        422,
        "LEVEL_BELOW_FLOOR",
    )
    assert error["details"] == {"min": "secret" if attached else "directed"}
    if attached:
        _error(api.client.put(url, json={"confidentiality": "directed"}), 422, "LEVEL_BELOW_FLOOR")
        assert (
            api.client.delete(f"{API}/{book_id}/links/{api.records['secret'].id}").status_code
            == 204
        )
    response = api.client.put(url, json={"confidentiality": "directed"})
    assert response.status_code == 200, response.text
    _error(api.client.put(url, json={"confidentiality": "general"}), 422, "LEVEL_BELOW_FLOOR")
    # Reattaching the secret original is also floor-checked and leaves the row detached.
    error = _error(
        api.client.post(
            f"{API}/{book_id}/links", json={"target_book_id": api.records["secret"].id}
        ),
        422,
        "LEVEL_BELOW_FLOOR",
    )
    assert error["details"] == {"min": "secret"}
    assert api.thread(book_id)[0]["attached"] is False


def test_attaching_secret_floor_is_atomic_and_relabel_respects_reference(
    links_api: LinkApi,
) -> None:
    api = links_api
    source, target = api.records["general"], api.records["secret"]
    url = f"{API}/{source.id}/links"
    error = _error(
        api.client.post(url, json={"target_book_id": target.id}), 422, "LEVEL_BELOW_FLOOR"
    )
    assert error["details"] == {"min": "secret"}
    assert api.db.scalar(select(BookLink)) is None
    source.confidentiality = "secret"
    api.db.commit()
    assert api.client.post(url, json={"target_book_id": target.id}).status_code == 201
    _error(
        api.client.put(f"{API}/{source.id}/confidentiality", json={"confidentiality": "personal"}),
        422,
        "LEVEL_BELOW_FLOOR",
    )
    assert api.client.delete(f"{url}/{target.id}").status_code == 204
    response = api.client.put(
        f"{API}/{source.id}/confidentiality", json={"confidentiality": "general"}
    )
    assert response.status_code == 200, response.text


@pytest.mark.parametrize("received", [True, False])
def test_reply_prefill_arabic_addressee_link_audit_and_printed_reference(
    links_api: LinkApi, received: bool
) -> None:
    api = links_api
    original = api.records["general"]
    api.as_user("receiver")
    if received:
        api.db.add_all(
            [
                BookRecipient(
                    book_id=original.id,
                    user_id=api.actor.id,
                    sent_by_user_id=api.users["creator"].id,
                    purpose="info",
                    created_at=datetime(2026, 9, 20),
                    revoked_at=datetime(2026, 9, 21),
                ),
                BookRecipient(
                    book_id=original.id,
                    user_id=api.actor.id,
                    sent_by_user_id=api.users["sender"].id,
                    purpose="reply",
                    created_at=datetime(2026, 9, 22),
                ),
            ]
        )
        api.db.commit()
    response = api.derive("general")
    assert response.status_code == 201, response.text
    new_id = response.json()["book_id"]
    book = api.db.get(Book, new_id)
    assert book.subject == "رد: موضوع أصلي"
    assert book.classification_code == original.classification_code
    assert book.confidentiality == original.confidentiality
    assert book.created_by_user_id == api.actor.id
    addressee = api.users["sender" if received else "creator"]
    text = "\n".join(p.text for p in api.working_document(new_id).paragraphs)
    assert f"اسم {'sender' if received else 'creator'}" in text
    assert "المرجع: HR-links-general في 21-09-2026" in text
    assert "نص أصلي للاقتباس" not in text, "Reply does not copy the original body"
    item = api.thread(new_id)[0]
    assert item["kind"] == "reply" and item["attached"] is True
    assert item["addressee_user_id"] == addressee.id and item["can_unlink"] is True
    audit = api.db.scalar(
        select(AuditLog).where(
            AuditLog.action == "book_derived",
            AuditLog.entity_id == str(new_id),
        )
    )
    assert audit is not None
    assert json.loads(audit.payload) == {
        "original_book_id": original.id,
        "kind": "reply",
        "attached": True,
    }


@pytest.mark.parametrize(
    "subject", ["رد: Mixed Case", " Re : Mixed Case", "الرد: Mixed Case", "رد: Re: Mixed Case"]
)
def test_reply_subject_has_one_prefix_without_case_folding(
    links_api: LinkApi, subject: str
) -> None:
    api = links_api
    api.records["general"].subject = subject
    api.db.commit()
    response = api.derive("general")
    assert response.status_code == 201, response.text
    assert api.db.get(Book, response.json()["book_id"]).subject == "رد: Mixed Case"


def test_based_on_copies_structured_general_book_body(links_api: LinkApi) -> None:
    api = links_api
    response = api.derive("general", kind="based_on", attach_reference=False)
    assert response.status_code == 201, response.text
    book_id = response.json()["book_id"]
    assert api.db.get(Book, book_id).subject == api.records["general"].subject
    text = "\n".join(p.text for p in api.working_document(book_id).paragraphs)
    assert "نص أصلي للاقتباس" in text
    assert "المرجع: HR-links-general في 21-09-2026" in text
    item = api.thread(book_id)[0]
    assert item["kind"] == "based_on" and item["attached"] is False
    assert item["addressee_user_id"] is None


@pytest.mark.parametrize("active_draft", [False, True])
def test_based_on_copies_word_authored_body_tables_images_and_links(
    links_api: LinkApi, tmp_path: Path, active_draft: bool
) -> None:
    from copy import deepcopy
    from io import BytesIO

    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from PIL import Image

    api = links_api
    response = api.derive("general")
    assert response.status_code == 201, response.text
    source_id = response.json()["book_id"]
    session = api.db.scalar(select(BookEditSession).where(BookEditSession.book_id == source_id))
    source = WordDocument(session.working_path)
    closing = next(p for p in source.paragraphs if p.text.strip() == "للتفضل بالعلم وإجراءاتكم")
    paragraph = closing.insert_paragraph_before("نص مؤلف في Word {{ literal }}")
    paragraph.runs[0].bold = True
    paragraph.style = source.styles.add_style("Authored body", 1)
    image = BytesIO()
    Image.new("RGB", (2, 2), "red").save(image, format="PNG")
    image.seek(0)
    paragraph.add_run().add_picture(image)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(
        qn("r:id"), source.part.relate_to("https://example.test/body", RT.HYPERLINK, True)
    )
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "رابط أصلي"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "خلية أصلية"
    closing._p.addprevious(table._tbl)
    source.part.relate_to(WordDocument().part.numbering_part, RT.NUMBERING)
    numbering = source.part.numbering_part.element
    abstract = deepcopy(WordDocument().part.numbering_part.element.find(qn("w:abstractNum")))
    abstract_id = (
        max(
            (int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))),
            default=-1,
        )
        + 1
    )
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    numbering.append(abstract)
    num = numbering.add_num(abstract_id)
    list_style = source.styles.add_style("Authored list", 1)
    list_style.element.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = num.numId
    closing.insert_paragraph_before("عنصر أصلي", style=list_style)
    source.save(session.working_path)
    source_book = api.db.get(Book, source_id)
    if not active_draft:
        document = Document(
            template_id="General Book",
            ref_number=source_book.ref_number,
            submission_id=f"authored-{source_id}",
            docx_path=session.working_path,
        )
        api.db.add(document)
        api.db.flush()
        api.db.add(
            BookVersion(
                book_id=source_id,
                version_no=1,
                template_id="General Book",
                document_id=document.id,
                fields={},
                status="none",
            )
        )
        session.state = "finished"
    api.db.commit()
    source_bytes = Path(session.working_path).read_bytes()
    response = api.client.post(f"{API}/{source_id}/derive", json={"kind": "based_on"})
    assert response.status_code == 201, response.text
    copied = api.working_document(response.json()["book_id"])
    body = next(p for p in copied.paragraphs if "نص مؤلف في Word" in p.text)
    assert "{{ literal }}" in body.text
    assert body.runs[0].bold and body.style.name == "Authored body"
    assert [table.cell(0, 0).text for table in copied.tables] == ["خلية أصلية"]
    image_id = body._p.find(".//" + qn("a:blip")).get(qn("r:embed"))
    assert copied.part.rels[image_id].target_part.blob == image.getvalue()
    link_id = body._p.find(qn("w:hyperlink")).get(qn("r:id"))
    assert copied.part.rels[link_id].target_ref == "https://example.test/body"
    copied_bullet = next(p for p in copied.paragraphs if p.text == "عنصر أصلي")
    num_id = copied_bullet.style.element.find(".//" + qn("w:numId")).get(qn("w:val"))
    assert (
        copied.part.numbering_part.element.find(f"./{qn('w:num')}[@{qn('w:numId')}='{num_id}']")
        is not None
    )
    assert sum(p.text.startswith("المرجع:") for p in copied.paragraphs) == 1
    assert any(f"المرجع: {source_book.ref_number} في " in p.text for p in copied.paragraphs)
    assert Path(session.working_path).read_bytes() == source_bytes


def test_based_on_non_general_book_copies_subject_only(links_api: LinkApi) -> None:
    api = links_api
    version = api.records["general"].versions[-1]
    version.template_id = "Report"
    api.db.commit()
    response = api.derive("general", kind="based_on")
    assert response.status_code == 201, response.text
    new_id = response.json()["book_id"]
    assert api.db.get(Book, new_id).subject == api.records["general"].subject
    assert "نص أصلي للاقتباس" not in "\n".join(
        p.text for p in api.working_document(new_id).paragraphs
    )


@pytest.mark.parametrize("paper_date", [None, "not-a-date"])
def test_reference_fallback_uses_local_creation_date(
    links_api: LinkApi, paper_date: str | None
) -> None:
    class OfficeDatetime(datetime):
        def astimezone(self, tz=None):
            return super().astimezone(tz or timezone(timedelta(hours=4)))

    api = links_api
    original = api.records["general"]
    version = original.versions[-1]
    version.template_id = "Report"
    version.fields = {"date": paper_date}
    api.db.commit()
    # Exercise the local midnight rollover without depending on the host timezone.
    original.created_at = OfficeDatetime(2026, 9, 21, 22, 30)
    assert original.created_at.replace(tzinfo=UTC).astimezone().day == 22
    response = api.derive("general", kind="based_on")
    assert response.status_code == 201, response.text
    text = "\n".join(p.text for p in api.working_document(response.json()["book_id"]).paragraphs)
    assert "المرجع: HR-links-general في 22-09-2026" in text


@pytest.mark.parametrize("kind", ["reply", "based_on"])
def test_word_original_reference_uses_printed_date_not_creation_date(
    links_api: LinkApi, tmp_path: Path, kind: str
) -> None:
    from app.core.docx_engine import DocxEngine

    api = links_api
    original = api.records["general"]
    authored = tmp_path / "authored.docx"
    DocxEngine(TEMPLATES).fill(
        "General Book",
        {
            "ref": original.ref_number,
            "date": "12-08-2025",
            "subject": original.subject,
            "body": "",
            "recipient_name": "",
            "cc": [],
        },
        authored,
    )
    version = original.versions[-1]
    version.fields = {}
    api.db.get(Document, version.document_id).docx_path = str(authored)
    api.db.commit()
    response = api.derive("general", kind=kind)
    assert response.status_code == 201, response.text
    text = "\n".join(p.text for p in api.working_document(response.json()["book_id"]).paragraphs)
    assert "المرجع: HR-links-general في 12-08-2025" in text


def test_send_personal_to_directed_transition_respects_secret_reference_floor(
    links_api: LinkApi,
) -> None:
    api = links_api
    main, reference = api.records["personal"], api.records["secret"]
    reference.confidentiality = "general"
    api.db.commit()
    response = api.client.post(f"{API}/{main.id}/links", json={"target_book_id": reference.id})
    assert response.status_code == 201, response.text
    reference.confidentiality = "secret"
    api.db.commit()
    error = _error(
        api.client.post(
            f"{API}/{main.id}/send",
            json={"user_ids": [api.users["receiver"].id], "purpose": "info", "make_directed": True},
        ),
        422,
        "LEVEL_BELOW_FLOOR",
    )
    assert error["details"] == {"min": "secret"}
    assert api.db.get(Book, main.id).confidentiality == "personal"
    assert api.db.scalar(select(BookRecipient)) is None


def test_derive_access_create_clearance_and_classification_checks(links_api: LinkApi) -> None:
    api = links_api
    api.as_user("receiver")
    _error(api.derive("secret"), 404, "BOOK_NOT_FOUND")
    _error(api.derive("general", confidentiality="secret"), 403, "SECRET_CLEARANCE_REQUIRED")
    perm_service.set_user_override(api.db, api.actor.id, "books.create", "deny")
    _error(api.derive("general"), 403, "FORBIDDEN")
    api.as_user("creator")
    api.records["general"].classification_code = None
    api.db.commit()
    _error(api.derive("general"), 422, "CLASSIFICATION_REQUIRED")
    response = api.derive("general", classification_code="5/1")
    assert response.status_code == 201, response.text
    api.as_user("reporter")
    _error(api.derive("general"), 403, "FORBIDDEN")


@pytest.mark.parametrize(
    ("capability", "error_code"),
    [
        ("books.create", "FORBIDDEN"),
        ("books.view", "FORBIDDEN"),
        ("books.category.GS", "RECORD_TYPE_FORBIDDEN"),
    ],
)
def test_read_grant_does_not_bypass_general_book_creation_permissions(
    links_api: LinkApi, capability: str, error_code: str
) -> None:
    api = links_api
    api.as_user("receiver")
    original = api.records["directed"]
    api.db.add(
        BookRecipient(
            book_id=original.id,
            user_id=api.actor.id,
            sent_by_user_id=api.users["sender"].id,
            purpose="reply",
        )
    )
    perm_service.set_user_override(api.db, api.actor.id, capability, "deny")
    ids_before = set(api.db.scalars(select(Book.id)))
    _error(api.derive("directed"), 403, error_code)
    assert set(api.db.scalars(select(Book.id))) == ids_before
    assert api.db.scalar(select(BookLink)) is None


def test_thread_omits_soft_deleted_links_and_deleted_root_is_hidden(links_api: LinkApi) -> None:
    api = links_api
    root, target = api.records["general"], api.records["directed"]
    response = api.client.post(f"{API}/{root.id}/links", json={"target_book_id": target.id})
    assert response.status_code == 201, response.text
    target.deleted_at = datetime(2026, 10, 1)
    api.db.commit()
    assert api.thread(root.id) == []
    _error(api.client.get(f"{API}/{target.id}/thread"), 404, "BOOK_NOT_FOUND")


def test_send_references_skips_uncleared_personal_existing_and_self(links_api: LinkApi) -> None:
    api = links_api
    main, reference = api.records["general"], api.records["secret"]
    # The original can become secret after a general record has referenced it.
    reference.confidentiality = "general"
    api.db.commit()
    for target in (reference, api.records["personal"], api.records["directed"]):
        response = api.client.post(f"{API}/{main.id}/links", json={"target_book_id": target.id})
        assert response.status_code == 201, response.text
    response = api.client.put(
        f"{API}/{reference.id}/confidentiality", json={"confidentiality": "secret"}
    )
    assert response.status_code == 200, response.text
    api.db.add(
        BookRecipient(
            book_id=api.records["directed"].id,
            user_id=api.users["receiver"].id,
            sent_by_user_id=api.actor.id,
            purpose="info",
        )
    )
    api.db.commit()
    user_ids = [api.users[name].id for name in ("receiver", "cleared", "creator")]
    response = api.client.post(
        f"{API}/{main.id}/send",
        json={"user_ids": user_ids, "purpose": "reply", "note": "نفس الملاحظة"},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert {row["user_id"] for row in result["sent"]} == set(user_ids[:-1])
    references = {item["book_id"]: item for item in result["references"]}
    secret = references[reference.id]
    assert secret["sent_user_ids"] == [api.users["cleared"].id]
    assert secret["skipped"] == [
        {"user_id": api.users["receiver"].id, "reason": "not_cleared"},
        {"user_id": api.actor.id, "reason": "self"},
    ]
    personal = references[api.records["personal"].id]
    assert personal["sent_user_ids"] == []
    assert [skip["reason"] for skip in personal["skipped"]] == ["personal", "personal", "self"]
    directed = references[api.records["directed"].id]
    assert directed["sent_user_ids"] == [api.users["cleared"].id]
    assert [skip["reason"] for skip in directed["skipped"]] == ["already_recipient", "self"]
    for book_id in (reference.id, api.records["directed"].id):
        sent = api.db.scalar(
            select(BookRecipient).where(
                BookRecipient.book_id == book_id,
                BookRecipient.user_id == api.users["cleared"].id,
            )
        )
        assert sent.purpose == "reply" and sent.note == "نفس الملاحظة"


def test_send_reference_opt_out_and_hidden_sender_references(links_api: LinkApi) -> None:
    api = links_api
    main = api.records["general"]
    for target in (api.records["directed"], api.records["personal"]):
        assert (
            api.client.post(
                f"{API}/{main.id}/links", json={"target_book_id": target.id}
            ).status_code
            == 201
        )
    response = api.client.post(
        f"{API}/{main.id}/send",
        json={
            "user_ids": [api.users["receiver"].id],
            "purpose": "info",
            "include_references": False,
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["references"] == []
    assert not api.db.scalars(select(BookRecipient).where(BookRecipient.book_id != main.id)).all()
    api.as_user("receiver")
    response = api.client.post(
        f"{API}/{main.id}/send", json={"user_ids": [api.users["cleared"].id], "purpose": "info"}
    )
    assert response.status_code == 200, response.text
    assert response.json()["references"] == []


def test_detached_derivation_stays_in_thread_but_is_not_sent(links_api: LinkApi) -> None:
    api = links_api
    response = api.derive("general", attach_reference=False)
    assert response.status_code == 201, response.text
    draft_id = response.json()["book_id"]
    response = api.client.post(
        f"{API}/{draft_id}/send",
        json={"user_ids": [api.users["receiver"].id], "purpose": "reply"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["references"] == []
    assert api.thread(draft_id)[0]["attached"] is False
    assert (
        api.db.scalar(
            select(BookRecipient).where(
                BookRecipient.book_id == api.records["general"].id,
            )
        )
        is None
    )
