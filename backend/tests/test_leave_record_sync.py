"""Leaves date edits append artifacts without restarting the record's workflow."""

import json
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import docx
import fitz
import pytest
from PIL import Image
from sqlalchemy import select

from app.api.errors import ValidationFailedError
from app.api.v1.books import _signed_pdf_url_of
from app.api.v1.documents import GenerateAttachmentSpec
from app.config import Settings, get_settings
from app.core import docx_engine, docx_render, signature_layout
from app.core import signature as signature_core
from app.db.models import (
    AuditLog,
    Book,
    BookApprovalStep,
    BookCategory,
    BookRevisionAccess,
    Document,
    Employee,
    Leave,
    Manager,
    SignatureArtifactRevision,
    Submitter,
    User,
)
from app.schemas.leave import LeaveUpdate
from app.services import (
    artifact_service,
    book_service,
    document_service,
    included_papers_service,
    leave_record_sync,
    leave_service,
    signature_placement_service,
    staging_service,
)


@pytest.fixture()
def env(db_session, tmp_path, monkeypatch):
    settings = Settings(data_dir=tmp_path / "data", templates_dir=document_service._TEMPLATES_DIR)
    monkeypatch.setenv("GSSG_DATA_DIR", str(settings.data_dir))
    get_settings.cache_clear()
    for module in (document_service, artifact_service, book_service):
        monkeypatch.setattr(module, "get_settings", lambda: settings)

    def convert(path):
        document = docx.Document(path)
        text = "\n".join(p.text for p in document.paragraphs)
        text += "\n" + "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
        pdf = fitz.open()
        page = pdf.new_page()
        # Keep the real filled DOCX content observable through a valid fake PDF.
        page.insert_textbox(
            fitz.Rect(20, 20, 575, 820), text.encode("ascii", "replace").decode(), fontsize=6
        )
        target = path.with_suffix(".pdf")
        pdf.save(target)
        pdf.close()
        return target

    monkeypatch.setattr(document_service, "convert_docx_to_pdf", convert)
    employee = Employee(id="G-SYNC", name_en="Sync Employee", status="Active")
    signer = User(
        email="sync@example.test",
        password_hash="x",
        role="admin",
        status="active",
        display_name="Sync Manager",
    )
    db_session.add_all([employee, signer, BookCategory(id="HR", prefix="HR")])
    db_session.flush()
    manager = Manager(name_en="Sync Manager", user_id=signer.id)
    db_session.add(manager)
    db_session.commit()
    signature = tmp_path / "signature.png"
    Image.new("RGBA", (100, 40), (0, 0, 0, 255)).save(signature)
    monkeypatch.setattr(
        book_service.user_signature_service, "resolve_signature", lambda _: signature
    )
    yield SimpleNamespace(
        db=db_session, settings=settings, signer=signer, employee=employee, manager=manager
    )
    get_settings.cache_clear()


def generate(
    env,
    *,
    start="20/10/2026",
    end="21/10/2026",
    count="2",
    template="Leave Application Form",
    leave_type="Annual Leave",
    attachments=None,
    embed_signature=None,
    submitter_id=None,
):
    result = document_service.generate_document(
        env.db,
        employee_id=env.employee.id,
        template_id=template,
        fields={
            "leave_type": leave_type,
            "start_date": start,
            "end_date": end,
            "total_days": count,
            "duration": count,
            "date": start,
            "admin_leaves_this_month": "1",
        },
        manager_id=env.manager.id,
        current_user=env.signer,
        attachments=attachments,
        embed_signature=embed_signature,
        submitter_id=submitter_id,
    )
    book = env.db.get_one(Book, result.book_id)
    document = env.db.get_one(Document, book.versions[-1].document_id)
    return book, env.db.get_one(Leave, document.leave_id)


def submit(env, book):
    book_service.submit_for_approval(
        env.db,
        book.id,
        submitted_by_user_id=env.signer.id,
        approver_user_id=env.signer.id,
        priority="Normal",
        reviewer_user_ids=[],
    )


def sign(env, book):
    submit(env, book)
    book_service.sign_book(env.db, book.id, user_id=env.signer.id, version_id=book.versions[-1].id)


def edit(env, leave):
    return leave_service.update_leave(
        env.db,
        leave.id,
        LeaveUpdate(start_date=date(2026, 10, 22), end_date=date(2026, 10, 25)),
        actor=env.signer.email,
    )


def text(env, path):
    with fitz.open(env.settings.data_dir / path) as pdf:
        return "\n".join(page.get_text() for page in pdf)


@pytest.mark.parametrize(
    "start,end,count,new_start,new_end,new_count",
    [
        ("20/10/2026", "21/10/2026", "2", "22/10/2026", "25/10/2026", "4"),
        ("2026-10-20", "2026-10-21", 2, "2026-10-22", "2026-10-25", 4),
        ("20-10-2026", "21-10-2026", "2", "22-10-2026", "25-10-2026", "4"),
    ],
)
def test_pending_edit_appends_form_and_companion_preserving_old_artifacts(
    env, start, end, count, new_start, new_end, new_count
):
    book, leave = generate(env, start=start, end=end, count=count)
    previous = book.versions[-1]
    old_doc = env.db.get_one(Document, previous.document_id)
    old_files = {
        p: (env.settings.data_dir / p).read_bytes()
        for p in (old_doc.docx_path, old_doc.pdf_path, old_doc.base_pdf_path)
    }
    state = book.approval_state
    edit(env, leave)
    current = book.versions[-1]
    assert current.id != previous.id
    assert current.trigger == "leave-edit"
    assert current.fields["start_date"] == new_start
    assert current.fields["end_date"] == new_end
    assert current.fields["total_days"] == new_count
    assert type(current.fields["total_days"]) is type(count)
    assert previous.fields["start_date"] == start
    assert previous.fields["end_date"] == end
    assert book.approval_state == state
    assert current.created_by_user_id == env.signer.id
    assert book_service.leave_owner(env.db, leave.id).id == book.id
    document = env.db.get_one(Document, current.document_id)
    assert document.leave_id == leave.id
    assert document.submission_id != old_doc.submission_id
    assert document.role == "primary"
    companions = list(
        env.db.scalars(
            select(Document).where(
                Document.submission_id == document.submission_id, Document.role == "companion"
            )
        )
    )
    assert len(companions) == 1
    assert companions[0].template_id == "Leave Undertaking"
    assert companions[0].leave_id is None
    assert new_start in text(env, companions[0].pdf_path)
    for path, data in old_files.items():
        assert (env.settings.data_dir / path).read_bytes() == data
    audit = env.db.scalar(
        select(AuditLog).where(
            AuditLog.entity_type == "book",
            AuditLog.entity_id == str(book.id),
            AuditLog.action == "book.leave_dates_synced",
        )
    )
    assert audit.actor == env.signer.email
    payload = json.loads(audit.payload)
    assert payload["actor"] == env.signer.email
    assert payload["from"] == {"start": "2026-10-20", "end": "2026-10-21", "days": 2}
    assert payload["to"] == {"start": "2026-10-22", "end": "2026-10-25", "days": 4}
    assert payload["source"] == "leaves"
    assert payload["note"] == "edited from Leaves"


def test_approved_record_with_generation_embedded_manager_signature_syncs(env):
    # The common production shape: approved at generation with the manager
    # signature embedded, never signed through sign_book.
    book, leave = generate(env)
    previous = book.versions[-1]
    previous.status = "approved"
    previous.manager_sig_embedded = True
    book.approval_state = "approved"
    env.db.commit()
    assert previous.signed_pdf_path is None and previous.signed_at is None
    edit(env, leave)
    current = book.versions[-1]
    assert current.id != previous.id
    assert current.status == book.approval_state == "approved"
    assert current.manager_sig_embedded is True
    assert current.signed_pdf_path is None and current.signed_by_user_id is None
    assert (current.fields["start_date"], current.fields["end_date"]) == (
        "22/10/2026",
        "25/10/2026",
    )
    document = env.db.get_one(Document, current.document_id)
    assert "25/10/2026" in text(env, document.pdf_path)
    assert env.db.scalar(
        select(AuditLog).where(
            AuditLog.entity_id == str(book.id), AuditLog.action == "book.leave_dates_synced"
        )
    )


@pytest.mark.parametrize("amend", [False, True])
def test_signed_record_keeps_signer_decisions_access_and_signature_registry(env, amend):
    book, leave = generate(env)
    assert book.versions[-1].revision_access == []
    sign(env, book)
    previous = book.versions[-1]
    assert leave.status == "Pending"
    old_signed = previous.signed_pdf_path
    old_signed_bytes = (env.settings.data_dir / old_signed).read_bytes()
    old_step = previous.approval_steps[0]
    step_id, decided_at = old_step.id, old_step.decided_at
    if amend:
        leave_service.update_leave(
            env.db, leave.id, LeaveUpdate(status="Approved"), actor=env.signer.email
        )
        leave_service.amend_approved_leave(
            env.db, leave.id, end_date=date(2026, 10, 25), reason="Longer", actor=env.signer.email
        )
    else:
        edit(env, leave)
    current = book.versions[-1]
    assert current.id != previous.id
    assert current.signed_by_user_id == previous.signed_by_user_id == env.signer.id
    assert current.signed_at == previous.signed_at
    assert current.status == book.approval_state == "approved"
    assert current.approval_context == previous.approval_context
    assert current.signed_pdf_path != old_signed
    assert "25/10/2026" in text(env, current.signed_pdf_path)
    assert (env.settings.data_dir / old_signed).read_bytes() == old_signed_bytes
    assert book_service.is_document_signed_locked(env.db, current) == (
        True,
        current.signed_pdf_path,
    )
    assert _signed_pdf_url_of(current) is not None
    steps = list(
        env.db.scalars(select(BookApprovalStep).where(BookApprovalStep.book_id == book.id))
    )
    assert len(steps) == 1
    assert steps[0].id == step_id
    assert steps[0].version_id == current.id
    assert steps[0].state == "approved"
    assert steps[0].assignee_user_id == env.signer.id
    assert steps[0].decided_at == decided_at
    for version in (previous, current):
        grant = env.db.scalar(
            select(BookRevisionAccess).where(BookRevisionAccess.version_id == version.id)
        )
        assert grant.state == "approved"
        assert grant.decided_at == decided_at
    if current.signature_revision:
        assert (
            env.db.scalar(
                select(SignatureArtifactRevision).where(
                    SignatureArtifactRevision.version_id == current.id,
                    SignatureArtifactRevision.revision == current.signature_revision,
                )
            )
            is not None
        )


def test_word_authored_record_is_untouched_and_audited(env):
    book, leave = generate(env)
    previous = book.versions[-1]
    previous.fields = {}
    env.db.commit()
    edit(env, leave)
    assert leave.days == 4
    assert book.versions == [previous]
    audit = env.db.scalar(
        select(AuditLog).where(
            AuditLog.entity_type == "book",
            AuditLog.entity_id == str(book.id),
            AuditLog.action == "book.leave_dates_not_synced",
        )
    )
    assert "word" in json.loads(audit.payload)["reason"].lower()


def test_manual_national_service_without_record_is_noop(env):
    leave = Leave(
        employee_id=env.employee.id,
        leave_type="National Service",
        start_date=date(2026, 10, 20),
        end_date=date(2026, 10, 21),
        days=2,
        status="Pending",
    )
    env.db.add(leave)
    env.db.commit()
    edit(env, leave)
    assert leave.days == 4
    assert list(env.db.scalars(select(Book))) == []


def test_pending_assignment_moves_to_current_revision_and_can_still_sign(env):
    book, leave = generate(env)
    submit(env, book)
    previous = book.versions[-1]
    step_id = previous.approval_steps[0].id
    edit(env, leave)
    current = book.versions[-1]
    assert book.approval_state == current.status == "pending"
    assert [(step.id, step.state) for step in current.approval_steps] == [(step_id, "pending")]
    assert previous.approval_steps == []
    assert [row.id for row in book_service.list_awaiting(env.db, user_id=env.signer.id)] == [
        book.id
    ]
    book_service.sign_book(env.db, book.id, user_id=env.signer.id, version_id=current.id)
    assert book.approval_state == "approved"
    assert "25/10/2026" in text(env, current.signed_pdf_path)


def test_duty_resumption_owner_is_not_replaced_by_application_sync(env):
    book, leave = generate(env)
    result = document_service.generate_document(
        env.db,
        employee_id=env.employee.id,
        template_id="Duty Resumption Form",
        fields={"date": "22/10/2026"},
        return_for_leave_id=leave.id,
        manager_id=env.manager.id,
        current_user=env.signer,
    )
    returned = env.db.get_one(Book, result.book_id)
    assert book_service.leave_owner(env.db, leave.id).id == returned.id
    document_count = env.db.query(Document).count()
    edit(env, leave)
    assert len(book.versions) == len(returned.versions) == 1
    assert book_service.leave_owner(env.db, leave.id).id == returned.id
    assert env.db.query(Document).count() == document_count
    audit = env.db.scalar(
        select(AuditLog).where(
            AuditLog.entity_type == "book",
            AuditLog.entity_id == str(returned.id),
            AuditLog.action == "book.leave_dates_not_synced",
        )
    )
    assert "mapping" in json.loads(audit.payload)["reason"]


@pytest.mark.parametrize(
    "unsafe",
    [
        "missing_signer",
        "missing_signature",
        "scan",
        "voided",
        "deleted",
        "unknown_template",
        "bad_date",
    ],
)
def test_unsafe_record_sync_is_skipped_without_blocking_leave_edit(env, monkeypatch, unsafe):
    book, leave = generate(env)
    if unsafe in ("missing_signer", "missing_signature", "scan"):
        sign(env, book)
    previous = book.versions[-1]
    if unsafe == "missing_signer":
        previous.signed_by_user_id = None
    elif unsafe == "missing_signature":
        monkeypatch.setattr(
            book_service.user_signature_service, "resolve_signature", lambda _: None
        )
    elif unsafe == "scan":
        previous.signed_base_pdf_path = "book_attachments/1/scanned.pdf"
    elif unsafe == "voided":
        book.voided_at = previous.created_at
    elif unsafe == "deleted":
        book.deleted_at = previous.created_at
    elif unsafe == "unknown_template":
        previous.template_id = "General Book"
    elif unsafe == "bad_date":
        previous.fields = {**previous.fields, "start_date": "not a date"}
    env.db.commit()
    state = book.approval_state
    edit(env, leave)
    assert leave.days == 4
    assert book.versions == [previous]
    assert book.approval_state == state
    if unsafe in ("voided", "deleted"):
        assert book_service.leave_owner(env.db, leave.id) is None
        return
    audit = env.db.scalar(
        select(AuditLog).where(
            AuditLog.entity_type == "book",
            AuditLog.entity_id == str(book.id),
            AuditLog.action == "book.leave_dates_not_synced",
        )
    )
    assert json.loads(audit.payload)["reason"]


@pytest.mark.parametrize(
    "template,count_key", [("Administrative Leave Form", "duration"), ("Leave Permit Form", None)]
)
def test_other_mapped_leave_forms_update_only_representable_dates(env, template, count_key):
    book, leave = generate(env, template=template)
    previous = book.versions[-1]
    # These record-group kinds do not expose date edits through Leaves today.
    # Exercise their mappings at the synchronization boundary itself.
    old_start, old_end, old_days = leave.start_date, leave.end_date, leave.days
    leave.start_date = date(2026, 10, 22)
    leave.end_date = date(2026, 10, 25) if count_key else date(2026, 10, 22)
    leave.days = 4 if count_key else 1
    leave_record_sync.sync_leave_record(
        env.db,
        leave,
        old_start=old_start,
        old_end=old_end,
        old_days=old_days,
        actor=env.signer.email,
    )
    env.db.commit()
    assert book.versions[-1].id != previous.id
    if count_key:
        assert book.versions[-1].fields[count_key] == "4"
        assert book.versions[-1].fields["start_date"] == "22/10/2026"
    else:
        assert book.versions[-1].fields["date"] == "22/10/2026"
        current = book.versions[-1]
        leave.end_date = date(2026, 10, 25)
        leave.days = 4
        leave_record_sync.sync_leave_record(
            env.db,
            leave,
            old_start=date(2026, 10, 22),
            old_end=date(2026, 10, 22),
            old_days=1,
            actor=env.signer.email,
        )
        env.db.commit()
        assert book.versions[-1] is current


def test_notes_only_or_unchanged_dates_do_not_append_record_revision(env):
    book, leave = generate(env)
    leave_service.update_leave(
        env.db, leave.id, LeaveUpdate(notes="Updated note"), actor=env.signer.email
    )
    assert len(book.versions) == 1
    leave_service.update_leave(
        env.db,
        leave.id,
        LeaveUpdate(start_date=leave.start_date, end_date=leave.end_date),
        actor=env.signer.email,
    )
    assert len(book.versions) == 1


@pytest.mark.parametrize("stage", ["primary", "companion", "signed"])
def test_pdf_conversion_failure_rolls_back_leave_revision_and_all_new_files(
    env, monkeypatch, stage
):
    book, leave = generate(env)
    if stage == "signed":
        sign(env, book)
    old_start, old_end, old_days = leave.start_date, leave.end_date, leave.days
    files_before = {path for path in env.settings.data_dir.rglob("*") if path.is_file()}
    versions_before = [version.id for version in book.versions]
    converter = document_service.convert_docx_to_pdf

    def convert(path):
        fail = (
            stage == "primary"
            or (stage == "companion" and "LeaveUndertaking" in path.name)
            or (stage == "signed" and "_signed" in path.name)
        )
        return None if fail else converter(path)

    monkeypatch.setattr(document_service, "convert_docx_to_pdf", convert)
    with pytest.raises(ValidationFailedError) as caught:
        edit(env, leave)
    assert caught.value.code == "GENERATION_PDF_FAILED"
    env.db.rollback()
    assert (leave.start_date, leave.end_date, leave.days) == (old_start, old_end, old_days)
    assert [version.id for version in book.versions] == versions_before
    assert {path for path in env.settings.data_dir.rglob("*") if path.is_file()} == files_before


def test_corrected_signature_placement_is_preserved_by_skipping_sync(env):
    book, leave = generate(env)
    sign(env, book)
    previous = book.versions[-1]
    previous.signature_revision = 2
    env.db.commit()
    edit(env, leave)
    assert leave.days == 4
    assert book.versions == [previous]
    audit = env.db.scalar(
        select(AuditLog).where(
            AuditLog.entity_type == "book",
            AuditLog.entity_id == str(book.id),
            AuditLog.action == "book.leave_dates_not_synced",
        )
    )
    assert "corrected signature placement" in json.loads(audit.payload)["reason"].lower()


def test_amendment_with_unchanged_end_date_does_not_create_revision(env):
    book, leave = generate(env)
    leave_service.update_leave(
        env.db, leave.id, LeaveUpdate(status="Approved"), actor=env.signer.email
    )
    leave_service.amend_approved_leave(
        env.db, leave.id, end_date=leave.end_date, reason="Clarification", actor=env.signer.email
    )
    assert len(book.versions) == 1


def test_pending_sign_gated_request_edit_still_settles_leave_when_current_revision_is_signed(env):
    book, leave = generate(env, leave_type="Emergency Leave")
    submit(env, book)
    assert leave.status == "Pending"
    edit(env, leave)
    current = book.versions[-1]
    assert current.fields["end_date"] == "25/10/2026"
    assert book.approval_state == "pending"
    book_service.sign_book(env.db, book.id, user_id=env.signer.id, version_id=current.id)
    assert book.approval_state == current.status == "approved"
    assert leave.status == "Approved"


@pytest.mark.parametrize("signed", [False, True])
def test_record_sync_preserves_included_papers_in_new_package(env, signed):
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((40, 80), "INCLUDED-PAPER")
        staged = staging_service.stage(pdf.tobytes(), "included.pdf")
    book, leave = generate(
        env,
        attachments=[
            GenerateAttachmentSpec(
                source="staged", staged_token=staged.token, original_name="included.pdf"
            ),
        ],
    )
    if signed:
        sign(env, book)
    previous = book.versions[-1]
    old_document = env.db.get_one(Document, previous.document_id)
    previous_path = previous.signed_pdf_path if signed else old_document.pdf_path
    old_bytes = (env.settings.data_dir / previous_path).read_bytes()
    edit(env, leave)
    current = book.versions[-1]
    new_document = env.db.get_one(Document, current.document_id)
    published = current.signed_pdf_path if signed else new_document.pdf_path
    assert "25/10/2026" in text(env, published)
    assert text(env, published).count("INCLUDED-PAPER") == 1
    assert (env.settings.data_dir / previous_path).read_bytes() == old_bytes
    assert "INCLUDED-PAPER" not in text(env, new_document.base_pdf_path)


@pytest.mark.parametrize(
    "failure", ["render", "generated_package", "signed_package", "signature_capture"]
)
def test_render_or_publish_exception_removes_all_attempt_files(env, monkeypatch, failure):
    book, leave = generate(env)
    if failure in ("signed_package", "signature_capture"):
        sign(env, book)
    old_dates = (leave.start_date, leave.end_date)
    previous_ids = [version.id for version in book.versions]
    files_before = {path for path in env.settings.data_dir.rglob("*") if path.is_file()}
    if failure == "render":
        produce = artifact_service.produce_from_template

        def fail_render(**kwargs):
            if kwargs["template_id"] == "Leave Undertaking":
                raise RuntimeError("template render failed")
            return produce(**kwargs)

        monkeypatch.setattr(artifact_service, "produce_from_template", fail_render)
    elif failure == "signature_capture":
        capture = signature_placement_service.capture_initial_revision

        def fail_capture(*args, **kwargs):
            capture(*args, **kwargs)
            raise RuntimeError("signature artifact registration failed")

        monkeypatch.setattr(signature_placement_service, "capture_initial_revision", fail_capture)
    else:
        write_bytes = included_papers_service._write_bytes_atomic
        calls = 0

        def fail_publish(path, data):
            nonlocal calls
            calls += 1
            if calls == (2 if failure == "generated_package" else 4):
                raise OSError("package publication failed")
            return write_bytes(path, data)

        monkeypatch.setattr(included_papers_service, "_write_bytes_atomic", fail_publish)
    with pytest.raises(ValidationFailedError) as caught:
        edit(env, leave)
    assert caught.value.code == "GENERATION_PDF_FAILED"
    # The inner savepoint already removed the record revision before outer rollback.
    assert [version.id for version in book.versions] == previous_ids
    assert {path for path in env.settings.data_dir.rglob("*") if path.is_file()} == files_before
    env.db.rollback()
    assert (leave.start_date, leave.end_date) == old_dates


class _Sept30(datetime):
    """Generation clock: the paper is issued on 30/09/2026, edited today."""

    @classmethod
    def now(cls, tz=None):
        return cls(2026, 9, 30, 10, 0, tzinfo=tz)


def generate_on_sept_30(env, monkeypatch, **kwargs):
    with monkeypatch.context() as clock:
        for module in (document_service, docx_engine, docx_render, artifact_service):
            clock.setattr(module, "datetime", _Sept30)
        return generate(env, **kwargs)


def give_employee_signature(env):
    path = signature_core.employee_signature_path(env.settings.vault_dir, env.employee.id)
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGBA", (120, 40), (0, 0, 80, 255)).save(path)
    return path


def drop_snapshot(env, book):
    """A paper generated before the ``_signing`` snapshot existed."""
    version = book.versions[-1]
    version.fields = {k: v for k, v in version.fields.items() if k != "_signing"}
    env.db.commit()


def roles(env, docx_path):
    path = env.settings.data_dir / docx_path
    return Counter(d.role for d in signature_layout.inspect_signature_drawings(path))


def companion_of(env, document):
    return env.db.scalar(
        select(Document).where(
            Document.submission_id == document.submission_id, Document.role == "companion"
        )
    )


def docx_text(env, docx_path):
    document = docx.Document(env.settings.data_dir / docx_path)
    cells = (c.text for t in document.tables for r in t.rows for c in r.cells)
    return "\n".join([*(p.text for p in document.paragraphs), *cells])


def test_legacy_paper_without_snapshot_keeps_its_employee_signature(env, monkeypatch):
    give_employee_signature(env)
    book, leave = generate_on_sept_30(env, monkeypatch, embed_signature={"employee": True})
    drop_snapshot(env, book)
    previous = book.versions[-1]
    old_doc = env.db.get_one(Document, previous.document_id)
    assert roles(env, old_doc.docx_path)["employee"] == 1
    edit(env, leave)
    current = book.versions[-1]
    assert current.id != previous.id
    # Render-only recovery: the record is not enrolled into the sign gate.
    assert "_signing" not in current.fields
    assert not book_service.is_enrolled(book)
    new_doc = env.db.get_one(Document, current.document_id)
    assert roles(env, new_doc.docx_path) == roles(env, old_doc.docx_path)
    assert roles(env, companion_of(env, new_doc).docx_path) == roles(
        env, companion_of(env, old_doc).docx_path
    )


def test_record_filed_by_a_submitter_syncs_with_the_same_signatures(env):
    # A picked submitter signs the employee cell; the form has no submitter
    # slot, so the submitter's signature appears only once, as ``employee``.
    give_employee_signature(env)
    submitter = Submitter(name=env.employee.name_en, employee_id=env.employee.id)
    env.db.add(submitter)
    env.db.commit()
    book, leave = generate(env, embed_signature={"employee": True}, submitter_id=submitter.id)
    previous = book.versions[-1]
    old_doc = env.db.get_one(Document, previous.document_id)
    assert previous.fields["_signing"]["submitter_id"] == submitter.id
    assert roles(env, old_doc.docx_path) == Counter({"employee": 1})
    edit(env, leave)
    current = book.versions[-1]
    assert current.version_no == previous.version_no + 1
    assert current.trigger == "leave-edit"
    new_doc = env.db.get_one(Document, current.document_id)
    assert roles(env, new_doc.docx_path) == roles(env, old_doc.docx_path)
    assert roles(env, companion_of(env, new_doc).docx_path) == roles(
        env, companion_of(env, old_doc).docx_path
    )
    assert "25/10/2026" in docx_text(env, new_doc.docx_path)
    actions = env.db.scalars(
        select(AuditLog.action).where(
            AuditLog.entity_type == "book", AuditLog.entity_id == str(book.id)
        )
    ).all()
    assert "book.leave_dates_synced" in actions
    assert "book.leave_dates_not_synced" not in actions


@pytest.mark.parametrize("legacy", [False, True])
def test_leave_edit_keeps_the_original_request_date(env, monkeypatch, legacy):
    book, leave = generate_on_sept_30(env, monkeypatch)
    if legacy:
        drop_snapshot(env, book)
    old_doc = env.db.get_one(Document, book.versions[-1].document_id)
    old_text = docx_text(env, old_doc.docx_path)
    assert "30/09/2026" in old_text
    edit(env, leave)
    new_doc = env.db.get_one(Document, book.versions[-1].document_id)
    new_text = docx_text(env, new_doc.docx_path)
    assert "25/10/2026" in new_text
    assert new_text.count("30/09/2026") == old_text.count("30/09/2026")
    today = datetime.now().strftime("%d/%m/%Y")
    if today not in old_text:
        assert today not in new_text
    companion_text = docx_text(env, companion_of(env, new_doc).docx_path)
    assert companion_text.count("30/09/2026") == docx_text(
        env, companion_of(env, old_doc).docx_path
    ).count("30/09/2026")


def test_legacy_signed_copy_is_reproduced_from_its_retained_source(env, monkeypatch):
    give_employee_signature(env)
    book, leave = generate_on_sept_30(env, monkeypatch, embed_signature={"employee": True})
    drop_snapshot(env, book)
    sign(env, book)
    previous = book.versions[-1]
    old_retained = env.db.scalar(
        select(SignatureArtifactRevision).where(SignatureArtifactRevision.version_id == previous.id)
    )
    edit(env, leave)
    current = book.versions[-1]
    assert current.id != previous.id
    assert current.signed_at == previous.signed_at
    assert book.approval_state == "approved"
    new_retained = env.db.scalar(
        select(SignatureArtifactRevision).where(SignatureArtifactRevision.version_id == current.id)
    )
    assert roles(env, new_retained.docx_path) == roles(env, old_retained.docx_path)
    assert "25/10/2026" in text(env, current.signed_pdf_path)


def test_signature_that_cannot_be_reproduced_skips_instead_of_unsigned_copy(env, monkeypatch):
    signature = give_employee_signature(env)
    book, leave = generate_on_sept_30(env, monkeypatch, embed_signature={"employee": True})
    drop_snapshot(env, book)
    previous = book.versions[-1]
    signature.unlink()
    files_before = {path for path in env.settings.data_dir.rglob("*") if path.is_file()}
    edit(env, leave)
    assert leave.days == 4
    assert book.versions == [previous]
    assert {path for path in env.settings.data_dir.rglob("*") if path.is_file()} == files_before
    audit = env.db.scalar(
        select(AuditLog).where(
            AuditLog.entity_id == str(book.id), AuditLog.action == "book.leave_dates_not_synced"
        )
    )
    assert "signature" in json.loads(audit.payload)["reason"].lower()


@pytest.mark.parametrize("damage", ["missing", "corrupt", "unreadable"])
def test_unreadable_record_paper_skips_sync_and_saves_the_leave(env, monkeypatch, damage):
    book, leave = generate(env)
    previous = book.versions[-1]
    original = env.settings.data_dir / env.db.get_one(Document, previous.document_id).docx_path
    if damage == "missing":
        original.unlink()
    elif damage == "corrupt":
        original.write_bytes(b"not a zip archive")
    else:
        real_open = Path.open

        def denied(self, *args, **kwargs):
            if self == original:
                raise PermissionError(13, "Permission denied", str(self))
            return real_open(self, *args, **kwargs)

        monkeypatch.setattr(Path, "open", denied)
        real_docx = leave_record_sync.DocxDocument

        def denied_docx(source, *args, **kwargs):
            if Path(str(source)) == original:
                raise PermissionError(13, "Permission denied", str(source))
            return real_docx(source, *args, **kwargs)

        monkeypatch.setattr(leave_record_sync, "DocxDocument", denied_docx)
        monkeypatch.setattr(
            signature_layout,
            "inspect_signature_drawings",
            lambda p: (
                (_ for _ in ()).throw(PermissionError(13, "Permission denied"))
                if Path(p) == original
                else []
            ),
        )
    converted = []
    monkeypatch.setattr(
        document_service, "convert_docx_to_pdf", lambda path: converted.append(path)
    )
    files_before = {path for path in env.settings.data_dir.rglob("*") if path.is_file()}
    edit(env, leave)
    env.db.commit()
    env.db.refresh(leave)
    assert (leave.end_date, leave.days) == (date(2026, 10, 25), 4)
    assert book.versions == [previous]
    assert converted == []  # skipped before any Word conversion
    assert {path for path in env.settings.data_dir.rglob("*") if path.is_file()} == files_before
    audit = env.db.scalar(
        select(AuditLog).where(
            AuditLog.entity_id == str(book.id), AuditLog.action == "book.leave_dates_not_synced"
        )
    )
    assert "cannot be read" in json.loads(audit.payload)["reason"]
