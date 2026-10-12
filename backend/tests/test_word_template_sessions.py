"""Word template authoring, draft capture and live metadata contract."""

from datetime import UTC, datetime
from pathlib import Path

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.api.deps import get_current_user
from app.config import Settings
from app.core.book_template_retokenize import _header_copies
from app.db import session as session_mod
from app.db.models import Base, Book, BookCategory, BookEditSession, User
from app.db.session import attach_sqlite_pragmas, get_db
from app.main import create_app
from app.services import artifact_service, book_template_service, perm_service, word_book_service


@pytest.fixture()
def authoring(monkeypatch, tmp_path):
    engine = create_engine(
        f"sqlite:///{tmp_path / 'routes.db'}", connect_args={"check_same_thread": False}
    )
    attach_sqlite_pragmas(engine, wal=False)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    monkeypatch.setattr(session_mod, "engine", engine)
    monkeypatch.setattr(session_mod, "SessionLocal", factory)
    settings = Settings(
        data_dir=tmp_path / "data",
        templates_dir=Path(__file__).resolve().parents[1] / "templates",
        public_base_url="http://word.test/",
    )
    for service in (word_book_service, book_template_service, artifact_service):
        monkeypatch.setattr(service, "get_settings", lambda: settings)
    with factory() as db:
        perm_service.seed_role_defaults(db)
        db.add(BookCategory(id="GS", prefix="GS"))
        user = User(email="author@test.ae", password_hash="x", role="admin", status="active")
        other = User(email="other@test.ae", password_hash="x", role="admin", status="active")
        db.add_all([user, other])
        db.commit()
        app = create_app()
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        with TestClient(app) as client:
            yield db, client, app, user, other
    engine.dispose()


def _save(db, session):
    session.last_put_at = datetime.now(UTC).replace(tzinfo=None)
    db.commit()


def _create_book(client, **extra):
    response = client.post(
        "/api/v1/books/word-sessions",
        json={"classification_code": "5/1", "subject": "Actual subject", **extra},
    )
    assert response.status_code == 201, response.text
    return response.json()["book_id"]


def test_template_session_round_trip(authoring):
    db, client, _app, _user, _other = authoring
    response = client.post("/api/v1/books/word-templates/sessions", json={"name": "Reusable"})
    assert response.status_code == 201, response.text
    info = response.json()
    assert info["name"] == "Reusable.docx"
    assert info["word_url"] == f"ms-word:ofe|u|http://word.test/dav/{info['token']}/Reusable.docx"
    assert db.query(Book).count() == 0
    session = db.get(BookEditSession, info["id"])
    working = Path(session.working_path)
    doc = Document(working)
    doc.add_paragraph("The reusable body survives Word authoring.")
    doc.save(working)
    _save(db, session)
    response = client.post(f"/api/v1/books/word-templates/sessions/{session.id}/finish")
    assert response.status_code == 201, response.text
    assert session.state == "finished"
    assert not working.parent.exists()
    assert "Reusable.docx" in [t.name for t in book_template_service.list_templates()]
    book_id = _create_book(client, template_name="Reusable.docx", cc=["New CC"])
    detail = client.get(f"/api/v1/books/{book_id}/word-sessions").json()
    assert detail["template_name"] == "Reusable.docx"
    assert detail["subject"] == "Actual subject"
    assert detail["cc"] == ["New CC"]
    new_session = db.query(BookEditSession).filter_by(book_id=book_id, state="active").one()
    rendered = Document(new_session.working_path)
    assert "The reusable body survives Word authoring." in "\n".join(
        p.text for p in rendered.paragraphs
    )
    header = "\n".join(p.text for copy in _header_copies(rendered) for p in copy)
    assert f"الرقم: {detail['ref_number']}" in header
    assert "1/1/GSSG/1" not in header
    assert "موضوع القالب" not in "\n".join(p.text for p in rendered.paragraphs)


def test_save_draft_requires_word_save_and_preserves_current_body(authoring):
    db, client, _app, _user, _other = authoring
    book_id = _create_book(client)
    url = f"/api/v1/books/{book_id}/save-as-template"
    response = client.post(url, json={"name": "Draft"})
    assert response.status_code == 409, response.text
    assert response.json()["error"]["code"] == "NO_SOURCE_DOCX"
    session = db.query(BookEditSession).filter_by(book_id=book_id, state="active").one()
    document = Document(session.working_path)
    document.add_paragraph("Current draft body captured before Finish.")
    document.save(session.working_path)
    _save(db, session)
    response = client.post(url, json={"name": "Draft"})
    assert response.status_code == 201, response.text
    assert session.state == "active"
    stored = Document(book_template_service.resolve_template_path("Draft"))
    assert "Current draft body captured before Finish." in "\n".join(
        p.text for p in stored.paragraphs
    )


def test_word_detail_reads_current_docx_and_missing_session(authoring):
    db, client, _app, _user, _other = authoring
    book_id = _create_book(client)
    session = db.query(BookEditSession).filter_by(book_id=book_id, state="active").one()
    session.template_name = "Started.docx"
    document = Document(session.working_path)
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("الموضوع:"):
            paragraph.text = "الموضوع: Edited in Word"
        if paragraph.text.startswith("السيد"):
            paragraph.text = "السيد \\ Current recipient المحترم"
    document.add_paragraph("• نسخة إلى: First\n• نسخة إلى: Second")
    document.add_paragraph("نسخة إلى: Third")
    document.save(session.working_path)
    _save(db, session)
    response = client.get(f"/api/v1/books/{book_id}/word-sessions")
    assert response.status_code == 200, response.text
    detail = response.json()
    assert detail["subject"] == "Edited in Word"
    assert detail["recipient_name"] == "Current recipient"
    assert detail["cc"] == ["First", "Second", "Third"]
    assert detail["template_name"] == "Started.docx"
    assert detail["created_at"] and detail["last_put_at"]
    session.state = "discarded"
    db.commit()
    response = client.get(f"/api/v1/books/{book_id}/word-sessions")
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "NO_ACTIVE_SESSION"


def test_template_session_owner_saves_collision_and_discard(authoring):
    db, client, app, user, other = authoring
    response = client.post("/api/v1/books/word-templates/sessions", json={"name": "Private"})
    assert response.status_code == 201, response.text
    session_id = response.json()["id"]
    url = f"/api/v1/books/word-templates/sessions/{session_id}"
    response = client.post(f"{url}/finish")
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "NO_SAVES_YET"
    assert client.get(url).status_code == 200
    app.dependency_overrides[get_current_user] = lambda: other
    for method, suffix in (("get", ""), ("get", "/preview"), ("post", "/finish"), ("delete", "")):
        response = getattr(client, method)(url + suffix)
        assert response.status_code == 404, response.text
        assert response.json()["error"]["code"] == "TEMPLATE_SESSION_NOT_FOUND"
    app.dependency_overrides[get_current_user] = lambda: user
    session = db.get(BookEditSession, session_id)
    working_dir = Path(session.working_path).parent
    assert client.delete(url).status_code == 204
    assert session.state == "discarded" and not working_dir.exists()
    assert client.get(url).status_code == 404
    book_template_service.resolve_template_path("Existing").write_bytes(b"existing")
    response = client.post("/api/v1/books/word-templates/sessions", json={"name": "Existing"})
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "TEMPLATE_EXISTS"


def test_template_preview_reuses_converter_and_requires_capability(authoring, monkeypatch):
    db, client, _app, user, _other = authoring
    response = client.post("/api/v1/books/word-templates/sessions", json={"name": "Preview"})
    assert response.status_code == 201
    session = db.get(BookEditSession, response.json()["id"])
    url = f"/api/v1/books/word-templates/sessions/{session.id}/preview"
    assert client.get(url).status_code == 409
    _save(db, session)

    def convert(source):
        output = source.with_suffix(".pdf")
        with fitz.open() as pdf:
            pdf.new_page()
            pdf.save(output)
        return output

    monkeypatch.setattr(word_book_service, "convert_docx_to_pdf", convert)
    response = client.get(url)
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert response.headers["cache-control"] == "no-store"
    assert client.get(url + "?encoding=base64").status_code == 200
    user.role = "operator"
    db.commit()
    assert client.get(url).status_code == 403
