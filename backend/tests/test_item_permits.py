"""Item permits (إدخال مواد): service + API behavior and the letter body builder."""

from __future__ import annotations

from collections.abc import Iterator
from typing import Any

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.api.errors import NotFoundError, ValidationFailedError
from app.core.arabic_rtl import html_to_docx
from app.core.permit_letter import _access_chips, build_item_permit_letter_html
from app.db.models import (
    AuditLog,
    Book,
    BookApprovalStep,
    BookCategory,
    BookRecipient,
    BookVersion,
    Employee,
    ItemPermit,
    Manager,
    User,
    UserPermission,
)
from app.db.models import (
    Document as GeneratedDocument,
)
from app.db.session import get_db
from app.main import create_app
from app.schemas.item_permit import ItemPermitCreate, ItemPermitUpdate
from app.services import document_service, item_permit_service


@pytest.fixture()
def env(api_db: Session, tmp_path: Any, monkeypatch: pytest.MonkeyPatch) -> Iterator[Session]:
    """Tmp data dir + stubbed PDF chain, with an employee and an operator."""
    from app.config import Settings, get_settings

    monkeypatch.setenv("GSSG_DATA_DIR", str(tmp_path / "data"))
    get_settings.cache_clear()
    settings = Settings(data_dir=tmp_path / "data")
    monkeypatch.setattr(document_service, "get_settings", lambda: settings)
    monkeypatch.setattr(document_service, "convert_docx_to_pdf", lambda p: None)
    api_db.add(BookCategory(id="GS", prefix="GS"))
    api_db.add(Employee(id="G3082", name_en="Ali", name_ar="علي", position_ar="مهندس"))
    api_db.add(User(email="op@x.ae", password_hash="x", role="admin", status="active"))
    api_db.commit()
    yield api_db
    get_settings.cache_clear()


def _linked_manager(db: Session) -> Manager:
    u = User(email="mgr@x.ae", password_hash="x", role="admin", status="active")
    db.add(u)
    db.commit()
    m = Manager(name_en="Boss", user_id=u.id)
    db.add(m)
    db.commit()
    return m


def _payload(**kw: Any) -> ItemPermitCreate:
    base: dict[str, Any] = dict(
        employee_id="G3082",
        items=[{"name": "كابل شبكة", "quantity": 3}, {"name": "<b>مفك</b>", "quantity": 1}],
        send_for_approval=False,
    )
    base.update(kw)
    return ItemPermitCreate(**base)


def _latest_version(db: Session, book_id: int):
    book = db.get_one(Book, book_id)
    return book, max(book.versions, key=lambda v: v.version_no)


def test_create_generates_1_5_letter_on_security_permit_paper(env: Session) -> None:
    row = item_permit_service.create_item_permit(
        env, _payload(), user=env.query(User).filter_by(email="op@x.ae").one()
    )
    book, latest = _latest_version(env, row.book_id)
    assert book.ref_number.startswith("1/5/")
    assert latest.template_id == "Security Permit"
    assert latest.fields["subject"] == "تصريح إدخال مواد"
    assert latest.fields["recipient_name"] == "مسؤول وحدة التفتيش"
    body = latest.fields["body"]
    assert "المنطقة الحمراء" in body and "مبنى مركز الإصلاح والتأهيل الوثبة - 2" in body
    assert "كابل شبكة" in body and "G3082" in body and "مهندس" in body and "علي" in body
    read = item_permit_service.to_read(row, db=env)
    assert read.book_ref == book.ref_number
    assert read.employee_name == "علي" and read.employee_title == "مهندس"


def test_update_regenerates_letter(env: Session) -> None:
    row = item_permit_service.create_item_permit(
        env, _payload(), user=env.query(User).filter_by(email="op@x.ae").one()
    )
    ref = env.get_one(Book, row.book_id).ref_number
    row = item_permit_service.update_item_permit(
        env,
        row.id,
        ItemPermitUpdate(zones=["green"], items=[{"name": "طابعة", "quantity": 2}]),
        user=env.query(User).filter_by(email="op@x.ae").one(),
    )
    book, latest = _latest_version(env, row.book_id)
    assert book.ref_number == ref and len(book.versions) == 1  # draft re-rendered in place
    assert "المنطقة الخضراء" in latest.fields["body"] and "طابعة" in latest.fields["body"]
    assert "كابل شبكة" not in latest.fields["body"]
    # A signed letter is never edited in place: the edit appends a new version.
    book.approval_state = "approved"
    env.commit()
    item_permit_service.update_item_permit(
        env,
        row.id,
        ItemPermitUpdate(site="مبنى آخر"),
        user=env.query(User).filter_by(email="op@x.ae").one(),
    )
    book, latest = _latest_version(env, row.book_id)
    assert len(book.versions) == 2 and "مبنى آخر" in latest.fields["body"]


def test_pending_letter_resubmits_after_edit(env: Session) -> None:
    mgr = _linked_manager(env)
    row = item_permit_service.create_item_permit(
        env,
        _payload(manager_id=mgr.id, send_for_approval=True),
        user=env.query(User).filter_by(email="op@x.ae").one(),
    )
    assert item_permit_service.to_read(row, db=env).approval_state == "pending"
    row = item_permit_service.update_item_permit(
        env,
        row.id,
        ItemPermitUpdate(site="مبنى آخر"),
        user=env.query(User).filter_by(email="op@x.ae").one(),
    )
    book, latest = _latest_version(env, row.book_id)
    assert book.approval_state == "pending"
    assert [s.state for s in latest.approval_steps] == ["pending"]


def test_manual_submit_surfaces_chain_errors(env: Session) -> None:
    m = Manager(name_en="No Login")
    env.add(m)
    env.commit()
    row = item_permit_service.create_item_permit(
        env, _payload(manager_id=m.id), user=env.query(User).filter_by(email="op@x.ae").one()
    )
    with pytest.raises(ValidationFailedError):
        item_permit_service.submit_item_permit_book(
            env, row.id, user=env.query(User).filter_by(email="op@x.ae").one()
        )


@pytest.mark.parametrize(
    "bad",
    [
        {"items": []},
        {"items": [{"name": "x", "quantity": 0}]},
        {"zones": ["blue"]},
        {"zones": []},
    ],
)
def test_schema_rejects_invalid_payload(bad: dict[str, Any]) -> None:
    with pytest.raises(ValidationError):
        _payload(**bad)


def test_unknown_employee_rejected(env: Session) -> None:
    with pytest.raises(ValidationFailedError) as exc:
        item_permit_service.create_item_permit(
            env,
            _payload(employee_id="G0000"),
            user=env.query(User).filter_by(email="op@x.ae").one(),
        )
    assert exc.value.code == "ITEM_PERMIT_EMPLOYEE_NOT_FOUND"


def test_soft_delete_hides_from_list_and_get(env: Session) -> None:
    row = item_permit_service.create_item_permit(
        env, _payload(), user=env.query(User).filter_by(email="op@x.ae").one()
    )
    item_permit_service.soft_delete_item_permit(
        env, row.id, user=env.query(User).filter_by(email="op@x.ae").one()
    )
    assert item_permit_service.list_item_permits(env) == ([], 0)
    with pytest.raises(NotFoundError):
        item_permit_service.get_item_permit(env, row.id)


def _client(db: Session, user: User) -> TestClient:
    app = create_app()
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: user
    return TestClient(app, raise_server_exceptions=True)


def test_api_create_list_search_get_delete(env: Session) -> None:
    c = _client(env, env.query(User).filter_by(email="op@x.ae").one())
    r = c.post(
        "/api/v1/item-permits",
        json={"employee_id": "G3082", "items": [{"name": "كابل شبكة", "quantity": 3}]},
    )
    assert r.status_code == 201, r.text
    created = r.json()
    assert created["zones"] == ["red"] and created["employee_name"] == "علي"
    assert c.get("/api/v1/item-permits", params={"q": "كابل"}).json()["total"] == 1
    assert c.get("/api/v1/item-permits", params={"q": "nomatch"}).json()["total"] == 0
    assert c.get(f"/api/v1/item-permits/{created['id']}").status_code == 200
    assert c.delete(f"/api/v1/item-permits/{created['id']}").status_code == 204
    assert c.get(f"/api/v1/item-permits/{created['id']}").status_code == 404


def test_api_requires_permits_capabilities(env: Session) -> None:
    viewer = User(email="v@x.ae", password_hash="x", role="operator", status="active")
    denied = User(email="d@x.ae", password_hash="x", role="manager", status="active")
    env.add_all([viewer, denied])
    env.commit()
    env.add(UserPermission(user_id=denied.id, capability="permits.view", effect="deny"))
    env.commit()
    r = _client(env, denied).get("/api/v1/item-permits")
    assert r.status_code == 403 and r.json()["error"]["details"]["capability"] == "permits.view"
    c = _client(env, viewer)
    assert c.get("/api/v1/item-permits").status_code == 200
    r = c.post("/api/v1/item-permits", json={})
    assert r.status_code == 403 and r.json()["error"]["details"]["capability"] == "permits.create"


@pytest.mark.parametrize("level", ["personal", "directed", "secret"])
@pytest.mark.parametrize("operation", ["update", "submit", "delete"])
def test_item_permit_mutations_require_linked_letter_write_access(
    env: Session, level: str, operation: str
) -> None:
    creator = env.query(User).filter_by(email="op@x.ae").one()
    editor = User(email="editor@x.ae", password_hash="x", role="manager", status="active")
    env.add(editor)
    env.commit()
    manager = _linked_manager(env)
    row = item_permit_service.create_item_permit(
        env, _payload(manager_id=manager.id, send_for_approval=operation == "update"), user=creator
    )
    book = env.get_one(Book, row.book_id)
    book.confidentiality = level
    if level == "directed":
        env.add(
            BookRecipient(
                book_id=book.id,
                user_id=editor.id,
                sent_by_user_id=creator.id,
                purpose="info",
            )
        )
    env.commit()

    def snapshot():
        env.expire_all()
        return {
            model.__tablename__: [
                tuple(record)
                for record in env.execute(select(*model.__table__.columns).order_by(model.id))
            ]
            for model in (
                ItemPermit,
                Book,
                BookVersion,
                BookApprovalStep,
                GeneratedDocument,
                AuditLog,
            )
        }

    before = snapshot()
    output_dir = document_service.get_settings().data_dir / "output"
    files_before = {p: p.read_bytes() for p in output_dir.rglob("*") if p.is_file()}
    path = f"/api/v1/item-permits/{row.id}"

    def request(user):
        client = _client(env, user)
        if operation == "update":
            return client.patch(path, json={"site": "Changed site"})
        if operation == "submit":
            return client.post(f"{path}/submit-approval")
        return client.delete(path)

    response = request(editor)
    assert response.status_code == 404, response.text
    assert snapshot() == before
    assert {p: p.read_bytes() for p in output_dir.rglob("*") if p.is_file()} == files_before
    response = request(creator)
    assert response.status_code == (204 if operation == "delete" else 200), response.text
    env.expire_all()
    if operation == "update":
        assert row.site == "Changed site"
        assert "Changed site" in book.versions[-1].fields["body"]
        assert book.approval_state == "pending"
    elif operation == "submit":
        assert book.approval_state == "pending"
        assert book.submitted_by_user_id == creator.id
        assert book.versions[-1].approval_steps
    else:
        assert row.deleted_at is not None


def _permit_operator(db: Session) -> User:
    """An operator granted the permit register but not generic record editing."""
    user = User(email="permits@x.ae", password_hash="x", role="operator", status="active")
    db.add(user)
    db.commit()
    for capability in ("permits.create", "permits.edit", "permits.delete"):
        db.add(UserPermission(user_id=user.id, capability=capability, effect="grant"))
    db.commit()
    return user


@pytest.mark.parametrize("operation", ["update", "regenerate", "submit", "delete"])
def test_permit_grants_carry_linked_letter_edits_without_books_edit(
    env: Session, operation: str
) -> None:
    """Permit grants (split from books.* by migration 0078) delegate edits to the
    linked letter; its level and write policy still apply. The outcome is
    all-or-nothing: the permit row and the letter body never disagree."""
    creator = env.query(User).filter_by(email="op@x.ae").one()
    operator = _permit_operator(env)
    manager = _linked_manager(env)
    row = item_permit_service.create_item_permit(
        env,
        _payload(manager_id=manager.id, send_for_approval=operation == "regenerate"),
        user=creator,
    )
    book = env.get_one(Book, row.book_id)
    client = _client(env, operator)
    path = f"/api/v1/item-permits/{row.id}"
    if operation in ("update", "regenerate"):
        response = client.patch(path, json={"site": "Changed site"})
    elif operation == "submit":
        response = client.post(f"{path}/submit-approval")
    else:
        response = client.delete(path)

    assert response.status_code == (204 if operation == "delete" else 200), response.text
    env.expire_all()
    if operation in ("update", "regenerate"):
        assert row.site == "Changed site"
        assert "Changed site" in book.versions[-1].fields["body"]
        assert book.approval_state == ("pending" if operation == "regenerate" else "none")
    elif operation == "submit":
        assert book.approval_state == "pending"
        assert book.submitted_by_user_id == operator.id
    else:
        assert row.deleted_at is not None


@pytest.mark.parametrize("operation", ["update", "delete"])
def test_permit_with_deleted_letter_stays_editable(env: Session, operation: str) -> None:
    """A draft letter deleted from Records no longer gates its permit; the
    permit capability still does. The deleted letter is left untouched."""
    creator = env.query(User).filter_by(email="op@x.ae").one()
    row = item_permit_service.create_item_permit(env, _payload(), user=creator)
    book = env.get_one(Book, row.book_id)
    assert _client(env, creator).delete(f"/api/v1/books/{book.id}").status_code == 204
    letter_before = [(v.id, v.fields) for v in book.versions]
    path = f"/api/v1/item-permits/{row.id}"
    viewer = User(email="viewer@x.ae", password_hash="x", role="operator", status="active")
    env.add(viewer)
    env.commit()

    def request(user: User):
        client = _client(env, user)
        if operation == "update":
            return client.patch(path, json={"site": "Changed site"})
        return client.delete(path)

    denied = request(viewer)
    assert denied.status_code == 403, denied.text
    capability = "permits.edit" if operation == "update" else "permits.delete"
    assert denied.json()["error"]["details"]["capability"] == capability
    response = request(_permit_operator(env))
    assert response.status_code == (200 if operation == "update" else 204), response.text
    env.expire_all()
    if operation == "update":
        assert row.site == "Changed site"
    else:
        assert row.deleted_at is not None
    assert book.deleted_at is not None
    assert [(v.id, v.fields) for v in book.versions] == letter_before


# ─── letter body builder (pure) ────────────────────────────────────────────────

_EMP = {"id": "G1", "title": "فني", "name": "سالم"}


def _td(*cells: str) -> str:
    return "".join(f'<td style="text-align:center">{c}</td>' for c in cells)


def test_letter_zone_mapping_numbering_and_no_approval_line() -> None:
    items: list[dict[str, object]] = [
        {"name": "أ", "quantity": 2},
        {"name": "ب", "quantity": 5},
    ]
    red = build_item_permit_letter_html(zones=["red"], site="الموقع", items=items, employee=_EMP)
    green = build_item_permit_letter_html(
        zones=["green"], site="الموقع", items=items, employee=_EMP
    )
    assert _access_chips(["red"]) + " في الموقع" in red and "المنطقة الخضراء" not in red
    assert _access_chips(["green"]) + " في الموقع" in green and "المنطقة الحمراء" not in green
    assert _td("1", "أ", "2") in red
    assert _td("2", "ب", "5") in red
    assert '<th style="text-align:center">م</th>' in red
    assert _td("G1", "فني", "سالم") in red
    assert "الإعتماد" not in red


def test_letter_escapes_user_text() -> None:
    html = build_item_permit_letter_html(
        zones=["red"],
        site="<i>s</i>",
        items=[{"name": "<script>x</script>", "quantity": 1}],
        employee={"id": "G1", "title": "<b>t</b>", "name": "a&b"},
    )
    assert "<script>" not in html and "<i>s" not in html and "<b>t" not in html
    assert "&lt;script&gt;" in html and "a&amp;b" in html


def test_zones_default_normalize_and_letter_chips(env: Session) -> None:
    assert _payload().zones == ["red"]
    assert _payload(zones=["work_residence", "red", "red"]).zones == ["red", "work_residence"]
    row = item_permit_service.create_item_permit(
        env,
        _payload(zones=["work_residence", "green"]),
        user=env.query(User).filter_by(email="op@x.ae").one(),
    )
    assert row.zones == ["green", "work_residence"]
    _, latest = _latest_version(env, row.book_id)
    body = latest.fields["body"]
    assert f"{_access_chips(['green'])} و {_access_chips(['work_residence'])} في" in body
    assert "المنطقة الحمراء" not in body


def test_update_zones_regenerates_letter_and_null_is_unchanged(env: Session) -> None:
    row = item_permit_service.create_item_permit(
        env, _payload(), user=env.query(User).filter_by(email="op@x.ae").one()
    )
    row = item_permit_service.update_item_permit(
        env,
        row.id,
        ItemPermitUpdate(zones=["green", "red"]),
        user=env.query(User).filter_by(email="op@x.ae").one(),
    )
    assert row.zones == ["red", "green"]
    _, latest = _latest_version(env, row.book_id)
    assert " و ".join(_access_chips([z]) for z in ("red", "green")) in latest.fields["body"]
    row = item_permit_service.update_item_permit(
        env,
        row.id,
        ItemPermitUpdate(zones=None),
        user=env.query(User).filter_by(email="op@x.ae").one(),
    )
    assert row.zones == ["red", "green"]


def test_api_zones_validation(env: Session) -> None:
    c = _client(env, env.query(User).filter_by(email="op@x.ae").one())
    items = [{"name": "كابل", "quantity": 1}]
    for bad in ([], ["blue"]):
        r = c.post(
            "/api/v1/item-permits", json={"employee_id": "G3082", "items": items, "zones": bad}
        )
        assert r.status_code == 422, r.text
    r = c.post(
        "/api/v1/item-permits",
        json={"employee_id": "G3082", "items": items, "zones": ["work_residence", "red", "red"]},
    )
    assert r.status_code == 201, r.text
    assert r.json()["zones"] == ["red", "work_residence"]
    r = c.patch(f"/api/v1/item-permits/{r.json()['id']}", json={"zones": []})
    assert r.status_code == 422


def test_letter_spacers_and_shared_chip_markup() -> None:
    html = build_item_permit_letter_html(
        zones=["red", "green"], site="s", items=[{"name": "a", "quantity": 1}], employee=_EMP
    )
    assert html.count("<p></p>") == 2
    sentence, rest = html.split("</p>", 1)
    assert rest.startswith("<p></p><table")
    assert _access_chips(["red"]) in sentence and _access_chips(["green"]) in sentence


def test_letter_tables_and_cells_centered_in_docx() -> None:
    doc = Document()
    html_to_docx(
        build_item_permit_letter_html(
            zones=["red"], site="s", items=[{"name": "a", "quantity": 1}], employee=_EMP
        ),
        doc.add_paragraph(),
    )
    assert len(doc.tables) == 2
    for table in doc.tables:
        assert table._tbl.tblPr.find(qn("w:jc")).get(qn("w:val")) == "center"
        paragraphs = [p for row in table.rows for cell in row.cells for p in cell.paragraphs]
        assert paragraphs and all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in paragraphs)
