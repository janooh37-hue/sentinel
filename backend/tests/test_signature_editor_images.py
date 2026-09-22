"""API-level tests for the signature-editor image transport (approval-
signature-placement plan §3, signature-placement-repair plan §3): both
``/signature-editor/images/{id}`` and
``/signature-editor/candidates/{id}/image`` must serve raw PNG bytes
normally, and ``?encoding=base64`` must decode back to byte-identical
content — the contract ``PlacementCanvas`` relies on (via
``frontend/src/lib/api.ts``'s ``fetchPermitBlob``) to show real ink before
Save. Fixtures follow ``test_document_download_companion.py``'s local
``api_db`` (data-dir-aware) and dependency-override ``TestClient``
convention.
"""

from __future__ import annotations

import base64
import io
from pathlib import Path

import pytest
from docx import Document as DocxFile
from fastapi.testclient import TestClient
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.api.deps import get_current_user
from app.core.constants import DEFAULT_MANAGER_NAME
from app.core.docx_engine import stamp_signature_above_name
from app.core.signature_layout import inspect_signature_drawings
from app.db import session as session_mod
from app.db.models import (
    Base,
    Book,
    BookCategory,
    BookVersion,
    Document,
    SignatureArtifactRevision,
    User,
)
from app.db.session import attach_sqlite_pragmas, get_db
from app.main import create_app
from app.services import perm_service
from app.services import signature_placement_service as sps


@pytest.fixture()
def api_db(monkeypatch, tmp_path) -> Session:
    monkeypatch.setenv("GSSG_DATA_DIR", str(tmp_path))
    from app.config import get_settings

    get_settings.cache_clear()
    eng = create_engine(
        f"sqlite:///{tmp_path / 't.db'}",
        future=True,
        connect_args={"check_same_thread": False},
    )
    attach_sqlite_pragmas(eng, wal=False)
    Base.metadata.create_all(eng)
    TS = sessionmaker(bind=eng, autoflush=False, expire_on_commit=False, future=True)
    monkeypatch.setattr(session_mod, "engine", eng)
    monkeypatch.setattr(session_mod, "SessionLocal", TS)
    db = TS()
    perm_service.seed_role_defaults(db)
    try:
        yield db
    finally:
        db.close()
        eng.dispose()
        get_settings.cache_clear()


def _user(db: Session, role: str = "admin", email: str = "admin@x.ae") -> User:
    u = User(email=email, password_hash="x", role=role, status="active")
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


def _client(db: Session, user: User) -> TestClient:
    app = create_app()
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: user
    return TestClient(app, raise_server_exceptions=True)


def _sig_png(path: Path, color: tuple[int, int, int, int] = (0, 0, 200, 255)) -> Path:
    Image.new("RGBA", (60, 30), color).save(path)
    return path


def _tracked_record(db: Session, tmp_path: Path, *, revision: int = 1) -> tuple[Document, str]:
    """A book/version/document with an active ``SignatureArtifactRevision``
    pointing at a real tracked docx carrying one marked manager signature.
    Returns ``(document, signature_id)``."""
    if db.get(BookCategory, "GS") is None:
        db.add(BookCategory(id="GS", prefix="GS"))
        db.flush()
    book = Book(category_id="GS", ref_number="1/1/1", subject="test")
    db.add(book)
    db.flush()

    docx_path = tmp_path / "source.docx"
    DocxFile().save(str(docx_path))
    document = Document(
        template_id="General Book",
        ref_number=book.ref_number,
        docx_path=docx_path.relative_to(tmp_path).as_posix(),
        submission_id="t-img",
        role="primary",
    )
    db.add(document)
    db.flush()
    version = BookVersion(
        book_id=book.id,
        version_no=1,
        template_id="General Book",
        fields={},
        document_id=document.id,
        signature_revision=revision,
    )
    db.add(version)
    db.flush()

    tracked_docx = tmp_path / "tracked.docx"
    d2 = DocxFile()
    d2.add_paragraph("Body")
    d2.add_paragraph(DEFAULT_MANAGER_NAME)
    d2.save(str(tracked_docx))
    sig = _sig_png(tmp_path / "sig.png")
    assert stamp_signature_above_name(
        tracked_docx, str(sig), [DEFAULT_MANAGER_NAME], size_mm=30.0, boldness=1
    )
    signature_id = inspect_signature_drawings(tracked_docx)[0].signature_id

    row = SignatureArtifactRevision(
        version_id=version.id,
        revision=revision,
        source_kind=sps.SOURCE_KIND_APPROVAL,
        signer_user_id=None,
        docx_path=tracked_docx.relative_to(tmp_path).as_posix(),
        primary_pdf_path=None,
        published_pdf_path=None,
        previous_published_pdf_path=None,
        docx_sha256=sps._sha256_file(tracked_docx),
        manifest=[],
        action=sps.ACTION_INITIAL,
        signature_id=signature_id,
        before_geometry=None,
        after_geometry=None,
        actor_user_id=None,
    )
    db.add(row)
    db.commit()
    return document, signature_id


def _legacy_record(db: Session, tmp_path: Path) -> tuple[Document, str]:
    """A document whose own docx is the ONE verified legacy source, carrying
    a single unmarked image drawing. Returns ``(document, docpr_name)``."""
    if db.get(BookCategory, "GS") is None:
        db.add(BookCategory(id="GS", prefix="GS"))
        db.flush()
    book = Book(category_id="GS", ref_number="2/2/2", subject="legacy test")
    db.add(book)
    db.flush()

    docx_path = tmp_path / "legacy.docx"
    d = DocxFile()
    d.add_paragraph("Body")
    sig = _sig_png(tmp_path / "legacy_sig.png", color=(200, 0, 0, 255))
    d.add_paragraph().add_run().add_picture(str(sig), width=None)
    d.save(str(docx_path))

    document = Document(
        template_id="General Book",
        ref_number=book.ref_number,
        docx_path=docx_path.relative_to(tmp_path).as_posix(),
        submission_id="t-legacy",
        role="primary",
    )
    db.add(document)
    db.flush()
    version = BookVersion(
        book_id=book.id,
        version_no=1,
        template_id="General Book",
        fields={},
        document_id=document.id,
        signature_revision=0,
    )
    db.add(version)
    db.commit()

    from app.services.signature_placement_service import discover_legacy_candidates

    _, found = discover_legacy_candidates(document)
    assert len(found) == 1
    return document, found[0].candidate_id


# ---------------------------------------------------------------------------
# /signature-editor/images/{signature_id}
# ---------------------------------------------------------------------------


def test_signature_image_raw_and_base64_decode_identically(api_db, tmp_path) -> None:
    document, signature_id = _tracked_record(api_db, tmp_path)
    admin = _user(api_db)
    client = _client(api_db, admin)

    raw = client.get(
        f"/api/v1/documents/{document.id}/signature-editor/images/{signature_id}",
        params={"signature_revision": 1},
    )
    assert raw.status_code == 200
    assert raw.headers["content-type"].startswith("image/png")
    Image.open(io.BytesIO(raw.content)).verify()  # Pillow accepts the raw bytes

    encoded = client.get(
        f"/api/v1/documents/{document.id}/signature-editor/images/{signature_id}",
        params={"signature_revision": 1, "encoding": "base64"},
    )
    assert encoded.status_code == 200
    assert encoded.headers["content-type"].startswith("text/plain")
    assert encoded.headers["x-content-type-options"] == "nosniff"
    assert encoded.headers["cache-control"] == "no-store"
    decoded = base64.b64decode(encoded.text)
    assert decoded == raw.content  # byte-identical to the raw response
    Image.open(io.BytesIO(decoded)).verify()  # and still a valid image


def test_signature_image_stale_revision_conflicts(api_db, tmp_path) -> None:
    document, signature_id = _tracked_record(api_db, tmp_path)
    admin = _user(api_db)
    client = _client(api_db, admin)

    resp = client.get(
        f"/api/v1/documents/{document.id}/signature-editor/images/{signature_id}",
        params={"signature_revision": 2},  # active revision is 1
    )
    assert resp.status_code == 409
    assert resp.json()["error"]["code"] == "SIGNATURE_REVISION_CONFLICT"


def test_signature_image_unauthorized_user_gets_no_bytes(api_db, tmp_path) -> None:
    document, signature_id = _tracked_record(api_db, tmp_path)
    outsider = _user(api_db, role="operator", email="outsider@x.ae")
    client = _client(api_db, outsider)

    resp = client.get(
        f"/api/v1/documents/{document.id}/signature-editor/images/{signature_id}",
        params={"signature_revision": 1},
    )
    assert resp.status_code == 403


# ---------------------------------------------------------------------------
# /signature-editor/candidates/{candidate_id}/image (legacy, admin-only)
# ---------------------------------------------------------------------------


def test_candidate_image_raw_and_base64_decode_identically(api_db, tmp_path) -> None:
    document, candidate_id = _legacy_record(api_db, tmp_path)
    admin = _user(api_db)
    client = _client(api_db, admin)

    raw = client.get(
        f"/api/v1/documents/{document.id}/signature-editor/candidates/{candidate_id}/image"
    )
    assert raw.status_code == 200
    assert raw.headers["content-type"].startswith("image/png")
    Image.open(io.BytesIO(raw.content)).verify()

    encoded = client.get(
        f"/api/v1/documents/{document.id}/signature-editor/candidates/{candidate_id}/image",
        params={"encoding": "base64"},
    )
    assert encoded.status_code == 200
    assert encoded.headers["content-type"].startswith("text/plain")
    assert encoded.headers["x-content-type-options"] == "nosniff"
    decoded = base64.b64decode(encoded.text)
    assert decoded == raw.content
    Image.open(io.BytesIO(decoded)).verify()


def test_candidate_image_non_admin_is_forbidden(api_db, tmp_path) -> None:
    document, candidate_id = _legacy_record(api_db, tmp_path)
    non_admin = _user(api_db, role="manager", email="mgr@x.ae")
    client = _client(api_db, non_admin)

    resp = client.get(
        f"/api/v1/documents/{document.id}/signature-editor/candidates/{candidate_id}/image"
    )
    assert resp.status_code == 403
