"""Printed grade marks in every header part of a record DOCX."""

import shutil
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.core.constants import TEMPLATE_FILES
from app.core.level_mark import LEVEL_MARK_TAG, apply_level_mark

TEMPLATES_DIR = Path(__file__).resolve().parents[1] / "templates"


def _header_elements(path: Path) -> list:
    doc = Document(str(path))
    parts = {}
    for section in doc.sections:
        for ref in section._sectPr.headerReference_lst:
            part = doc.part.rels[ref.get(qn("r:id"))].target_part
            parts[id(part)] = part.element
    return list(parts.values())


def _marks(hdr) -> list:
    return hdr.xpath(f'.//w:sdt[w:sdtPr/w:tag/@w:val="{LEVEL_MARK_TAG}"]')


def _text(mark) -> str:
    return "".join(t.text for t in mark.iter(qn("w:t")))


@pytest.fixture
def general_book(tmp_path: Path) -> Path:
    path = tmp_path / "book.docx"
    shutil.copy(TEMPLATES_DIR / TEMPLATE_FILES["General Book"], path)
    return path


@pytest.mark.parametrize(
    ("grade", "text"),
    [
        ("secret", "سري"),
        ("top_secret", "سري للغاية"),
        ("addressee_only", "سري للغاية — لعنايته فقط"),
    ],
)
def test_mark_in_every_header_part(general_book: Path, grade: str, text: str) -> None:
    assert apply_level_mark(general_book, grade)
    headers = _header_elements(general_book)
    assert len(headers) == 3  # even, default, first
    for hdr in headers:
        (mark,) = _marks(hdr)
        assert _text(mark) == text
        assert hdr[0] is mark  # first block: positioned frame at the top


def test_run_properties_exact(general_book: Path) -> None:
    assert apply_level_mark(general_book, "secret")
    (mark,) = _marks(_header_elements(general_book)[0])
    (rpr,) = mark.iter(qn("w:rPr"))
    fonts = rpr.find(qn("w:rFonts"))
    assert {fonts.get(qn(f"w:{k}")) for k in ("ascii", "hAnsi", "cs")} == {"Calibri"}
    assert rpr.find(qn("w:sz")).get(qn("w:val")) == "18"
    assert rpr.find(qn("w:szCs")).get(qn("w:val")) == "18"
    assert rpr.find(qn("w:u")).get(qn("w:val")) == "single"
    assert rpr.find(qn("w:color")).get(qn("w:val")) == "000000"
    assert rpr.find(qn("w:rtl")) is not None
    assert rpr.find(qn("w:cs")) is not None


def test_repeat_and_switch_keep_one_mark_then_remove(general_book: Path) -> None:
    assert apply_level_mark(general_book, "secret")
    assert apply_level_mark(general_book, "secret")
    for grade, text in (
        ("top_secret", "سري للغاية"),
        ("addressee_only", "سري للغاية — لعنايته فقط"),
    ):
        assert apply_level_mark(general_book, grade)
        for hdr in _header_elements(general_book):
            (mark,) = _marks(hdr)
            assert _text(mark) == text
    assert apply_level_mark(general_book, "unclassified")
    assert all(not _marks(hdr) for hdr in _header_elements(general_book))


def test_document_without_header_gets_one(tmp_path: Path) -> None:
    path = tmp_path / "bare.docx"
    Document().save(str(path))
    assert apply_level_mark(path, "secret")
    (hdr,) = _header_elements(path)
    assert len(_marks(hdr)) == 1
    assert apply_level_mark(path, "unclassified")
    assert not _marks(_header_elements(path)[0])


def test_later_section_first_page_inherits_mark_without_moving_letterhead(tmp_path: Path) -> None:
    path = tmp_path / "sections.docx"
    doc = Document()
    first = doc.sections[0]
    first.header.paragraphs[0].text = "Existing letterhead"
    first.footer.paragraphs[0].text = "Gray submitter footer"
    distance = first.header_distance
    later = doc.add_section()
    later.different_first_page_header_footer = True
    doc.save(str(path))

    assert apply_level_mark(path, "secret")
    stamped = Document(str(path))
    assert stamped.sections[0].header_distance == distance
    assert stamped.sections[0].header.paragraphs[0].text == "Existing letterhead"
    assert stamped.sections[0].footer.paragraphs[0].text == "Gray submitter footer"
    for section in stamped.sections:
        assert len(_marks(section.header._element)) == 1
    assert len(_marks(stamped.sections[1].first_page_header._element)) == 1


def test_reapply_restores_mark_format_after_word_edit(general_book: Path) -> None:
    assert apply_level_mark(general_book, "secret")
    doc = Document(str(general_book))
    (mark,) = _marks(doc.sections[0].header._element)
    next(mark.iter(qn("w:sz"))).set(qn("w:val"), "40")
    next(mark.iter(qn("w:color"))).set(qn("w:val"), "FF0000")
    doc.save(str(general_book))

    assert apply_level_mark(general_book, "secret")
    (mark,) = _marks(_header_elements(general_book)[0])
    assert next(mark.iter(qn("w:sz"))).get(qn("w:val")) == "18"
    assert next(mark.iter(qn("w:color"))).get(qn("w:val")) == "000000"


@pytest.mark.parametrize(
    "template_id",
    [
        "General Book",
        "Report",
        "Inmate Conduct Violations",
        "Leave Application Form",
        "Leave Permit Form",
        "Security Permit",
    ],
)
def test_real_templates_keep_letterhead_and_footer(tmp_path: Path, template_id: str) -> None:
    path = tmp_path / "form.docx"
    shutil.copy(TEMPLATES_DIR / TEMPLATE_FILES[template_id], path)
    before = Document(str(path))
    original_headers = {
        part.partname: part.element.xml
        for part in before.part.related_parts.values()
        if part.partname.startswith("/word/header")
    }
    original_footers = {
        part.partname: part.element.xml
        for part in before.part.related_parts.values()
        if part.partname.startswith("/word/footer")
    }
    assert apply_level_mark(path, "secret")
    stamped = Document(str(path))
    for part in stamped.part.related_parts.values():
        if part.partname in original_headers:
            for mark in _marks(part.element):
                part.element.remove(mark)
            assert part.element.xml == original_headers[part.partname]
        elif part.partname in original_footers:
            assert part.element.xml == original_footers[part.partname]


def test_unreadable_docx_reports_failure(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip")
    assert not apply_level_mark(path, "secret")
