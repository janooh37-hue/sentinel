"""Synthetic XLSX vehicle import parsing, review, and atomic-apply contracts."""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XlsxImage
from PIL import Image
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.config import get_settings
from app.core.vehicle_xlsx import (
    CANONICAL_COLUMNS,
    MAX_COMPRESSED_BYTES,
    VehicleXlsxError,
    build_vehicle_import_template,
    normalize_import_values,
    parse_vehicle_workbook,
)
from app.db.models import AuditLog, User, Vehicle, VehicleFile, VehicleSite
from app.db.session import get_db
from app.main import create_app
from app.schemas.vehicle import VehicleProfileScan
from app.services import vehicle_import_service, vehicle_service


@pytest.fixture()
def import_user(api_db: Session) -> User:
    user = User(
        email="vehicle-import@test.ae",
        password_hash="x",
        role="admin",
        status="active",
    )
    api_db.add(user)
    api_db.commit()
    api_db.refresh(user)
    return user


@pytest.fixture()
def import_client(
    api_db: Session,
    import_user: User,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> TestClient:
    settings = get_settings()
    monkeypatch.setattr(settings, "data_dir", tmp_path / "data")
    settings.ensure_dirs()
    app = create_app()
    app.dependency_overrides[get_db] = lambda: api_db
    app.dependency_overrides[get_current_user] = lambda: import_user
    return TestClient(app, raise_server_exceptions=True)


def _site(db: Session, name: str) -> VehicleSite:
    row = VehicleSite(name_ar=name, name_en=name, active=True)
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _vehicle(
    db: Session,
    site: VehicleSite,
    *,
    plate_code: str,
    plate_number: str,
    notes_ar: str | None = None,
    archived: bool = False,
) -> Vehicle:
    row = Vehicle(
        plate_code=plate_code,
        plate_number=plate_number,
        traffic_code="1180021637",
        type_ar="مركبة اختبار",
        type_en="مركبة اختبار",
        class_ar="فئة اختبار",
        class_en="فئة اختبار",
        vin=f"SYNTHETIC{plate_number}",
        site_id=site.id,
        license_start=date(2025, 1, 1),
        license_expiry=date(2027, 1, 1),
        notes_ar=notes_ar,
        archived_at=datetime(2026, 1, 1) if archived else None,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _jpeg(colour: str) -> bytes:
    output = io.BytesIO()
    with Image.new("RGB", (12, 8), colour) as image:
        image.save(output, format="JPEG")
    return output.getvalue()


def _add_image(sheet: Any, anchor: str, data: bytes) -> None:
    source = io.BytesIO(data)
    image = XlsxImage(source)
    image.width = 12
    image.height = 8
    sheet.add_image(image, anchor)


def _workbook_bytes(workbook: Any) -> bytes:
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _inject_vml_note_image(data: bytes, *, row_number: int, image_data: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        members = {name: source.read(name) for name in source.namelist()}

    sheet_name = "xl/worksheets/sheet1.xml"
    rel_name = "xl/worksheets/_rels/sheet1.xml.rels"
    sheet_xml = members[sheet_name]
    if b"xmlns:r=" not in sheet_xml:
        sheet_xml = sheet_xml.replace(
            b"<worksheet ",
            b'<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ',
            1,
        )
    sheet_xml = sheet_xml.replace(
        b"</worksheet>",
        (
            b'<legacyDrawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/'
            b'2006/relationships" r:id="rIdVehicleVml"/></worksheet>'
        ),
    )
    members[sheet_name] = sheet_xml

    relationships = ET.fromstring(members[rel_name])
    ET.SubElement(
        relationships,
        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
        {
            "Id": "rIdVehicleVml",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing",
            "Target": "../drawings/vehicleImport.vml",
        },
    )
    members[rel_name] = ET.tostring(relationships, encoding="utf-8", xml_declaration=True)
    members["xl/drawings/vehicleImport.vml"] = f"""<xml xmlns:v="urn:schemas-microsoft-com:vml"
 xmlns:o="urn:schemas-microsoft-com:office:office"
 xmlns:x="urn:schemas-microsoft-com:office:excel">
 <v:shape id="vehicle-note" type="#_x0000_t202">
  <v:fill o:relid="rIdVehicleImage" type="frame"/>
  <x:ClientData ObjectType="Note"><x:Row>{row_number - 1}</x:Row><x:Column>5</x:Column></x:ClientData>
 </v:shape>
</xml>""".encode()
    members["xl/drawings/_rels/vehicleImport.vml.rels"] = b"""<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
 <Relationship Id="rIdVehicleImage" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/vml-license.jpeg"/>
</Relationships>"""
    members["xl/media/vml-license.jpeg"] = image_data

    types = ET.fromstring(members["[Content_Types].xml"])
    if not any(node.get("Extension") == "vml" for node in types):
        ET.SubElement(
            types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
            {
                "Extension": "vml",
                "ContentType": "application/vnd.openxmlformats-officedocument.vmlDrawing",
            },
        )
    members["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8", xml_declaration=True)

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name, raw in members.items():
            target.writestr(name, raw)
    return output.getvalue()


def _legacy_workbook(*, class_name: str = "مركبة خفيفة") -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fleet"
    headers = [
        "ت",
        "الرقم \\ الفئة",
        "الرمز المروري",
        "تاريخ الترخيص",
        "انتهاء الترخيص",
        "نوع المركبة",
        "رقم القاعدة",
        "ملاحظات",
        "عدد النزلاء",
        "عدد الركاب",
        "ملحقات المركبة",
        "ملاحظات",
        "تفاصيل العقد",
    ]
    sheet.cell(1, 1, "كشف تفصيلي لمركبات الموقع الأول -2")
    for column, header in enumerate(headers, start=1):
        sheet.cell(2, column, header)

    def vehicle_row(
        row: int, plate: str, *, notes: str | None = None, traffic: object = 1180021637
    ) -> None:
        values = [
            row - 2,
            plate,
            traffic,
            "01/01/2025",
            "01/01/2027",
            "مركبة اختبار",
            f"SYNTHETIC-{row}",
            class_name,
            "٠ نزيل",  # noqa: RUF001 — Arabic-Indic digit is the point of this fixture
            0,
            None,
            notes,
            None,
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column, value)

    vehicle_row(3, "10 \\ 10001")
    vehicle_row(4, "10 \\ 10002")
    vehicle_row(5, "10 \\ 10003")
    vehicle_row(6, "10 \\ 10004")
    vehicle_row(7, "10 \\ 10004")
    vehicle_row(8, "not-a-plate")
    sheet.cell(10, 1, "كشف تفصيلي لمركبات الموقع الثاني -2")
    for column, header in enumerate(headers, start=1):
        sheet.cell(11, column, header)
    vehicle_row(12, "20 \\ 20001")
    normal_image = _jpeg("blue")
    _add_image(sheet, "F4", normal_image)
    data = _workbook_bytes(workbook)
    return _inject_vml_note_image(data, row_number=3, image_data=_jpeg("red"))


def _standard_workbook(
    rows: list[dict[str, object]], image_rows: list[tuple[int, str]] | None = None
) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vehicles"
    for column, header in enumerate(CANONICAL_COLUMNS, start=1):
        sheet.cell(1, column, header)
    for row_number, values in enumerate(rows, start=2):
        for column, header in enumerate(CANONICAL_COLUMNS, start=1):
            sheet.cell(row_number, column, values.get(header))
    for row_number, colour in image_rows or []:
        _add_image(sheet, f"X{row_number}", _jpeg(colour))
    return _workbook_bytes(workbook)


def _standard_row(plate_number: str, **overrides: object) -> dict[str, object]:
    values: dict[str, object] = {
        "plate_code": "7",
        "plate_number": plate_number,
        "traffic_code": "1180021637",
        "type_ar": "مركبة اختبار",
        "type_en": "Synthetic vehicle",
        "class_ar": "فئة اختبار",
        "class_en": "Synthetic class",
        "vin": f"VIN{plate_number}",
        "license_start": "2025-01-01",
        "license_expiry": "2027-01-01",
    }
    values.update(overrides)
    return values


def _preview_payload(
    inspection: dict[str, Any],
    mappings: dict[str, int],
    *,
    roles: dict[str, str] | None = None,
    reviewed: set[str] | None = None,
    confirmed: set[str] | None = None,
) -> dict[str, Any]:
    roles = roles or {}
    reviewed = reviewed or set()
    confirmed = confirmed or set()
    rows = []
    for row in inspection["rows"]:
        image_ids = row["image_ids"]
        row_roles = {image_id: roles[image_id] for image_id in image_ids if image_id in roles}
        license_ids = [
            image_id
            for image_id in image_ids
            if row_roles.get(image_id) == "license"
            or next(image for image in inspection["images"] if image["image_id"] == image_id)[
                "kind"
            ]
            == "license"
        ]
        photo_ids = [image_id for image_id in image_ids if row_roles.get(image_id) == "photo"]
        rows.append(
            {
                "row_id": row["row_id"],
                "excluded": False,
                "values": row["values"],
                "image_ids": image_ids,
                "image_roles": row_roles,
                "photo_action": "use_imported" if photo_ids else None,
                "primary_image_id": photo_ids[0] if photo_ids else None,
                "license_action": "use_imported" if license_ids else None,
                "license_image_id": license_ids[0] if license_ids else None,
                "ocr_reviewed_image_ids": [value for value in license_ids if value in reviewed],
                "ocr_manual_image_ids": [],
                "ocr_identity_confirmed_image_ids": [
                    value for value in license_ids if value in confirmed
                ],
            }
        )
    return {"site_mappings": mappings, "rows": rows, "excluded_image_ids": []}


def _inspect(client: TestClient, data: bytes, filename: str = "synthetic.xlsx") -> dict[str, Any]:
    response = client.post(
        "/api/v1/vehicles/imports/inspect",
        files={
            "file": (
                filename,
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_legacy_mixed_outcomes_images_atomic_apply_and_idempotent_reimport(
    import_client: TestClient,
    api_db: Session,
) -> None:
    first_site = _site(api_db, "First")
    second_site = _site(api_db, "Second")
    active = _vehicle(api_db, first_site, plate_code="10", plate_number="10002", notes_ar="keep me")
    archived = _vehicle(api_db, first_site, plate_code="10", plate_number="10003", archived=True)
    workbook = _legacy_workbook()

    inspection = _inspect(import_client, workbook)
    assert len(inspection["sections"]) == 2
    assert len(inspection["images"]) == 2
    assert {image["kind"] for image in inspection["images"]} == {"license"}
    assert all(image["row_id"] is not None for image in inspection["images"])
    served = import_client.get(inspection["images"][0]["url"])
    assert served.status_code == 200
    assert served.headers["content-type"] == "image/jpeg"
    assert served.headers["x-content-type-options"] == "nosniff"

    mappings = {
        inspection["sections"][0]["id"]: first_site.id,
        inspection["sections"][1]["id"]: second_site.id,
    }
    payload = _preview_payload(inspection, mappings)
    preview = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/preview", json=payload
    )
    assert preview.status_code == 200, preview.text
    rows = {row["row_id"]: row for row in preview.json()["rows"]}
    actions = [row["action"] for row in preview.json()["rows"]]
    assert actions.count("create") == 2
    assert actions.count("update") == 1
    assert actions.count("archived") == 1
    assert actions.count("invalid") == 3
    archived_row = next(row for row in rows.values() if row["vehicle_id"] == archived.id)
    assert archived_row["action"] == "archived"
    assert any(error["code"] == "VEHICLE_IMPORT_ARCHIVED_MATCH" for error in archived_row["errors"])
    active_row = next(row for row in rows.values() if row["vehicle_id"] == active.id)
    assert active_row["values"]["notes_ar"] == "keep me"
    assert active_row["current_photo_url"] is None
    assert all(image["kind"] == "license" for row in rows.values() for image in row["images"])

    selected = [
        row["row_id"]
        for row in preview.json()["rows"]
        if row["action"] in {"create", "update", "unchanged"}
    ]
    confirmed = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/confirm",
        json={"revision": preview.json()["revision"], "row_ids": selected},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json() == {
        "created": 2,
        "updated": 1,
        "unchanged": 0,
        "images_added": 2,
        "images_skipped": 0,
        "vehicle_ids": confirmed.json()["vehicle_ids"],
    }
    api_db.expire_all()
    # 2 pre-existing (active + archived, neither re-created) + 2 newly created.
    assert api_db.scalar(select(func.count()).select_from(Vehicle)) == 4
    assert api_db.scalar(select(func.count()).select_from(VehicleFile)) == 2
    assert api_db.get(Vehicle, active.id).notes_ar == "keep me"
    assert api_db.get(Vehicle, archived.id).archived_at is not None
    assert all(file.kind == "license" for file in api_db.scalars(select(VehicleFile)).all())
    assert all(vehicle.photo_file_id is None for vehicle in api_db.scalars(select(Vehicle)).all())

    second_inspection = _inspect(import_client, workbook)
    second_payload = _preview_payload(second_inspection, mappings)
    second_preview = import_client.post(
        f"/api/v1/vehicles/imports/{second_inspection['token']}/preview",
        json=second_payload,
    )
    assert second_preview.status_code == 200, second_preview.text
    second_rows = [
        row
        for row in second_preview.json()["rows"]
        if row["action"] in {"create", "update", "unchanged"}
    ]
    assert second_rows and {row["action"] for row in second_rows} == {"unchanged"}
    second_confirm = import_client.post(
        f"/api/v1/vehicles/imports/{second_inspection['token']}/confirm",
        json={
            "revision": second_preview.json()["revision"],
            "row_ids": [row["row_id"] for row in second_rows],
        },
    )
    assert second_confirm.status_code == 200, second_confirm.text
    assert second_confirm.json()["created"] == 0
    assert second_confirm.json()["updated"] == 0
    assert second_confirm.json()["images_added"] == 0
    assert second_confirm.json()["images_skipped"] == 2
    assert api_db.scalar(select(func.count()).select_from(Vehicle)) == 4
    assert api_db.scalar(select(func.count()).select_from(VehicleFile)) == 2


def test_confirm_rejects_stale_vehicle_and_rolls_back(
    import_client: TestClient, api_db: Session
) -> None:
    site = _site(api_db, "Stale")
    vehicle = _vehicle(api_db, site, plate_code="7", plate_number="30001", notes_ar="before")
    workbook = _standard_workbook([_standard_row("30001", notes_ar="from sheet")])
    inspection = _inspect(import_client, workbook)
    preview = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/preview",
        json=_preview_payload(inspection, {inspection["sections"][0]["id"]: site.id}),
    )
    assert preview.status_code == 200
    vehicle.notes_ar = "concurrent change"
    api_db.commit()

    response = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/confirm",
        json={
            "revision": preview.json()["revision"],
            "row_ids": [preview.json()["rows"][0]["row_id"]],
        },
    )
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "VEHICLE_IMPORT_STALE"
    api_db.expire_all()
    assert api_db.get(Vehicle, vehicle.id).notes_ar == "concurrent change"
    assert api_db.scalar(select(func.count()).select_from(Vehicle)) == 1


def test_confirm_same_token_twice_is_claimed(import_client: TestClient, api_db: Session) -> None:
    site = _site(api_db, "Once")
    inspection = _inspect(import_client, _standard_workbook([_standard_row("40001")]))
    preview = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/preview",
        json=_preview_payload(inspection, {inspection["sections"][0]["id"]: site.id}),
    ).json()
    body = {"revision": preview["revision"], "row_ids": [preview["rows"][0]["row_id"]]}
    first = import_client.post(f"/api/v1/vehicles/imports/{inspection['token']}/confirm", json=body)
    second = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/confirm", json=body
    )
    assert first.status_code == 200
    assert second.status_code == 409
    assert second.json()["error"]["code"] == "VEHICLE_IMPORT_TOKEN_CLAIMED"
    assert api_db.scalar(select(func.count()).select_from(Vehicle)) == 1


def test_scanned_identity_conflict_requires_explicit_confirmation(
    import_client: TestClient,
    api_db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    site = _site(api_db, "OCR")
    inspection = _inspect(import_client, _legacy_workbook())
    image = inspection["images"][0]
    monkeypatch.setattr(
        vehicle_import_service.vehicle_profile_scan_service,
        "scan_vehicle_profile",
        lambda _data: VehicleProfileScan(
            plate_code="99",
            plate_number="99999",
            warnings=["OCR_REVIEW_REQUIRED"],
        ),
    )
    scanned = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/images/{image['image_id']}/scan"
    )
    assert scanned.status_code == 200
    mappings = {section["id"]: site.id for section in inspection["sections"]}
    blocked = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/preview",
        json=_preview_payload(inspection, mappings),
    )
    assert blocked.status_code == 200
    blocked_row = next(
        row
        for row in blocked.json()["rows"]
        if image["image_id"] in [item["image_id"] for item in row["images"]]
    )
    assert blocked_row["action"] == "invalid"
    assert any(
        error["code"] == "VEHICLE_IMPORT_SCAN_IDENTITY_CONFLICT" for error in blocked_row["errors"]
    )

    accepted = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/preview",
        json=_preview_payload(
            inspection,
            mappings,
            reviewed={image["image_id"]},
            confirmed={image["image_id"]},
        ),
    )
    assert accepted.status_code == 200
    accepted_row = next(
        row
        for row in accepted.json()["rows"]
        if image["image_id"] in [item["image_id"] for item in row["images"]]
    )
    assert not any(
        error["code"] == "VEHICLE_IMPORT_SCAN_IDENTITY_CONFLICT" for error in accepted_row["errors"]
    )


def test_mid_batch_file_failure_rolls_back_db_and_only_new_files(
    import_client: TestClient,
    api_db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    site = _site(api_db, "Atomic")
    owner = _vehicle(api_db, site, plate_code="8", plate_number="50000")
    existing = vehicle_service.store_file(
        api_db,
        owner.id,
        kind="license",
        filename="existing.jpeg",
        data=_jpeg("white"),
        media_type="image/jpeg",
    )
    existing_path = (get_settings().data_dir / existing.path).resolve()
    existing_bytes = existing_path.read_bytes()
    workbook = _standard_workbook(
        [_standard_row("50001"), _standard_row("50002")],
        image_rows=[(2, "red"), (3, "blue")],
    )
    inspection = _inspect(import_client, workbook)
    roles = {image["image_id"]: "license" for image in inspection["images"]}
    preview = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/preview",
        json=_preview_payload(
            inspection,
            {inspection["sections"][0]["id"]: site.id},
            roles=roles,
        ),
    ).json()
    original = vehicle_service._store_file_record
    calls = 0

    def fail_second(*args: Any, **kwargs: Any) -> Any:
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError("synthetic second-file failure")
        return original(*args, **kwargs)

    monkeypatch.setattr(vehicle_service, "_store_file_record", fail_second)
    with pytest.raises(RuntimeError, match="synthetic second-file failure"):
        import_client.post(
            f"/api/v1/vehicles/imports/{inspection['token']}/confirm",
            json={
                "revision": preview["revision"],
                "row_ids": [row["row_id"] for row in preview["rows"]],
            },
        )
    api_db.expire_all()
    assert api_db.scalar(select(func.count()).select_from(Vehicle)) == 1
    assert api_db.scalar(select(func.count()).select_from(VehicleFile)) == 1
    assert existing_path.read_bytes() == existing_bytes
    vehicle_files = get_settings().data_dir / "vehicle_files"
    assert [path for path in vehicle_files.rglob("*") if path.is_file()] == [existing_path]


def test_generated_template_roundtrip_keeps_zero_dates_and_both_image_roles(
    import_client: TestClient,
    api_db: Session,
) -> None:
    site = _site(api_db, "Template")
    response = import_client.get("/api/v1/vehicles/imports/template")
    assert response.status_code == 200
    assert (
        response.headers["content-disposition"]
        == 'attachment; filename="vehicle-import-template.xlsx"'
    )
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Vehicles"]
    values = _standard_row(
        "00123",
        plate_code="04",
        inmate_capacity=0,
        passenger_capacity="٠",  # noqa: RUF001 — Arabic-Indic digit is the point of this fixture
        license_start=date(2025, 2, 1),
        license_expiry="01/02/2027",
        insurance_expiry="2026-12-31",
    )
    for column, name in enumerate(CANONICAL_COLUMNS, start=1):
        sheet.cell(2, column, values.get(name))
    _add_image(sheet, "X2", _jpeg("green"))
    _add_image(sheet, "Y2", _jpeg("yellow"))
    populated = _workbook_bytes(workbook)

    inspection = _inspect(import_client, populated)
    assert [image["kind"] for image in inspection["images"]] == [None, None]
    roles = {
        inspection["images"][0]["image_id"]: "photo",
        inspection["images"][1]["image_id"]: "license",
    }
    preview = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/preview",
        json=_preview_payload(
            inspection,
            {inspection["sections"][0]["id"]: site.id},
            roles=roles,
        ),
    )
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["action"] == "create"
    assert row["values"]["plate_code"] == "04"
    assert row["values"]["plate_number"] == "00123"
    assert row["values"]["inmate_capacity"] == 0
    assert row["values"]["passenger_capacity"] == 0
    assert row["values"]["license_start"] == "2025-02-01"
    confirmed = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/confirm",
        json={"revision": preview.json()["revision"], "row_ids": [row["row_id"]]},
    )
    assert confirmed.status_code == 200, confirmed.text
    api_db.expire_all()
    created = api_db.get(Vehicle, confirmed.json()["vehicle_ids"][0])
    assert created is not None
    assert created.inmate_capacity == 0
    assert created.passenger_capacity == 0
    assert created.license_start == date(2025, 2, 1)
    assert created.insurance_expiry == date(2026, 12, 31)
    assert created.photo_file_id is not None
    assert created.license_file_id is not None
    assert {file.kind for file in api_db.scalars(select(VehicleFile)).all()} == {"photo", "license"}
    imported_audits = api_db.scalars(
        select(AuditLog).where(AuditLog.action == "vehicle.imported")
    ).all()
    assert len(imported_audits) == 1


def test_core_legacy_detection_anchor_resolution_and_normalization() -> None:
    parsed = parse_vehicle_workbook(_legacy_workbook())
    assert parsed.layout == "legacy"
    assert len(parsed.sections) == 2
    assert len(parsed.images) == 2
    assert all(image.row_id is not None and image.kind == "license" for image in parsed.images)
    normalized, errors = normalize_import_values(
        {
            "combined_plate": "۱۴ \\ ۰۰۱۲۳",
            "traffic_code": "١١٨٠٠٢١٦٣٧",
            "license_start": "01\\02\\2025",
            "license_expiry": "2027-02-01",
            "inmate_capacity": "٠ نزيل",  # noqa: RUF001 — Arabic-Indic digit
        }
    )
    assert errors == []
    assert normalized["plate_code"] == "14"
    assert normalized["plate_number"] == "00123"
    assert normalized["traffic_code"] == "1180021637"
    assert normalized["license_start"] == "2025-02-01"
    assert normalized["inmate_capacity"] == 0
    _, invalid = normalize_import_values({"license_start": "31/02/2027"})
    assert invalid == [("license_start", "Invalid calendar date.")]


@pytest.mark.parametrize(
    ("legacy_class", "expected_ar", "expected_en"),
    [
        ("بيك اب  ثقيل", "بيك أب ثقيل", "Heavy pickup"),
        ("بيك اب ثقيل", "بيك أب ثقيل", "Heavy pickup"),
        ("فرع الامن", "فرع الأمن", "Security branch"),
    ],
)
def test_legacy_class_aliases_resolve_to_exact_bilingual_presets(
    legacy_class: str,
    expected_ar: str,
    expected_en: str,
) -> None:
    parsed = parse_vehicle_workbook(_legacy_workbook(class_name=legacy_class))
    assert parsed.rows[0].values["class_ar"] == expected_ar
    assert parsed.rows[0].values["class_en"] == expected_en


def test_legacy_unknown_class_is_preserved_for_custom_review(
    import_client: TestClient,
    api_db: Session,
) -> None:
    custom_class = "فئة مخصصة للمهمة"
    site = _site(api_db, "Custom class")
    inspection = _inspect(import_client, _legacy_workbook(class_name=custom_class))
    inspected_row = inspection["rows"][0]
    assert inspected_row["raw"]["class_ar"] == custom_class
    assert inspected_row["values"]["class_ar"] == custom_class
    assert inspected_row["values"]["class_en"] is None

    response = import_client.post(
        f"/api/v1/vehicles/imports/{inspection['token']}/preview",
        json=_preview_payload(
            inspection,
            {section["id"]: site.id for section in inspection["sections"]},
        ),
    )
    assert response.status_code == 200, response.text
    preview_row = response.json()["rows"][0]
    assert preview_row["action"] == "invalid"
    assert any(
        error["code"] == "VEHICLE_IMPORT_REQUIRED_FIELD" and error["field"] == "class_en"
        for error in preview_row["errors"]
    )


def test_core_formula_corrupt_and_oversized_rejections() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for column, header in enumerate(CANONICAL_COLUMNS, start=1):
        sheet.cell(1, column, header)
    sheet.cell(2, 2, "=1+1")
    sheet.cell(2, 3, "1180021637")
    parsed = parse_vehicle_workbook(_workbook_bytes(workbook))
    assert any(
        warning.code == "VEHICLE_IMPORT_INVALID_FIELD" and warning.field == "plate_number"
        for warning in parsed.warnings
    )
    with pytest.raises(VehicleXlsxError, match=r"valid \.xlsx") as corrupt:
        parse_vehicle_workbook(b"not a zip")
    assert corrupt.value.code == "VEHICLE_IMPORT_BAD_FILE"
    with pytest.raises(VehicleXlsxError) as oversized:
        parse_vehicle_workbook(b"x" * (MAX_COMPRESSED_BYTES + 1))
    assert oversized.value.code == "VEHICLE_IMPORT_TOO_LARGE"


def test_generated_template_has_canonical_order_instructions_and_no_formulas() -> None:
    workbook = load_workbook(io.BytesIO(build_vehicle_import_template()), data_only=False)
    assert workbook.sheetnames == ["Vehicles", "Instructions - التعليمات"]
    assert tuple(cell.value for cell in workbook["Vehicles"][1]) == CANONICAL_COLUMNS
    assert all(
        cell.data_type != "f"
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
    )
    assert workbook["Vehicles"].max_row == 1
    instructions = workbook["Instructions - التعليمات"]
    overview_en = [cell.value for cell in instructions["B"] if isinstance(cell.value, str)]
    overview_ar = [cell.value for cell in instructions["C"] if isinstance(cell.value, str)]
    assert any("blank cell preserves the existing stored value" in value for value in overview_en)
    assert any("الخلية الفارغة القيمة الحالية المخزنة كما هي" in value for value in overview_ar)

    header_row = next(
        cell.row for cell in instructions["A"] if cell.value == "Canonical column / اسم العمود"
    )
    documented_columns = tuple(
        instructions.cell(row=row_number, column=1).value
        for row_number in range(header_row + 1, instructions.max_row + 1)
    )
    assert documented_columns == CANONICAL_COLUMNS
    required_columns = {
        "plate_number",
        "traffic_code",
        "type_ar",
        "type_en",
        "class_ar",
        "class_en",
        "license_start",
        "license_expiry",
    }
    for row_number, column_name in enumerate(CANONICAL_COLUMNS, start=header_row + 1):
        assert instructions.cell(row_number, 2).value
        assert instructions.cell(row_number, 3).value
        expected = "Yes / نعم" if column_name in required_columns else "No / لا"
        assert instructions.cell(row_number, 4).value == expected

    values = _standard_row("00123")
    for column, name in enumerate(CANONICAL_COLUMNS, start=1):
        workbook["Vehicles"].cell(2, column, values.get(name))
    parsed = parse_vehicle_workbook(_workbook_bytes(workbook))
    assert parsed.layout == "standard"
    assert parsed.rows[0].values["plate_number"] == "00123"
