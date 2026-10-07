"""Regression: the rendered transfer letter must not repeat the General Book
paper's own greeting/closing (duty_service -> document_service -> DocxEngine)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.config import Settings
from app.db.models import BookCategory, Document, Employee
from app.schemas.duty import DutyTransferMove
from app.services import artifact_service, document_service, duty_service


@pytest.fixture()
def gen_env(db_session, tmp_path, monkeypatch):
    settings = Settings(data_dir=tmp_path / "data")
    monkeypatch.setattr(document_service, "get_settings", lambda: settings)
    monkeypatch.setattr(artifact_service, "get_settings", lambda: settings)
    monkeypatch.setattr(document_service, "convert_docx_to_pdf", lambda p: None)
    db_session.add(BookCategory(id="GS", prefix="GS"))
    db_session.add_all(
        [
            Employee(
                id="G3309",
                name_en="Majid",
                name_ar="ماجد",
                position_ar="حارس أمن",
                duty_unit="السرية الخامسة",
                duty_post="تفتيش",
            ),
            Employee(
                id="G4030",
                name_en="Saif",
                name_ar="سيف",
                position_ar="حارس أمن",
                duty_unit="السرية الثانية",
                duty_post="التفتيش",
            ),
        ]
    )
    db_session.commit()
    return db_session, settings


def test_rendered_swap_letter_has_greeting_and_closing_once(gen_env, admin_user):
    db, settings = gen_env
    result = duty_service.transfer(
        db,
        moves=[
            DutyTransferMove(employee_id="G3309", to_unit="السرية الثانية", to_post="التفتيش"),
            DutyTransferMove(employee_id="G4030", to_unit="السرية الخامسة", to_post="تفتيش"),
        ],
        cc=["مدراء الأفرع"],
        current_user=admin_user,
    )
    docx = Path(db.get(Document, result.document_id).docx_path)
    if not docx.is_absolute():
        docx = Path(settings.data_dir) / docx
    doc = DocxDocument(str(docx))

    text = "\n".join(p.text for p in doc.paragraphs)
    assert text.count("يطيب لنا أن نتقدم") == 1
    assert text.count("للتفضل بالعلم") == 1
    assert text.count("تحية طيبة") == 1
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3  # header + 2 data rows
    assert all(len(r.cells) == 5 for r in table.rows)
