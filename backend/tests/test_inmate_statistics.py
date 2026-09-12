"""The monthly inmate violation register: projection, overlay, seal, export.

Every test here asserts a rule the wayfinder map fixed: which occurrences a
month counts, which table an occurrence lands in, what a close refuses, and what
a sealed month keeps saying once its sources move underneath it.
"""

from __future__ import annotations

import io
import json
from datetime import date, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker

from app.api.deps import get_current_user
from app.api.errors import ValidationFailedError
from app.config import get_settings
from app.db import session as session_mod
from app.db.models import (
    AuditLog,
    Base,
    Book,
    BookCategory,
    BookVersion,
    Employee,
    InmateViolationManualRow,
    InmateViolationPeriod,
    InmateViolationStatRow,
    User,
)
from app.db.session import attach_sqlite_pragmas, get_db
from app.db.workforce_models import DutyAssignmentEvent
from app.main import create_app
from app.services import inmate_statistics_service as register
from app.services import perm_service

TEMPLATE_ID = register.TEMPLATE_ID
XLSX_PREFIX = "application/vnd.openxmlformats-officedocument.spreadsheetml"

#: 2026-08 is over, so it may be closed; 2026-09 stands in for a running month.
CLOSED_YEAR, CLOSED_MONTH = 2026, 8
TODAY = date(2026, 9, 10)


@pytest.fixture()
def api_db(monkeypatch, tmp_path) -> Session:
    # The close-time workbook copy lands under the data root; keep every test
    # out of the live one.
    monkeypatch.setenv("GSSG_DATA_DIR", str(tmp_path / "data"))
    get_settings.cache_clear()
    db_file = tmp_path / "register.db"
    eng = create_engine(
        f"sqlite:///{db_file}", future=True, connect_args={"check_same_thread": False}
    )
    attach_sqlite_pragmas(eng, wal=False)
    Base.metadata.create_all(eng)
    TestSession = sessionmaker(bind=eng, autoflush=False, expire_on_commit=False, future=True)
    monkeypatch.setattr(session_mod, "engine", eng)
    monkeypatch.setattr(session_mod, "SessionLocal", TestSession)
    db = TestSession()
    perm_service.seed_role_defaults(db)
    db.add(BookCategory(id="NAT", prefix="IV"))
    db.commit()
    try:
        yield db
    finally:
        db.close()


@pytest.fixture()
def db_session(api_db) -> Session:
    """Shadow the conftest fixture so client and setup share one database."""

    return api_db


def _client_for(api_db: Session, role: str, email: str) -> TestClient:
    user = User(email=email, password_hash="x", role=role, status="active")
    api_db.add(user)
    api_db.commit()
    api_db.refresh(user)
    app = create_app()
    app.dependency_overrides[get_db] = lambda: api_db
    app.dependency_overrides[get_current_user] = lambda: user
    client = TestClient(app, raise_server_exceptions=True)
    client.user_id = user.id  # type: ignore[attr-defined]
    return client


@pytest.fixture()
def admin_client(api_db: Session) -> TestClient:
    return _client_for(api_db, "admin", "admin@x.ae")


@pytest.fixture()
def operator_client(api_db: Session) -> TestClient:
    """A manager: reads and writes the register, but may not close a month."""

    return _client_for(api_db, "manager", "mgr@x.ae")


def _actor(db: Session, *, role: str = "admin", email: str = "actor@x.ae") -> User:
    user = User(email=email, password_hash="x", role=role, status="active")
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


def _approve_test_month(db: Session, year: int, month: int, *, actor: User, today: date):
    """Exercise the real three-person service chain for existing seal tests."""
    if not actor.employee_id:
        employee = Employee(
            id=f"GACT{actor.id}", name_en="Synthetic manager", name_ar="مدير تجريبي"
        )
        db.add(employee)
        db.flush()
        actor.employee_id = employee.id
        db.commit()
    suffix = len(db.scalars(select(User)).all())
    prep = _actor(db, email=f"prep{suffix}@example.test")
    reviewer = _actor(db, email=f"review{suffix}@example.test")
    for user in (prep, reviewer):
        employee = Employee(id=f"GACT{user.id}", name_en="Synthetic actor", name_ar="موظف تجريبي")
        db.add(employee)
        db.flush()
        user.employee_id = employee.id
    db.commit()
    view = register.build_month(db, year, month)
    view = register.prepare_month(
        db,
        year,
        month,
        actor=prep,
        expected_version=view.workflow["version"],
        expected_projection_fingerprint=view.projection_fingerprint,
        reviewer_user_id=reviewer.id,
    )
    view = register.review_month(
        db,
        year,
        month,
        actor=reviewer,
        expected_version=view.workflow["version"],
        submission_id=view.workflow["active_submission_id"],
        manager_user_id=actor.id,
    )
    return register.approve_month(
        db,
        year,
        month,
        actor=actor,
        expected_version=view.workflow["version"],
        submission_id=view.workflow["active_submission_id"],
        today=today,
    )


def _reporter(
    db: Session,
    employee_id: str = "G4603",
    *,
    duty_unit: str | None = "السرية الثانية",
    name_ar: str = "عبدالله محمد الكعبي",
) -> Employee:
    employee = Employee(
        id=employee_id,
        name_en="ABDULLA M ALKAABI",
        name_ar=name_ar,
        duty_unit=duty_unit,
    )
    db.add(employee)
    db.commit()
    return employee


_REF = {"n": 0}


def _record(
    db: Session,
    *,
    report_date: str | None,
    inmates: list[dict[str, str]] | None = None,
    details: str = "<p>مخالفة</p>",
    reporter_id: str | None = "G4603",
    ref: str | None = None,
    versions: int = 1,
    deleted: bool = False,
    voided: bool = False,
    imported_names: list[str] | None = None,
) -> Book:
    """One filed Inmate Conduct Violations Record with its current version."""

    _REF["n"] += 1
    book = Book(
        category_id="NAT",
        ref_number=ref or f"IV-2026-{_REF['n']:04d}",
        subject="مخالفة سلوكية",
        direction="outgoing",
        approval_state="approved",
        deleted_at=datetime(2026, 8, 20) if deleted else None,
        voided_at=datetime(2026, 8, 20) if voided else None,
        created_at=datetime(2026, 8, 20, 7, 15),
    )
    db.add(book)
    db.flush()
    for version_no in range(1, versions + 1):
        if imported_names is not None:
            fields: dict[str, object] = {
                "report_date": report_date,
                "inmate_names": imported_names,
                "subject": "مخالفة سلوكية",
                "imported_approved": True,
            }
        else:
            fields = {
                "report_date": report_date,
                "report_time": "23:40",
                "reporter_id": reporter_id,
                "violation_details": details,
                "inmates": inmates or [],
            }
        db.add(
            BookVersion(
                book_id=book.id,
                version_no=version_no,
                template_id=TEMPLATE_ID,
                fields=fields,
                status="approved",
            )
        )
    db.commit()
    db.refresh(book)
    return book


def _inmate(
    name: str,
    *,
    nationality: str = "الامارات",
    uid: str = "784-1993-4471820-3",
    wing: str = "3A",
    holding_no: str = "H-2291",
) -> dict[str, str]:
    return {
        "name": name,
        "nationality": nationality,
        "wing": wing,
        "uid": uid,
        "holding_no": holding_no,
    }


def _by_name(month: register.MonthRegister, name: str) -> register.RegisterEntry:
    return next(entry for entry in month.entries if entry.name == name)


# --------------------------------------------------------------------------- #
# what a month counts
# --------------------------------------------------------------------------- #


def test_month_is_the_calendar_month_of_the_report_date(db_session: Session) -> None:
    _reporter(db_session)
    _record(
        db_session,
        report_date="2026-08-05",
        inmates=[_inmate("راشد سعيد المنصوري")],
    )

    august = register.build_month(db_session, 2026, 8)
    september = register.build_month(db_session, 2026, 9)

    # Filed at 07:15 on 2026-08-20 (``Book.created_at``), incident on 08-05.
    assert [entry.name for entry in august.entries] == ["راشد سعيد المنصوري"]
    assert september.entries == ()


def test_only_the_current_version_of_a_live_record_contributes(db_session: Session) -> None:
    _reporter(db_session)
    revised = _record(
        db_session,
        report_date="2026-08-05",
        inmates=[_inmate("راشد سعيد المنصوري")],
        versions=2,
    )
    _record(db_session, report_date="2026-08-06", inmates=[_inmate("محمد علي")], deleted=True)
    _record(db_session, report_date="2026-08-07", inmates=[_inmate("سعيد جمعة")], voided=True)

    month = register.build_month(db_session, 2026, 8)

    assert [entry.name for entry in month.entries] == ["راشد سعيد المنصوري"]
    assert month.entries[0].handle == f"{revised.id}:2:0"


def test_a_record_with_no_usable_date_is_flagged_not_dropped(db_session: Session) -> None:
    _reporter(db_session)
    undated = _record(
        db_session,
        report_date="",
        inmates=[_inmate("بدون تاريخ")],
        ref="IV-2026-0141",
    )

    month = register.build_month(db_session, 2026, 8)

    assert month.entries == ()
    assert [(item.book_id, item.ref_number) for item in month.uncounted] == [
        (undated.id, "IV-2026-0141")
    ]


def test_every_occurrence_counts_including_repeats(db_session: Session) -> None:
    _reporter(db_session)
    _record(
        db_session,
        report_date="2026-08-05",
        inmates=[_inmate("راشد سعيد المنصوري"), _inmate("راشد سعيد المنصوري")],
    )

    month = register.build_month(db_session, 2026, 8)

    assert month.counts["total"] == 2
    assert [entry.row_no for entry in month.entries] == [1, 2]


# --------------------------------------------------------------------------- #
# populations, ordinals, order
# --------------------------------------------------------------------------- #


def test_nationality_decides_the_population_and_ordinals_restart(db_session: Session) -> None:
    _reporter(db_session)
    _record(
        db_session,
        report_date="2026-08-05",
        inmates=[
            _inmate("راشد سعيد المنصوري", nationality="الامارات"),
            _inmate("ANWAR HOSSAIN", nationality="بنغلاديش"),
            _inmate("MUHAMMAD IMRAN", nationality="باكستان"),
        ],
    )

    month = register.build_month(db_session, 2026, 8)

    assert _by_name(month, "راشد سعيد المنصوري").population == register.POPULATION_CITIZENS
    assert _by_name(month, "ANWAR HOSSAIN").population == register.POPULATION_EXPATS
    assert _by_name(month, "راشد سعيد المنصوري").row_no == 1
    # Each printed table is read alone, so each starts at 1.
    assert _by_name(month, "ANWAR HOSSAIN").row_no == 1
    assert _by_name(month, "MUHAMMAD IMRAN").row_no == 2
    assert month.counts == {"citizens": 1, "expats": 2, "pending": 0, "total": 3}


def test_unmatched_history_stays_an_expat_with_an_incomplete_mark(db_session: Session) -> None:
    _reporter(db_session)
    _record(
        db_session,
        report_date="2026-08-05",
        inmates=[_inmate("مجهول الجنسية", nationality="الإ")],
    )

    entry = register.build_month(db_session, 2026, 8).entries[0]

    assert entry.population == register.POPULATION_EXPATS
    assert entry.nationality_label == register.UNSPECIFIED
    assert "nationality" in entry.incomplete_marks
    # Terminal, not completable: the bridge resolved, it just found nothing.
    assert entry.missing == []


def test_an_entry_with_no_nationality_has_no_population(db_session: Session) -> None:
    _reporter(db_session)
    _record(
        db_session,
        report_date="2026-08-05",
        inmates=[_inmate("قيد الإكمال", nationality="")],
    )

    entry = register.build_month(db_session, 2026, 8).entries[0]

    assert entry.population == register.POPULATION_PENDING
    assert entry.missing == ["nationality"]


def test_details_are_flattened_plain_text_repeated_per_row(db_session: Session) -> None:
    _reporter(db_session)
    _record(
        db_session,
        report_date="2026-08-05",
        inmates=[_inmate("الأول"), _inmate("الثاني")],
        details="<p>سطر أول</p><p>سطر ثانٍ &amp; ثالث</p>",
    )

    month = register.build_month(db_session, 2026, 8)

    assert month.entries[0].details_text == "سطر أول\nسطر ثانٍ & ثالث"
    assert month.entries[1].details_text == month.entries[0].details_text


# --------------------------------------------------------------------------- #
# السربة
# --------------------------------------------------------------------------- #


def test_duty_unit_resolves_as_of_the_violation_date(db_session: Session) -> None:
    _reporter(db_session, duty_unit="السرية الخامسة")
    mover = _actor(db_session, email="mover@x.ae")
    db_session.add_all(
        [
            DutyAssignmentEvent(
                employee_id="G4603",
                event_type="baseline",
                to_unit="السرية الأولى",
                effective_at=datetime(2026, 1, 1),
            ),
            DutyAssignmentEvent(
                employee_id="G4603",
                event_type="transfer",
                actor_user_id=mover.id,
                to_unit="السرية الخامسة",
                effective_at=datetime(2026, 9, 1),
            ),
        ]
    )
    db_session.commit()
    _record(db_session, report_date="2026-08-05", inmates=[_inmate("راشد")])

    entry = register.build_month(db_session, 2026, 8).entries[0]

    # A reporter who transfers later does not re-label their earlier entries.
    assert entry.duty_unit == "السرية الأولى"


def test_office_hours_prints_verbatim_and_a_missing_unit_guards(db_session: Session) -> None:
    _reporter(db_session, "G4029", duty_unit="الدوام الرسمي")
    _reporter(db_session, "G9999", duty_unit=None, name_ar="بدون وحدة")
    _record(db_session, report_date="2026-08-05", inmates=[_inmate("أ")], reporter_id="G4029")
    _record(db_session, report_date="2026-08-06", inmates=[_inmate("ب")], reporter_id="G9999")

    month = register.build_month(db_session, 2026, 8)

    assert _by_name(month, "أ").duty_unit == "الدوام الرسمي"
    assert _by_name(month, "ب").duty_unit == register.UNSPECIFIED


# --------------------------------------------------------------------------- #
# manual entries
# --------------------------------------------------------------------------- #


def test_manual_entry_derives_its_duty_unit_and_sorts_after_papered_ones(
    db_session: Session,
) -> None:
    _reporter(db_session)
    _record(db_session, report_date="2026-08-05", inmates=[_inmate("الورقي")])
    register.create_manual_row(
        db_session,
        2026,
        8,
        actor=_actor(db_session),
        name="اليدوي",
        violation_date=date(2026, 8, 5),
        reason="تقرير ورقي مفقود",
        nationality_label="الإمارات",
        reporter_id="G4603",
        details_text="مشادة كلامية",
    )

    month = register.build_month(db_session, 2026, 8)

    assert [entry.name for entry in month.entries] == ["الورقي", "اليدوي"]
    manual = _by_name(month, "اليدوي")
    assert manual.handle.startswith("manual:")
    assert manual.duty_unit == "السرية الثانية"
    assert manual.source_ref_number is None
    assert manual.manual_reason == "تقرير ورقي مفقود"


def test_manual_entry_without_details_blocks_the_close(db_session: Session) -> None:
    actor = _actor(db_session)
    register.create_manual_row(
        db_session,
        2026,
        8,
        actor=actor,
        name="بلا تفاصيل",
        violation_date=date(2026, 8, 9),
        reason="ورقة مفقودة",
        nationality_label="الإمارات",
    )

    month = register.build_month(db_session, 2026, 8)

    assert month.entries[0].missing == ["wing", "details"]


def test_a_manual_entry_shadowing_a_record_warns_but_never_merges(db_session: Session) -> None:
    _reporter(db_session)
    _record(db_session, report_date="2026-08-05", inmates=[_inmate("راشد  سعيد")])
    register.create_manual_row(
        db_session,
        2026,
        8,
        actor=_actor(db_session),
        name="راشد سعيد",
        violation_date=date(2026, 8, 5),
        reason="أدخل يدويًا قبل ورود التقرير",
        nationality_label="الإمارات",
        details_text="تفاصيل",
    )

    month = register.build_month(db_session, 2026, 8)

    assert month.counts["total"] == 2
    assert [entry.handle for entry in month.duplicates] == [_by_name(month, "راشد سعيد").handle]


def test_manual_entry_operations_are_audited(db_session: Session) -> None:
    actor = _actor(db_session)
    row = register.create_manual_row(
        db_session,
        2026,
        8,
        actor=actor,
        name="اسم أول",
        violation_date=date(2026, 8, 9),
        reason="سبب",
    )
    register.update_manual_row(db_session, row.id, actor=actor, changes={"name": "اسم ثانٍ"})
    register.delete_manual_row(db_session, row.id, actor=actor)

    actions = [
        (log.action, log.entity_id)
        for log in db_session.scalars(select(AuditLog).order_by(AuditLog.id))
    ]
    assert actions == [
        ("inmate_violation_manual_row_created", f"manual:{row.id}"),
        ("inmate_violation_manual_row_updated", f"manual:{row.id}"),
        ("inmate_violation_manual_row_deleted", f"manual:{row.id}"),
    ]
    updated = json.loads(
        db_session.scalars(
            select(AuditLog.payload).where(AuditLog.action == "inmate_violation_manual_row_updated")
        ).one()
    )
    assert updated["changed"]["name"] == {"from": "اسم أول", "to": "اسم ثانٍ"}
    assert db_session.get(InmateViolationManualRow, row.id) is None


def test_a_manual_entry_must_fall_inside_its_month(db_session: Session) -> None:
    with pytest.raises(Exception) as raised:
        register.create_manual_row(
            db_session,
            2026,
            8,
            actor=_actor(db_session),
            name="خارج الشهر",
            violation_date=date(2026, 9, 1),
            reason="سبب",
        )
    assert "INMATE_REGISTER_DATE_OUTSIDE_MONTH" in str(raised.value.code)  # type: ignore[union-attr]


# --------------------------------------------------------------------------- #
# completing an uploaded approved copy
# --------------------------------------------------------------------------- #


def test_an_uncompleted_import_projects_one_pending_entry_per_name(db_session: Session) -> None:
    book = _record(
        db_session,
        report_date="2026-08-11",
        imported_names=["حمد علي المزروعي", "سالم راشد"],
    )

    month = register.build_month(db_session, 2026, 8)

    assert [entry.name for entry in month.entries] == ["حمد علي المزروعي", "سالم راشد"]
    assert {entry.population for entry in month.entries} == {register.POPULATION_PENDING}
    assert all(entry.completion_book_id == book.id for entry in month.entries)
    # Underivable, not empty: an import carries no reporter until completion.
    assert month.entries[0].duty_unit == register.UNSPECIFIED


def test_completion_merges_a_namespaced_block_and_leaves_the_import_intact(
    db_session: Session,
) -> None:
    _reporter(db_session)
    book = _record(db_session, report_date="2026-08-11", imported_names=["حمد علي المزروعي"])

    register.complete_import(
        db_session,
        book.id,
        actor=_actor(db_session),
        report_time="23:40",
        reporter_id="G4603",
        violation_details="<p>تفاصيل من الصورة</p>",
        inmates=[{"name": "حمد علي المزروعي", "nationality": "الإمارات", "uid": "", "wing": ""}],
    )

    version = max(book.versions, key=lambda item: item.version_no)
    db_session.refresh(version)
    fields = version.fields or {}
    assert fields["imported_approved"] is True
    assert fields["inmate_names"] == ["حمد علي المزروعي"]
    assert fields["report_date"] == "2026-08-11"
    assert fields["completion"]["reporter_id"] == "G4603"

    entry = register.build_month(db_session, 2026, 8).entries[0]
    assert entry.population == register.POPULATION_CITIZENS
    assert entry.duty_unit == "السرية الثانية"
    assert entry.details_text == "تفاصيل من الصورة"
    # Empty UID remains optional; a draft wing must be completed for approval.
    assert (entry.uid, entry.wing) == ("", "")
    assert entry.missing == ["wing"]

    logged = db_session.scalars(
        select(AuditLog).where(AuditLog.action == "inmate_violation_import_completed")
    ).all()
    assert len(logged) == 1


def test_completion_refuses_without_reporter_or_narrative(db_session: Session) -> None:
    book = _record(db_session, report_date="2026-08-11", imported_names=["حمد"])
    actor = _actor(db_session)

    with pytest.raises(ValidationFailedError):
        register.complete_import(
            db_session,
            book.id,
            actor=actor,
            report_time="23:40",
            reporter_id="",
            violation_details="تفاصيل",
            inmates=[{"name": "حمد"}],
        )


# --------------------------------------------------------------------------- #
# close, reopen, re-close
# --------------------------------------------------------------------------- #


def _closable_month(db: Session) -> None:
    _reporter(db)
    _record(
        db,
        report_date="2026-08-05",
        inmates=[_inmate("راشد سعيد المنصوري"), _inmate("ANWAR HOSSAIN", nationality="نيبال")],
    )


def test_a_running_month_cannot_be_closed(db_session: Session) -> None:
    _reporter(db_session)
    _record(db_session, report_date="2026-09-03", inmates=[_inmate("راشد")])

    with pytest.raises(Exception) as raised:
        _approve_test_month(db_session, 2026, 9, actor=_actor(db_session), today=TODAY)

    assert "INMATE_REGISTER_MONTH_NOT_ENDED" in str(raised.value.code)  # type: ignore[union-attr]
    assert register.first_closable_date(2026, 9) == date(2026, 10, 1)


def test_a_get_never_creates_a_period_row(db_session: Session) -> None:
    _closable_month(db_session)

    register.build_month(db_session, CLOSED_YEAR, CLOSED_MONTH)

    assert db_session.scalars(select(InmateViolationPeriod)).all() == []


def test_close_freezes_the_entries_and_audits(db_session: Session) -> None:
    _closable_month(db_session)
    actor = _actor(db_session)

    month = _approve_test_month(db_session, CLOSED_YEAR, CLOSED_MONTH, actor=actor, today=TODAY)

    assert month.closed is True
    assert month.closed_by == actor.id
    assert month.counts["total"] == 2
    frozen = db_session.scalars(select(InmateViolationStatRow)).all()
    assert {row.row_handle for row in frozen} == {entry.handle for entry in month.entries}
    assert {row.duty_unit for row in frozen} == {"السرية الثانية"}
    logged = db_session.scalars(
        select(AuditLog).where(AuditLog.action == "inmate_violation_month_closed")
    ).all()
    assert json.loads(logged[0].payload or "{}")["row_count"] == 2


def test_a_sealed_month_is_never_re_derived(db_session: Session) -> None:
    _closable_month(db_session)
    _approve_test_month(
        db_session, CLOSED_YEAR, CLOSED_MONTH, actor=_actor(db_session), today=TODAY
    )

    # The source paper is deleted after the seal.
    book = db_session.scalars(select(Book)).first()
    assert book is not None
    book.deleted_at = datetime(2026, 9, 2)
    db_session.commit()

    month = register.build_month(db_session, CLOSED_YEAR, CLOSED_MONTH)

    assert month.counts["total"] == 2
    assert month.entries[0].source_ref_number == book.ref_number


def test_final_approval_refuses_pending_completion_without_bypass(db_session: Session) -> None:
    _reporter(db_session)
    _record(db_session, report_date="2026-08-11", imported_names=["نزيل تجريبي"])
    with pytest.raises(Exception) as raised:
        _approve_test_month(
            db_session, CLOSED_YEAR, CLOSED_MONTH, actor=_actor(db_session), today=TODAY
        )
    assert raised.value.code == "INMATE_REGISTER_INCOMPLETE_ENTRIES"
    assert not register.build_month(db_session, CLOSED_YEAR, CLOSED_MONTH).closed


def test_a_closed_month_admits_no_write(db_session: Session) -> None:
    _closable_month(db_session)
    actor = _actor(db_session)
    _approve_test_month(db_session, CLOSED_YEAR, CLOSED_MONTH, actor=actor, today=TODAY)

    with pytest.raises(Exception) as raised:
        register.create_manual_row(
            db_session,
            CLOSED_YEAR,
            CLOSED_MONTH,
            actor=actor,
            name="متأخر",
            violation_date=date(2026, 8, 9),
            reason="سبب",
        )

    assert "INMATE_REGISTER_MONTH_CLOSED" in str(raised.value.code)  # type: ignore[union-attr]


def test_an_occurrence_edited_into_a_closed_month_is_listed_outside_the_seal(
    db_session: Session,
) -> None:
    _closable_month(db_session)
    _approve_test_month(
        db_session, CLOSED_YEAR, CLOSED_MONTH, actor=_actor(db_session), today=TODAY
    )
    late = _record(db_session, report_date="2026-08-30", inmates=[_inmate("وصل بعد الإغلاق")])

    month = register.build_month(db_session, CLOSED_YEAR, CLOSED_MONTH)

    assert month.counts["total"] == 2
    assert [item.name for item in month.arrived_after_close] == ["وصل بعد الإغلاق"]
    assert month.arrived_after_close[0].source_book_id == late.id


def test_reopen_keeps_the_previous_seal_and_re_close_never_downgrades_the_duty_unit(
    db_session: Session,
) -> None:
    _closable_month(db_session)
    actor = _actor(db_session)
    _approve_test_month(db_session, CLOSED_YEAR, CLOSED_MONTH, actor=actor, today=TODAY)

    version = register.build_month(db_session, CLOSED_YEAR, CLOSED_MONTH).workflow["version"]
    reopened = register.reopen_month(
        db_session, CLOSED_YEAR, CLOSED_MONTH, actor=actor, expected_version=version, reason="تصحيح"
    )

    assert reopened.closed is False
    # The previous seal survives a reopen; entry presence never signals closure.
    assert len(db_session.scalars(select(InmateViolationStatRow)).all()) == 2

    # The reporter's placement is gone by the time the month is re-closed —
    # a purged duty event must not silently rewrite a sealed unit.
    employee = db_session.get(Employee, "G4603")
    assert employee is not None
    employee.duty_unit = None
    db_session.commit()

    re_closed = _approve_test_month(db_session, CLOSED_YEAR, CLOSED_MONTH, actor=actor, today=TODAY)

    assert {entry.duty_unit for entry in re_closed.entries} == {"السرية الثانية"}
    assert db_session.scalars(
        select(AuditLog.action).where(AuditLog.action == "inmate_violation_month_reopened")
    ).all() == ["inmate_violation_month_reopened"]


# --------------------------------------------------------------------------- #
# awaiting close
# --------------------------------------------------------------------------- #


def test_awaiting_close_lists_ended_months_with_entries_and_no_seal(db_session: Session) -> None:
    _reporter(db_session)
    _record(db_session, report_date="2026-07-04", inmates=[_inmate("يوليو")])
    _record(db_session, report_date="2026-08-05", imported_names=["أغسطس قيد الإكمال"])
    _record(db_session, report_date="2026-09-03", inmates=[_inmate("سبتمبر الجاري")])
    _approve_test_month(db_session, 2026, 7, actor=_actor(db_session), today=TODAY)

    months = register.awaiting_close(db_session, today=TODAY)

    # July is sealed; September has not ended.
    assert [(item.year, item.month) for item in months] == [(2026, 8)]
    assert months[0].pending_count == 1
    assert months[0].closable is False


# --------------------------------------------------------------------------- #
# the routes
# --------------------------------------------------------------------------- #


def test_month_route_serves_entries_counts_and_warnings(
    api_db: Session, operator_client: TestClient
) -> None:
    _closable_month(api_db)

    response = operator_client.get("/api/v1/inmate-violations/statistics/2026/8")

    assert response.status_code == 200
    body = response.json()
    assert body["counts"] == {"citizens": 1, "expats": 1, "pending": 0, "total": 2}
    assert body["closed"] is False
    assert "can_close" not in body
    assert body["workflow"]["state"] == "draft"
    assert body["first_closable_date"] == "2026-09-01"
    assert body["entries"][0]["id"].endswith(":0")


def test_direct_close_is_removed_and_reopen_rejects_missing_body(
    api_db: Session, operator_client: TestClient, admin_client: TestClient
) -> None:
    _closable_month(api_db)
    for client in (operator_client, admin_client):
        assert client.post(
            "/api/v1/inmate-violations/statistics/2026/8/close", json={}
        ).status_code in {404, 405}
    assert (
        admin_client.post("/api/v1/inmate-violations/statistics/2026/8/reopen").status_code == 422
    )


def test_manual_row_routes_round_trip_and_return_the_month(
    api_db: Session, operator_client: TestClient
) -> None:
    _reporter(api_db)

    created = operator_client.post(
        "/api/v1/inmate-violations/statistics/2026/8/manual-rows",
        json={
            "name": "اليدوي",
            "violation_date": "2026-08-09",
            "reason": "لا يوجد تقرير",
            "nationality_label": "الإمارات",
            "reporter_id": "G4603",
            "details_text": "تفاصيل",
        },
    )
    assert created.status_code == 201
    entry = created.json()["entries"][0]
    row_id = entry["manual"]["row_id"]
    assert entry["duty_unit"] == "السرية الثانية"

    patched = operator_client.patch(
        f"/api/v1/inmate-violations/statistics/manual-rows/{row_id}",
        json={"name": "اليدوي المعدل"},
    )
    assert patched.status_code == 200
    assert patched.json()["entries"][0]["name"] == "اليدوي المعدل"

    deleted = operator_client.delete(f"/api/v1/inmate-violations/statistics/manual-rows/{row_id}")
    assert deleted.status_code == 200
    assert deleted.json()["counts"]["total"] == 0


def test_tasks_route_is_accessible_to_non_admins(
    api_db: Session, operator_client: TestClient
) -> None:
    assert operator_client.get("/api/v1/inmate-violations/statistics/tasks").status_code == 200
    assert (
        "/api/v1/inmate-violations/statistics/awaiting-close"
        not in operator_client.app.openapi()["paths"]
    )


def test_nationality_list_route_carries_the_closed_list(operator_client: TestClient) -> None:
    body = operator_client.get("/api/v1/inmate-violations/nationalities").json()

    labels = {item["label_ar"]: item for item in body["items"]}
    assert labels["الإمارات"]["code"] == "AE"
    assert labels["بدون"]["selectable"] is True
    assert labels[register.UNSPECIFIED]["selectable"] is False


def test_export_downloads_a_workbook_without_closing_the_month(
    api_db: Session, operator_client: TestClient
) -> None:
    _closable_month(api_db)

    response = operator_client.get("/api/v1/inmate-violations/statistics/2026/8/export")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_PREFIX)
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    assert "المواطنون" in workbook.sheetnames
    # Unlike the timesheet, a download is never a seal.
    assert api_db.scalars(select(InmateViolationPeriod)).all() == []


def test_close_materialises_one_workbook_copy(api_db: Session, admin_client: TestClient) -> None:
    _closable_month(api_db)

    view = _approve_test_month(
        api_db, 2026, 8, actor=api_db.get(User, admin_client.user_id), today=TODAY
    )

    period = api_db.scalars(select(InmateViolationPeriod)).one()
    assert (
        period.export_path
        == f"inmate_violations/2026-08-submission-{view.workflow['active_submission_id']}.xlsx"
    )
    stored = get_settings().data_dir / period.export_path
    assert stored.is_file()
